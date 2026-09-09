# ไฟล์จริงที่ "หน้าตาไม่เหมือนไฟล์ตัวอย่าง" แล้วทำให้ข้อมูลเพี้ยน/หลุดแบบเงียบ ๆ
#
# เทสชุดนี้เกิดจากการยิงไฟล์จริงหลายรูปแบบเมื่อ 09/09/2026 แล้วเจอของที่พังโดยไม่มีอะไรฟ้อง
#
#   ① ชีทที่ผู้ใช้ "ซ่อนไว้" ใน Excel (เงินเดือน ต้นทุน สูตรราคา)
#      excel-csv และ excel-to-pdf เคยส่งออกให้ครบทุกตัวอักษรโดยไม่ถามและไม่บอก
#      = ทำข้อมูลที่ตั้งใจซ่อนหลุดออกไป ซึ่งร้ายกว่าแปลงไม่สำเร็จเสียอีก
#
#   ② รูปถ่ายจากมือถือที่หมุนด้วยแท็ก EXIF Orientation
#      images-to-pdf ฝังไบต์ JPEG ดิบลง PDF ตรง ๆ ซึ่ง pdf-lib ไม่รู้จักแท็กนี้เลย
#      รูปแนวตั้งจึงกลายเป็นแนวนอนล้มตะแคง (วัดได้จากสัดส่วนหน้า PDF)
#
#   ③ ย่อหน้าไทยยาว ๆ ที่ไม่มีเว้นวรรค (เอกสารราชการ/สัญญาเขียนแบบนี้เป็นปกติ)
#      jsPDF ตัดบรรทัดด้วย split(" ") เท่านั้น ทั้งย่อหน้าจึงเป็น "คำเดียว" แล้วถูกหั่น
#      ตรงไหนก็ได้ วัดจริงแล้ว "และ" ถูกฉีกเป็น "แล" + "ะ" และ "ค่าใช้จ่าย" ถูกฉีกจน
#      ค้นหาใน PDF ไม่เจอเลย · เกณฑ์ตรวจใช้ 2 ชั้น: บรรทัดต้องไม่จบด้วยสระนำโดด ๆ
#      (เ แ โ ใ ไ ซึ่งไม่มีทางเป็นตัวสุดท้ายของคำไทยได้) และคำทดสอบต้องค้นหาเจอครบ
#
#   ④ PDF ที่หน้าถูกหมุนไว้ (/Rotate 90/180/270 — ปกติมากสำหรับไฟล์สแกน)
#      pdf-page-numbers วางเลขด้วยพิกัดดิบซึ่งคนละแกนกับที่ตาเห็น เลขจึงไปโผล่ขอบซ้าย
#      ขอบขวา หรือหัวกระดาษ · ‼️ PyMuPDF คืนพิกัดข้อความแบบไม่คิด rotation ให้
#      ต้องคูณ page.rotation_matrix เองก่อนวัด ไม่งั้นจะอ่านผลผิดทั้งชุด
#
#   ⑤ PDF ที่ยังค้นหาข้อความได้ ถูกส่งเข้าเครื่องบีบอัด
#      เดิมวาดทุกหน้าใหม่เป็นภาพให้เอง = ไฟล์มักใหญ่ขึ้นและชั้นข้อความหายทั้งไฟล์
#      ตอนนี้ต้องบีบแบบไม่วาดใหม่ก่อน และถ้าทำไม่ได้ต้องหยุดถาม ไม่ทำลายของเงียบ ๆ
#
#   ⑥ รวมไฟล์ Word ที่ต่างมี เชิงอรรถ / อ้างอิงท้ายเรื่อง / คอมเมนต์
#      ทุกไฟล์ .docx เริ่มนับ w:id ของของพวกนี้ใหม่ที่เลขต่ำเหมือนกันหมด ถ้าไม่ออกเลขใหม่ตอนรวม
#      ตัวอ้างอิงของไฟล์หลังจะไปชี้เนื้อของไฟล์แรก = เนื้อหาสลับกันโดยไม่มีอะไรฟ้อง
#      และเนื้อของไฟล์หลังหายไปเลย (คอมเมนต์หายสนิท เชิงอรรถไปนอนเป็นไฟล์กำพร้าใน word/media/)
#
#   ⑨ หัวกระดาษ/ท้ายกระดาษของ Word (มักมีเลขที่หนังสือกับชั้นความลับ)
#      mammoth ไม่ส่งส่วนนี้มาให้เลย เดิมจึงหายทั้งหมดเงียบ ๆ ตอนแปลงเป็น PDF
#      และเลขหน้าใน Word เป็น field ถ้าไม่แทนค่าจะได้เลข 1 ซ้ำทุกหน้า
#
#   ⑩ รูปแบบรายการมีเลข (numbering.xml) ตอนรวมไฟล์ Word
#      ทุกไฟล์เริ่มนับ w:numId ใหม่ที่ 1 เหมือนกัน ถ้าไม่ออกเลขใหม่ตอนรวม ไฟล์หลังจะถูกบังคับ
#      ให้ใช้รูปแบบของไฟล์แรก (เช่นตั้งใจเป็น ก. ข. ค. แต่กลายเป็น 1. 2. 3.) โดยข้อความไม่หาย
#      จึงไม่มีอะไรฟ้อง ตัวนิยามซ้อนสองชั้น numId ชี้ไป abstractNumId ต้องออกเลขใหม่ทั้งคู่
#
#   ⑪ PDF ที่เข้ารหัสไว้ (รวมไฟล์ที่ล็อกแค่สิทธิ์ ซึ่งเปิดอ่านได้ปกติไม่ต้องใส่รหัส)
#      ไลบรารีที่ใช้เขียนไฟล์ไม่มีโค้ดถอดรหัสเลย ถ้าฝืนทำต่อจะได้ไฟล์ที่หน้าว่างเปล่า
#      หรือเปิดไม่ขึ้นเลย ทั้งที่สถานะขึ้นว่าสำเร็จ ต้องหยุดแล้วบอกวิธีแก้แทน
#
# ‼️ หน้าเว็บมี Content-Security-Policy ที่ไม่มี unsafe-eval — wait_for_function
#    ต้องส่งสตริงที่เป็นฟังก์ชันลูกศรเท่านั้น (มีเทสจับกฎนี้ใน accepts.test.mjs)
#
# รัน: ../.venv/bin/python tests/browser_realfiles.py     (เปิด/ปิด server ให้เอง)
#      ตั้ง FK_BASE=http://localhost:8899 ถ้ามี server อยู่แล้ว
import os
import subprocess
import sys
import pathlib
import shutil
import socket
import tempfile
import time

import re
import zipfile

import fitz
import openpyxl
from docx import Document
from PIL import Image
from playwright.sync_api import sync_playwright

ROOT = pathlib.Path(__file__).resolve().parent.parent
TMP = pathlib.Path(tempfile.mkdtemp(prefix="filekit_realfiles_"))
DL = TMP / "dl"
DL.mkdir()

P, F = 0, []


def ck(name, got, want):
    global P
    ok = got == want
    if ok:
        P += 1
    else:
        F.append(f"{name}\n      ได้    : {got!r}\n      ควรได้ : {want!r}")
    print(f"  {'✅' if ok else '❌'} {name}")


def free_port():
    s = socket.socket()
    s.bind(("127.0.0.1", 0))
    port = s.getsockname()[1]
    s.close()
    return port


def make_hidden_workbook():
    """xlsx ที่มีชีทเห็นปกติ 1 ชีท, ซ่อน 1 ชีท, ซ่อนลึก 1 ชีท"""
    wb = openpyxl.Workbook()
    s1 = wb.active
    s1.title = "รายชื่อ"
    s1["A1"], s1["A2"] = "ชื่อ", "สมชาย"
    s2 = wb.create_sheet("ลับ")
    s2["A1"], s2["A2"] = "เงินเดือน", "ห้ามเผยแพร่"
    s2.sheet_state = "hidden"
    s3 = wb.create_sheet("ลับมาก")
    s3["A1"] = "รหัสผ่านระบบ"
    s3.sheet_state = "veryHidden"
    p = TMP / "มีชีทซ่อน.xlsx"
    wb.save(p)
    return p


def make_rotated_photo():
    """JPEG ที่เก็บพิกเซลเป็นแนวนอน แล้วสั่งหมุนเป็นแนวตั้งด้วย EXIF (แบบที่มือถือถ่ายจริง)"""
    im = Image.new("RGB", (800, 600), (240, 240, 240))
    im.paste(Image.new("RGB", (200, 60), (20, 20, 200)), (0, 0))
    ex = Image.Exif()
    ex[0x0112] = 6          # 6 = หมุนตามเข็ม 90 องศาตอนแสดงผล
    p = TMP / "ถ่ายแนวตั้ง.jpg"
    im.save(p, exif=ex)
    return p


FIND_WORDS = ["ไม่เหมือนใคร", "ประสบการณ์", "ค่าใช้จ่าย", "รับประกัน"]

# ‼️ กฎที่ตัดสินได้แน่นอนว่า "ฉีกคำ" โดยไม่ต้องมีตัวตัดคำภาษาไทยในฝั่งเทส
#   · สระนำ (เ แ โ ใ ไ) เขียนไว้ "หน้า" พยัญชนะ จึงเป็นตัวท้ายบรรทัดไม่ได้เด็ดขาด
#   · สระบน/ล่าง วรรณยุกต์ ทัณฑฆาต และสระ ะ า ำ ต้องเกาะพยัญชนะที่มาก่อนเสมอ
#     จึงเป็นตัวแรกของบรรทัดไม่ได้เด็ดขาด
BAD_LINE_END = "เแโใไ"
BAD_LINE_START = "ะัาำิีึืุู็่้๊๋์ํๆฺๅ"


def broken_words(lines):
    return ([x for x in lines if x[-1] in BAD_LINE_END]
            + [x for x in lines if x[0] in BAD_LINE_START])

THAI_PARA = (
    "บริษัทของเราให้บริการที่ไม่เหมือนใครในตลาดเพราะเราใส่ใจรายละเอียดทุกขั้นตอน"
    "และมีทีมงานที่มีประสบการณ์มากกว่าสิบปีพร้อมให้คำปรึกษาแก่ลูกค้าทุกท่าน"
    "โดยไม่มีค่าใช้จ่ายเพิ่มเติมและรับประกันความพึงพอใจเต็มร้อย"
)


def make_thai_docx():
    """.docx ที่มีย่อหน้าไทยยาวไม่มีเว้นวรรคเลย แบบเอกสารราชการ/สัญญาจริง"""
    d = Document()
    d.add_paragraph(THAI_PARA)
    p = TMP / "ย่อหน้าไทย.docx"
    d.save(p)
    return p


def make_narrow_table():
    """xlsx หลายคอลัมน์จนแต่ละช่องแคบ บังคับให้ต้องตัดบรรทัดในเซลล์"""
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "งาน"
    ws.append(["ลำดับ", "รายการ", "หมายเหตุ", "ผู้รับผิดชอบ", "สถานะงาน",
               "หน่วยงาน", "งบประมาณ", "วันครบกำหนด"])
    ws.append([1, "ค่าบริการรายเดือน", THAI_PARA, "ฝ่ายบัญชีและการเงิน",
               "อยู่ระหว่างดำเนินการตรวจสอบเอกสาร", "สำนักงานใหญ่",
               "หนึ่งแสนสองหมื่นบาทถ้วน", "สามสิบเอ็ดธันวาคม"])
    ws.append([2, "ค่าติดตั้งอุปกรณ์", "ไม่มีค่าใช้จ่ายเพิ่มเติมเพราะรวมอยู่ในสัญญาแล้ว",
               "ฝ่ายเทคโนโลยีสารสนเทศ", "เสร็จสมบูรณ์แล้วรอปิดงาน", "สาขาภาคเหนือ",
               "ห้าหมื่นบาทถ้วน", "สิบห้ามกราคม"])
    p = TMP / "ตารางแคบ.xlsx"
    wb.save(p)
    return p


def make_rotated_pdf():
    """PDF ที่แต่ละหน้าถูกหมุนไว้คนละองศา แบบไฟล์สแกนที่ผ่านการหมุนแก้มาแล้ว"""
    d = fitz.open()
    for rot in (0, 90, 270, 180):
        page = d.new_page(width=595, height=842)
        page.insert_text((70, 100), f"rot {rot}", fontsize=18)
        page.set_rotation(rot)
    p = TMP / "หมุนสลับ.pdf"
    d.save(p)
    d.close()
    return p


def make_text_pdf(compact):
    """PDF ที่มีชั้นข้อความจริง · compact=True คือบันทึกแบบบีบโครงสร้างมาแล้ว
    (บีบแบบไม่เสียข้อความต่อไม่ได้ เครื่องมือต้องหยุดถามแทนที่จะวาดใหม่เป็นภาพ)"""
    d = fitz.open()
    for i in range(6):
        page = d.new_page()
        for k in range(28):
            page.insert_text((60, 80 + k * 22),
                             f"Line {k + 1} of page {i + 1} - searchable text that must survive",
                             fontsize=11)
    p = TMP / ("บีบมาแล้ว.pdf" if compact else "ข้อความล้วน.pdf")
    # use_objstms=1 คือรูปแบบเดียวกับที่ pdf-lib บันทึกออกมา (เช่นไฟล์ที่ผ่าน pdf-merge มาแล้ว)
    # บันทึกซ้ำจึงไม่เล็กลงอีก = เคสที่เครื่องมือต้องหยุดถาม ไม่ใช่วาดใหม่เป็นภาพให้เอง
    if compact:
        d.save(p, garbage=4, deflate=True, use_objstms=1)
    else:
        d.save(p)
    d.close()
    return p


def pages_with_text(path):
    doc = fitz.open(path)
    n = sum(1 for i in range(doc.page_count) if doc[i].get_text().strip())
    doc.close()
    return n


W_NS = "http://schemas.openxmlformats.org/wordprocessingml/2006/main"
R_NS = "http://schemas.openxmlformats.org/officeDocument/2006/relationships"


def make_docx_with_notes(tag):
    """.docx ที่มีทั้งเชิงอรรถ อ้างอิงท้ายเรื่อง และคอมเมนต์ โดยตั้ง w:id ให้ตรงกันทุกไฟล์
    (แบบที่ Word ทำจริง คือทุกไฟล์เริ่มนับใหม่จากเลขต่ำ) เพื่อบังคับให้ชนกันตอนรวม"""
    doc = ('<?xml version="1.0" encoding="UTF-8" standalone="yes"?>'
           f'<w:document xmlns:w="{W_NS}" xmlns:r="{R_NS}"><w:body>'
           '<w:p><w:commentRangeStart w:id="1"/>'
           f'<w:r><w:t>เนื้อความ {tag}</w:t></w:r><w:commentRangeEnd w:id="1"/>'
           '<w:r><w:commentReference w:id="1"/></w:r>'
           '<w:r><w:footnoteReference w:id="2"/></w:r>'
           '<w:r><w:endnoteReference w:id="2"/></w:r></w:p>'
           '<w:sectPr><w:pgSz w:w="11906" w:h="16838"/></w:sectPr></w:body></w:document>')
    fn = ('<?xml version="1.0" encoding="UTF-8" standalone="yes"?>'
          f'<w:footnotes xmlns:w="{W_NS}">'
          '<w:footnote w:type="separator" w:id="-1"><w:p><w:r><w:separator/></w:r></w:p></w:footnote>'
          f'<w:footnote w:id="2"><w:p><w:r><w:t>เชิงอรรถ {tag}</w:t></w:r></w:p></w:footnote></w:footnotes>')
    en = ('<?xml version="1.0" encoding="UTF-8" standalone="yes"?>'
          f'<w:endnotes xmlns:w="{W_NS}">'
          '<w:endnote w:type="separator" w:id="-1"><w:p><w:r><w:separator/></w:r></w:p></w:endnote>'
          f'<w:endnote w:id="2"><w:p><w:r><w:t>ท้ายเรื่อง {tag}</w:t></w:r></w:p></w:endnote></w:endnotes>')
    cm = ('<?xml version="1.0" encoding="UTF-8" standalone="yes"?>'
          f'<w:comments xmlns:w="{W_NS}">'
          '<w:comment w:id="1" w:author="ผู้ตรวจ" w:date="2026-09-09T00:00:00Z">'
          f'<w:p><w:r><w:t>คอมเมนต์ {tag}</w:t></w:r></w:p></w:comment></w:comments>')
    over = ('<Override PartName="/word/{0}.xml" ContentType="application/vnd.openxmlformats-'
            'officedocument.wordprocessingml.{0}+xml"/>')
    ct = ('<?xml version="1.0" encoding="UTF-8" standalone="yes"?>'
          '<Types xmlns="http://schemas.openxmlformats.org/package/2006/content-types">'
          '<Default Extension="rels" ContentType="application/vnd.openxmlformats-package.relationships+xml"/>'
          '<Default Extension="xml" ContentType="application/xml"/>'
          '<Override PartName="/word/document.xml" ContentType="application/vnd.openxmlformats-'
          'officedocument.wordprocessingml.document.main+xml"/>'
          + "".join(over.format(k) for k in ("footnotes", "endnotes", "comments")) + '</Types>')
    rels = ('<?xml version="1.0" encoding="UTF-8" standalone="yes"?>'
            '<Relationships xmlns="http://schemas.openxmlformats.org/package/2006/relationships">'
            f'<Relationship Id="rId1" Type="{R_NS}/officeDocument" Target="word/document.xml"/></Relationships>')
    drels = ('<?xml version="1.0" encoding="UTF-8" standalone="yes"?>'
             '<Relationships xmlns="http://schemas.openxmlformats.org/package/2006/relationships">'
             f'<Relationship Id="rId3" Type="{R_NS}/footnotes" Target="footnotes.xml"/>'
             f'<Relationship Id="rId4" Type="{R_NS}/endnotes" Target="endnotes.xml"/>'
             f'<Relationship Id="rId5" Type="{R_NS}/comments" Target="comments.xml"/></Relationships>')
    path = TMP / f"บันทึก-{tag}.docx"
    with zipfile.ZipFile(path, "w", zipfile.ZIP_DEFLATED) as z:
        z.writestr("[Content_Types].xml", ct)
        z.writestr("_rels/.rels", rels)
        z.writestr("word/document.xml", doc)
        z.writestr("word/_rels/document.xml.rels", drels)
        z.writestr("word/footnotes.xml", fn)
        z.writestr("word/endnotes.xml", en)
        z.writestr("word/comments.xml", cm)
    return path


def make_pptx_with_table_and_chart():
    """สไลด์ที่มีทั้งตารางและกราฟ แบบสไลด์สรุปยอดขายจริง"""
    from pptx import Presentation
    from pptx.util import Inches
    from pptx.chart.data import CategoryChartData
    from pptx.enum.chart import XL_CHART_TYPE
    pr = Presentation()
    sl = pr.slides.add_slide(pr.slide_layouts[5])
    sl.shapes.title.text = "สรุปยอดขายรายไตรมาส"
    data = [["ไตรมาส", "ยอดขาย", "เติบโต"],
            ["ไตรมาสหนึ่ง", "1,250,000", "+4.2%"],
            ["ไตรมาสสอง", "1,410,000", "+12.8%"]]
    tb = sl.shapes.add_table(3, 3, Inches(0.5), Inches(1.8), Inches(6), Inches(1.5)).table
    for r in range(3):
        for c in range(3):
            tb.cell(r, c).text = data[r][c]
    cd = CategoryChartData()
    cd.categories = ["ก", "ข", "ค"]
    cd.add_series("ชุดหนึ่ง", (1.0, 2.0, 3.0))
    sl.shapes.add_chart(XL_CHART_TYPE.COLUMN_CLUSTERED, Inches(7), Inches(1.8),
                        Inches(2.5), Inches(2), cd)
    p = TMP / "นำเสนอ-มีตาราง.pptx"
    pr.save(p)
    return p


def make_docx_with_textbox():
    """.docx ที่มีกล่องข้อความซ้อนอยู่ในย่อหน้า (ข้อความคนละที่บนหน้ากระดาษ)"""
    V_NS = "urn:schemas-microsoft-com:vml"
    doc = ('<?xml version="1.0" encoding="UTF-8" standalone="yes"?>'
           f'<w:document xmlns:w="{W_NS}" xmlns:r="{R_NS}" xmlns:v="{V_NS}"><w:body>'
           '<w:p><w:r><w:t>ยอดรวมทั้งสิ้น</w:t></w:r>'
           '<w:r><w:pict><v:shape id="s1" style="width:200pt;height:60pt"><v:textbox>'
           '<w:txbxContent><w:p><w:r><w:t>หมายเหตุในกล่อง</w:t></w:r></w:p>'
           '</w:txbxContent></v:textbox></v:shape></w:pict></w:r></w:p>'
           '<w:p><w:r><w:t>บรรทัดสุดท้าย</w:t></w:r></w:p>'
           '<w:sectPr><w:pgSz w:w="11906" w:h="16838"/></w:sectPr></w:body></w:document>')
    ct = ('<?xml version="1.0" encoding="UTF-8" standalone="yes"?>'
          '<Types xmlns="http://schemas.openxmlformats.org/package/2006/content-types">'
          '<Default Extension="rels" ContentType="application/vnd.openxmlformats-package.relationships+xml"/>'
          '<Default Extension="xml" ContentType="application/xml"/>'
          '<Override PartName="/word/document.xml" ContentType="application/vnd.openxmlformats-'
          'officedocument.wordprocessingml.document.main+xml"/></Types>')
    rels = ('<?xml version="1.0" encoding="UTF-8" standalone="yes"?>'
            '<Relationships xmlns="http://schemas.openxmlformats.org/package/2006/relationships">'
            f'<Relationship Id="rId1" Type="{R_NS}/officeDocument" Target="word/document.xml"/></Relationships>')
    path = TMP / "มีกล่องข้อความ.docx"
    with zipfile.ZipFile(path, "w", zipfile.ZIP_DEFLATED) as z:
        z.writestr("[Content_Types].xml", ct)
        z.writestr("_rels/.rels", rels)
        z.writestr("word/document.xml", doc)
        z.writestr("word/_rels/document.xml.rels",
                   '<?xml version="1.0"?><Relationships xmlns="http://schemas.openxmlformats.org/'
                   'package/2006/relationships"/>')
    return path


def make_docx_with_header_footer():
    """.docx ที่มีหัวกระดาษ และท้ายกระดาษที่ใส่เลขหน้าอัตโนมัติ (field PAGE)
    แบบเดียวกับหนังสือราชการไทยที่ใส่เลขที่หนังสือไว้บนหัวทุกหน้า"""
    from docx import Document
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    d = Document()
    sec = d.sections[0]
    sec.header.paragraphs[0].text = "เอกสารลับ เลขที่ กค 0123/2569"
    fp = sec.footer.paragraphs[0]
    fp.text = "ฝ่ายบัญชี หน้า "
    fld = OxmlElement("w:fldSimple")
    fld.set(qn("w:instr"), " PAGE ")
    r = OxmlElement("w:r")
    t = OxmlElement("w:t")
    t.text = "1"
    r.append(t)
    fld.append(r)
    fp._p.append(fld)
    for i in range(60):
        d.add_paragraph(f"ย่อหน้าที่ {i + 1} เนื้อความทดสอบภาษาไทยยาวพอให้เอกสารขึ้นหลายหน้า")
    p = TMP / "มีหัวท้าย.docx"
    d.save(p)
    return p


def make_docx_numbered(tag, fmt):
    """.docx ที่มีรายการมีเลข ใช้ w:numId=1 เหมือนกันทุกไฟล์ (แบบที่ Word ทำจริง)
    แต่รูปแบบต่างกัน เพื่อบังคับให้นิยามชนกันตอนรวมไฟล์"""
    body = "".join(
        '<w:p><w:pPr><w:numPr><w:ilvl w:val="0"/><w:numId w:val="1"/></w:numPr></w:pPr>'
        f'<w:r><w:t>ข้อ {i} ของ {tag}</w:t></w:r></w:p>' for i in (1, 2))
    doc = ('<?xml version="1.0" encoding="UTF-8" standalone="yes"?>'
           f'<w:document xmlns:w="{W_NS}" xmlns:r="{R_NS}"><w:body>{body}'
           '<w:sectPr><w:pgSz w:w="11906" w:h="16838"/></w:sectPr></w:body></w:document>')
    num = ('<?xml version="1.0" encoding="UTF-8" standalone="yes"?>'
           f'<w:numbering xmlns:w="{W_NS}">'
           '<w:abstractNum w:abstractNumId="0"><w:multiLevelType w:val="singleLevel"/>'
           f'<w:lvl w:ilvl="0"><w:start w:val="1"/><w:numFmt w:val="{fmt}"/>'
           '<w:lvlText w:val="%1."/></w:lvl></w:abstractNum>'
           '<w:num w:numId="1"><w:abstractNumId w:val="0"/></w:num></w:numbering>')
    over = ('<Override PartName="/word/{0}.xml" ContentType="application/vnd.openxmlformats-'
            'officedocument.wordprocessingml.{0}+xml"/>')
    ct = ('<?xml version="1.0" encoding="UTF-8" standalone="yes"?>'
          '<Types xmlns="http://schemas.openxmlformats.org/package/2006/content-types">'
          '<Default Extension="rels" ContentType="application/vnd.openxmlformats-package.relationships+xml"/>'
          '<Default Extension="xml" ContentType="application/xml"/>'
          '<Override PartName="/word/document.xml" ContentType="application/vnd.openxmlformats-'
          'officedocument.wordprocessingml.document.main+xml"/>' + over.format("numbering") + '</Types>')
    rels = ('<?xml version="1.0" encoding="UTF-8" standalone="yes"?>'
            '<Relationships xmlns="http://schemas.openxmlformats.org/package/2006/relationships">'
            f'<Relationship Id="rId1" Type="{R_NS}/officeDocument" Target="word/document.xml"/></Relationships>')
    drels = ('<?xml version="1.0" encoding="UTF-8" standalone="yes"?>'
             '<Relationships xmlns="http://schemas.openxmlformats.org/package/2006/relationships">'
             f'<Relationship Id="rId9" Type="{R_NS}/numbering" Target="numbering.xml"/></Relationships>')
    path = TMP / f"รายการ-{tag}.docx"
    with zipfile.ZipFile(path, "w", zipfile.ZIP_DEFLATED) as z:
        z.writestr("[Content_Types].xml", ct)
        z.writestr("_rels/.rels", rels)
        z.writestr("word/document.xml", doc)
        z.writestr("word/_rels/document.xml.rels", drels)
        z.writestr("word/numbering.xml", num)
    return path


def make_owner_locked_pdf():
    """PDF ที่ล็อกเฉพาะสิทธิ์ (owner password) เปิดอ่านได้ปกติไม่ต้องใส่รหัส
    แต่ไลบรารีที่เราใช้เขียนไฟล์ถอดรหัสไม่ได้ ถ้าฝืนทำต่อจะได้ไฟล์เสีย"""
    d = fitz.open()
    for i in range(3):
        page = d.new_page(width=595, height=842)
        page.insert_text((60, 100), f"page {i + 1} important content", fontsize=16)
    p = TMP / "ล็อกสิทธิ์.pdf"
    d.save(p, encryption=fitz.PDF_ENCRYPT_AES_256, owner_pw="owner123",
           permissions=fitz.PDF_PERM_ACCESSIBILITY)
    d.close()
    return p


def press(pg, pattern):
    """กดปุ่มที่มองเห็นจริงด้วย evaluate — ระหว่างที่เครื่องมือทำงานหนัก .click() ของ Playwright
    จะรอจนงานเสร็จก่อนค่อยกด ทำให้ทดสอบผิดเคสโดยไม่รู้ตัว (บทเรียน 09/09/2026)"""
    return pg.evaluate(
        "(re) => { const b = [...document.querySelectorAll('button')]"
        ".find((b) => b.offsetParent && new RegExp(re).test(b.textContent)); "
        "if (b) { b.click(); return b.textContent.trim(); } return null; }",
        pattern,
    )


def main():
    base = os.environ.get("FK_BASE")
    server = None
    if not base:
        port = free_port()
        server = subprocess.Popen([sys.executable, "-m", "http.server", str(port)],
                                  cwd=str(ROOT), stdout=subprocess.DEVNULL, stderr=subprocess.DEVNULL)
        base = f"http://localhost:{port}"
        time.sleep(1.5)

    xlsx = make_hidden_workbook()
    photo = make_rotated_photo()
    docx = make_thai_docx()
    narrow = make_narrow_table()
    rotated = make_rotated_pdf()
    note_docs = [make_docx_with_notes("AAA"), make_docx_with_notes("BBB")]
    deck = make_pptx_with_table_and_chart()
    tbox = make_docx_with_textbox()
    hf_doc = make_docx_with_header_footer()
    num_docs = [make_docx_numbered("AAA", "decimal"), make_docx_numbered("BBB", "thaiLetters")]
    locked = make_owner_locked_pdf()
    text_pdf = make_text_pdf(False)
    compact_pdf = make_text_pdf(True)

    try:
        with sync_playwright() as pw:
            b = pw.chromium.launch()
            ctx = b.new_context(viewport={"width": 1360, "height": 950},
                                reduced_motion="no-preference", accept_downloads=True)
            pg = ctx.new_page()
            errs = []
            pg.on("pageerror", lambda e: errs.append(str(e)[:160]))
            pg.on("console", lambda m: errs.append(m.text[:160]) if m.type == "error" else None)

            # ── ① ชีทที่ซ่อนไว้ ต้องไม่หลุดออกไปเอง ────────────────────────────────
            print("\n── ① ชีทที่ผู้ใช้ซ่อนไว้ ต้องไม่ถูกส่งออกโดยไม่ถาม ──")

            pg.goto(f"{base}/#/excel-csv", wait_until="networkidle")
            pg.wait_for_selector(".dz")
            pg.locator(".dz input[type=file]").set_input_files(str(xlsx))
            pg.wait_for_timeout(600)
            press(pg, "แปลงไฟล์")
            pg.wait_for_timeout(2500)
            names = pg.evaluate("() => [...document.querySelectorAll('.result .r-name strong')]"
                                ".map((e) => e.textContent)")
            ck("excel-csv แปลงเฉพาะชีทที่เห็นปกติ ไม่แตะชีทที่ซ่อนไว้",
               [n for n in names if "ลับ" in n], [])
            ck("excel-csv ได้ไฟล์ของชีทที่เห็นปกติจริง", len(names), 1)
            warn = pg.locator(".note.warn").inner_text() if pg.locator(".note.warn").count() else ""
            ck("excel-csv บอกผู้ใช้ตรง ๆ ว่าข้ามชีทที่ซ่อนไว้กี่ชีท ชื่ออะไร",
               ("ลับ" in warn and "ลับมาก" in warn), True)

            # ผู้ใช้ที่ต้องการจริง ๆ ต้องยังสั่งได้ ไม่ใช่ปิดตายไปเลย
            press(pg, "แปลงชีทที่ซ่อนด้วย")
            pg.wait_for_timeout(2500)
            names2 = pg.evaluate("() => [...document.querySelectorAll('.result .r-name strong')]"
                                 ".map((e) => e.textContent)")
            ck("กดขอเองแล้วได้ชีทที่ซ่อนมาครบ", len([n for n in names2 if "ลับ" in n]), 2)

            pg.goto(f"{base}/#/excel-to-pdf", wait_until="networkidle")
            pg.wait_for_selector(".dz")
            pg.locator(".dz input[type=file]").set_input_files(str(xlsx))
            pg.wait_for_timeout(600)
            press(pg, "แปลงเป็น PDF")
            pg.wait_for_timeout(4000)
            with pg.expect_download(timeout=20000) as info:
                pg.locator(".result button").first.click()
            out = DL / "sheets.pdf"
            info.value.save_as(str(out))
            doc = fitz.open(out)
            text = "\n".join(doc[i].get_text() for i in range(doc.page_count))
            doc.close()
            ck("excel-to-pdf ไม่พิมพ์เนื้อหาชีทที่ซ่อนไว้ลงหน้ากระดาษ", "ห้ามเผยแพร่" in text, False)
            ck("excel-to-pdf ไม่พิมพ์ชีทที่ซ่อนลึกด้วย", "รหัสผ่านระบบ" in text, False)
            ck("excel-to-pdf ยังพิมพ์ชีทที่เห็นปกติครบ", "สมชาย" in text, True)
            warn2 = pg.locator(".note.warn").inner_text() if pg.locator(".note.warn").count() else ""
            ck("excel-to-pdf บอกผู้ใช้ว่าข้ามชีทที่ซ่อนไว้", "ลับ" in warn2, True)

            # ── ② รูปที่หมุนด้วย EXIF ต้องเข้า PDF ถูกทิศ ──────────────────────────
            print("\n── ② รูปถ่ายที่หมุนด้วยแท็ก EXIF ──")
            pg.goto(f"{base}/#/images-to-pdf", wait_until="networkidle")
            pg.wait_for_selector(".dz")
            pg.locator(".dz input[type=file]").set_input_files(str(photo))
            pg.wait_for_timeout(1200)
            press(pg, "สร้างไฟล์ PDF")
            pg.wait_for_timeout(4000)
            with pg.expect_download(timeout=20000) as info:
                pg.locator(".result button").last.click()
            out2 = DL / "photo.pdf"
            info.value.save_as(str(out2))
            doc = fitz.open(out2)
            r = doc[0].rect
            doc.close()
            # ไฟล์เก็บพิกเซลไว้ 800x600 (แนวนอน) แต่ EXIF สั่งหมุนเป็น 600x800 (แนวตั้ง)
            ck("images-to-pdf วางรูปถ่ายแนวตั้งเป็นหน้าแนวตั้ง ไม่ล้มตะแคง", r.height > r.width, True)

            # ── ③ ย่อหน้าไทยยาวไม่มีเว้นวรรค ต้องตัดบรรทัดที่ขอบคำ ───────────────────
            print("\n── ③ ตัดบรรทัดภาษาไทยโดยไม่ฉีกคำ ──")

            def lines_of(pdf_path):
                doc = fitz.open(pdf_path)
                page = doc[0]
                out = []
                for blk in page.get_text("dict")["blocks"]:
                    for ln in blk.get("lines", []):
                        t = "".join(sp["text"] for sp in ln["spans"]).strip()
                        if t:
                            out.append(t)
                hits = {w: len(page.search_for(w)) for w in FIND_WORDS}
                doc.close()
                return out, hits

            pg.goto(f"{base}/#/word-to-pdf", wait_until="networkidle")
            pg.wait_for_selector(".dz")
            pg.locator(".dz input[type=file]").set_input_files(str(docx))
            pg.wait_for_timeout(1200)
            press(pg, "แปลง|สร้าง")
            pg.wait_for_timeout(5000)
            with pg.expect_download(timeout=25000) as info:
                pg.locator(".result button").last.click()
            out3 = DL / "thai.pdf"
            info.value.save_as(str(out3))
            lines, hits = lines_of(out3)
            ck("word-to-pdf ไม่มีบรรทัดไหนขึ้นต้น/ลงท้ายด้วยรูปที่ยืนเดี่ยวไม่ได้ (= ฉีกคำแน่นอน)",
               broken_words(lines), [])
            ck("word-to-pdf ค้นหาคำในเอกสารที่สร้างเจอครบทุกคำ",
               [w for w, n in hits.items() if n == 0], [])

            pg.goto(f"{base}/#/excel-to-pdf", wait_until="networkidle")
            pg.wait_for_selector(".dz")
            pg.locator(".dz input[type=file]").set_input_files(str(narrow))
            pg.wait_for_timeout(900)
            press(pg, "แปลงเป็น PDF")
            pg.wait_for_timeout(5000)
            with pg.expect_download(timeout=25000) as info:
                pg.locator(".result button").first.click()
            out4 = DL / "narrow.pdf"
            info.value.save_as(str(out4))
            lines2, hits2 = lines_of(out4)
            # ถ้าคอลัมน์กว้างพอจนไม่มีเซลล์ไหนต้องตัดบรรทัดเลย เทสนี้จะผ่านแบบหลอก ๆ
            # (8 คอลัมน์ x 3 แถว = 24 ช่อง ถ้าบรรทัดมากกว่านั้นแปลว่ามีการตัดบรรทัดจริง)
            ck("มีเซลล์ที่ถูกตัดบรรทัดจริงในไฟล์ทดสอบ (ประชากรต้องไม่เป็นศูนย์)",
               len(lines2) > 24, True)
            ck("excel-to-pdf ตัดบรรทัดในเซลล์แคบโดยไม่ฉีกคำไทย",
               broken_words(lines2), [])
            ck("excel-to-pdf ค้นหาคำในตารางที่สร้างเจอครบทุกคำ",
               [w for w, n in hits2.items() if n == 0], [])

            # ── ④ PDF ที่หน้าถูกหมุนไว้ เลขหน้าต้องอยู่ขอบล่างเสมอ ─────────────────────
            print("\n── ④ เลขหน้าบน PDF ที่หมุนไว้ ──")
            pg.goto(f"{base}/#/pdf-page-numbers", wait_until="networkidle")
            pg.wait_for_selector(".dz")
            pg.locator(".dz input[type=file]").set_input_files(str(rotated))
            pg.wait_for_timeout(2000)
            press(pg, "ใส่เลขหน้า|เพิ่มเลขหน้า")
            pg.wait_for_timeout(5000)
            with pg.expect_download(timeout=25000) as info:
                pg.locator(".result button").first.click()
            out5 = DL / "numbered.pdf"
            info.value.save_as(str(out5))
            doc = fitz.open(out5)
            spots = []
            for i in range(doc.page_count):
                page = doc[i]
                r = page.rect
                box = None
                for blk in page.get_text("dict")["blocks"]:
                    for ln in blk.get("lines", []):
                        if "".join(sp["text"] for sp in ln["spans"]).strip().isdigit():
                            box = ln["bbox"]
                # ‼️ PyMuPDF คืนพิกัดแบบไม่คิด rotation ต้องแปลงเองก่อนถึงจะวัดตรงกับที่ตาเห็น
                if box:
                    # ‼️ อย่าใช้ชื่อ b ตรงนี้ มันไปทับตัวแปรเบราว์เซอร์จนปิดไม่ลงตอนจบ
                    vb = tuple(fitz.Rect(box) * page.rotation_matrix)
                    cy, cx = (vb[1] + vb[3]) / 2, (vb[0] + vb[2]) / 2
                    spots.append((page.rotation,
                                  "ล่าง" if cy > r.height * 0.8 else "ไม่ใช่ล่าง",
                                  "กลาง" if r.width * 0.3 < cx < r.width * 0.7 else "ไม่ใช่กลาง"))
                else:
                    spots.append((page.rotation, "ไม่พบเลข", "ไม่พบเลข"))
            doc.close()
            ck("pdf-page-numbers วางเลขที่ขอบล่างกลางหน้าเสมอ ไม่ว่าหน้าจะถูกหมุนไว้กี่องศา",
               spots, [(0, "ล่าง", "กลาง"), (90, "ล่าง", "กลาง"),
                       (270, "ล่าง", "กลาง"), (180, "ล่าง", "กลาง")])

            # ── ⑤ บีบอัด PDF ต้องไม่ทำลายชั้นข้อความเงียบ ๆ ────────────────────────────
            print("\n── ⑤ บีบอัด PDF ที่ยังค้นหาข้อความได้ ──")
            pg.goto(f"{base}/#/pdf-compress", wait_until="networkidle")
            pg.wait_for_selector(".dz")
            pg.locator(".dz input[type=file]").first.set_input_files(str(text_pdf))
            pg.wait_for_timeout(2500)
            press(pg, "บีบอัดไฟล์")
            pg.wait_for_selector(".results .result, .results .note.warn", timeout=40000)
            ck("บีบไฟล์ข้อความแล้วได้ไฟล์ออกมาจริง", pg.locator(".results .result").count() >= 1, True)
            with pg.expect_download(timeout=25000) as info:
                pg.locator(".results .result button").first.click()
            out6 = DL / "compressed.pdf"
            info.value.save_as(str(out6))
            ck("บีบแล้วชั้นข้อความยังอยู่ครบทุกหน้า", pages_with_text(out6), 6)
            ck("บีบแล้วไฟล์เล็กลงจริง", out6.stat().st_size < text_pdf.stat().st_size, True)
            ck("บอกผู้ใช้ว่าข้อความยังค้นหาได้",
               "ค้นหา" in pg.locator(".status").inner_text(), True)

            # ไฟล์ที่บีบโครงสร้างมาแล้ว: ห้ามวาดใหม่เป็นภาพให้เอง ต้องหยุดถามก่อน
            pg.goto(f"{base}/#/pdf-compress", wait_until="networkidle")
            pg.wait_for_selector(".dz")
            pg.locator(".dz input[type=file]").first.set_input_files(str(compact_pdf))
            pg.wait_for_timeout(2500)
            press(pg, "บีบอัดไฟล์")
            pg.wait_for_selector(".results .result, .results .note.warn", timeout=40000)
            ck("ไฟล์ที่บีบต่อไม่ได้โดยไม่เสียข้อความ ต้องไม่ยัดไฟล์ที่ข้อความหายให้เอง",
               pg.locator(".results .result").count(), 0)
            ck("ต้องมีปุ่มให้ผู้ใช้เลือกเองว่าจะยอมเสียข้อความไหม",
               pg.locator(".results .note.warn button").count(), 1)
            pg.evaluate("() => document.querySelector('.note.warn button').click()")
            pg.wait_for_selector(".results .result", timeout=40000)
            with pg.expect_download(timeout=25000) as info:
                pg.locator(".results .result button").first.click()
            out7 = DL / "forced.pdf"
            info.value.save_as(str(out7))
            ck("กดยอมแล้วยังบีบให้ได้จริง (ไม่ได้ปิดทางผู้ใช้)", out7.stat().st_size > 0, True)

            # ── ⑥ รวมไฟล์ Word ต้องไม่ทำเชิงอรรถ/คอมเมนต์ของไฟล์หลังหาย ────────────────
            print("\n── ⑥ รวมไฟล์ Word ที่มีเชิงอรรถ อ้างอิงท้ายเรื่อง และคอมเมนต์ ──")
            pg.goto(f"{base}/#/word-join", wait_until="networkidle")
            pg.wait_for_selector(".dz")
            pg.locator(".dz input[type=file]").first.set_input_files([str(x) for x in note_docs])
            pg.wait_for_timeout(1500)
            press(pg, "รวมไฟล์|รวมเป็น")
            pg.wait_for_selector(".result", timeout=60000)
            pg.wait_for_timeout(1200)
            with pg.expect_download(timeout=30000) as info:
                pg.locator(".result button").last.click()
            out8 = DL / "joined.docx"
            info.value.save_as(str(out8))
            with zipfile.ZipFile(out8) as z:
                names = z.namelist()
                body = z.read("word/document.xml").decode()
                parts = {n: z.read(n).decode() for n in
                         ("word/footnotes.xml", "word/endnotes.xml", "word/comments.xml") if n in names}
            for kind, part, item, ref, txt in [
                ("เชิงอรรถ", "word/footnotes.xml", "footnote", "footnoteReference", "เชิงอรรถ"),
                ("อ้างอิงท้ายเรื่อง", "word/endnotes.xml", "endnote", "endnoteReference", "ท้ายเรื่อง"),
                ("คอมเมนต์", "word/comments.xml", "comment", "commentReference", "คอมเมนต์"),
            ]:
                inner = parts.get(part, "")
                refs = re.findall(r'<w:%s[^>]*w:id="(-?\d+)"' % ref, body)
                ids = re.findall(r'<w:%s [^>]*w:id="(-?\d+)"' % item, inner)
                ck(f"{kind}: เนื้อของทั้ง 2 ไฟล์อยู่ครบ",
                   (f"{txt} AAA" in inner, f"{txt} BBB" in inner), (True, True))
                ck(f"{kind}: จุดอ้างอิง 2 จุดชี้คนละ id", len(refs) == 2 and refs[0] != refs[1], True)
                ck(f"{kind}: ทุกจุดอ้างอิงมีเนื้อรองรับจริง", all(r in ids for r in refs), True)
            rng = sorted(set(re.findall(r'<w:commentRange(?:Start|End)[^>]*w:id="(-?\d+)"', body)))
            cref = sorted(set(re.findall(r'<w:commentReference[^>]*w:id="(-?\d+)"', body)))
            ck("คอมเมนต์: หัวและท้ายช่วงที่คลุมไว้ถูกเปลี่ยนเลขตามไปด้วย", rng, cref)
            ck("ไม่มีไฟล์กำพร้าหลุดไปอยู่ใน word/media/",
               [n for n in names if n.startswith("word/media/") and n.endswith(".xml")], [])

            # ── ⑦ ตารางในสไลด์ และการแทนที่คำข้ามขอบเขตกล่องข้อความ ────────────────────
            print("\n── ⑦ ตารางในสไลด์ PowerPoint ──")
            pg.goto(f"{base}/#/powerpoint-to-pdf", wait_until="networkidle")
            pg.wait_for_selector(".dz")
            pg.locator(".dz input[type=file]").first.set_input_files(str(deck))
            pg.wait_for_timeout(1500)
            press(pg, "แปลง|สร้าง")
            pg.wait_for_selector(".result", timeout=60000)
            status = pg.locator(".status").inner_text()
            with pg.expect_download(timeout=30000) as info:
                pg.locator(".result button").last.click()
            out9 = DL / "deck.pdf"
            info.value.save_as(str(out9))
            doc = fitz.open(out9)
            deck_text = doc[0].get_text()
            doc.close()
            ck("ข้อความในตารางบนสไลด์ไม่หาย",
               [w for w in ("ไตรมาสหนึ่ง", "1,250,000", "+12.8%", "เติบโต") if w not in deck_text], [])
            ck("บอกผู้ใช้ว่ามีกราฟที่แปลงเป็นข้อความไม่ได้", "กราฟ" in status, True)

            print("\n── ⑧ แทนที่คำ ห้ามข้ามขอบเขตกล่องข้อความ ──")
            pg.goto(f"{base}/#/word-replace", wait_until="networkidle")
            pg.wait_for_selector(".dz")
            pg.locator(".dz input[type=file]").first.set_input_files(str(tbox))
            pg.wait_for_timeout(1200)
            boxes = pg.locator("input[type=text]")
            # คำนี้ "ไม่มีอยู่จริง" ในเอกสาร มันคร่อมรอยต่อระหว่างข้อความนอกกล่องกับในกล่อง
            boxes.nth(0).fill("ทั้งสิ้นหมายเหตุ")
            boxes.nth(1).fill("XXX")
            pg.wait_for_timeout(500)
            press(pg, "แทนที่|เริ่ม")
            pg.wait_for_timeout(3000)
            ck("ไม่แทนที่คำที่ไม่มีอยู่จริง (คร่อมขอบเขตกล่องข้อความ)",
               "0 จุด" in pg.locator(".status").inner_text(), True)
            if pg.locator(".result button").count():
                with pg.expect_download(timeout=30000) as info:
                    pg.locator(".result button").last.click()
                out10 = DL / "rep.docx"
                info.value.save_as(str(out10))
                with zipfile.ZipFile(out10) as z:
                    xml = z.read("word/document.xml").decode()
                ck("ข้อความในกล่องข้อความไม่ถูกลบทิ้ง", "หมายเหตุในกล่อง" in xml, True)
                ck("ข้อความนอกกล่องไม่ถูกลบทิ้ง", "ยอดรวมทั้งสิ้น" in xml, True)

            # ── ⑨ หัวกระดาษ/ท้ายกระดาษของ Word ต้องติดไปทุกหน้า ──────────────────────
            print("\n── ⑨ หัวกระดาษและท้ายกระดาษ ──")
            pg.goto(f"{base}/#/word-to-pdf", wait_until="networkidle")
            pg.wait_for_selector(".dz")
            pg.locator(".dz input[type=file]").first.set_input_files(str(hf_doc))
            pg.wait_for_timeout(1200)
            press(pg, "แปลง|สร้าง")
            pg.wait_for_selector(".result", timeout=60000)
            with pg.expect_download(timeout=30000) as info:
                pg.locator(".result button").last.click()
            out11 = DL / "headfoot.pdf"
            info.value.save_as(str(out11))
            doc = fitz.open(out11)
            pages = doc.page_count
            per_page = [doc[i].get_text() for i in range(pages)]
            top3 = []
            d0 = doc[0].get_text("dict")
            rows = sorted((ln["bbox"][1], "".join(sp["text"] for sp in ln["spans"]))
                          for blk in d0["blocks"] for ln in blk.get("lines", []))
            top3 = [t for _, t in rows[:2]]
            doc.close()
            ck("เอกสารทดสอบยาวหลายหน้าจริง (ประชากรต้องไม่เป็นศูนย์)", pages >= 2, True)
            ck("หัวกระดาษติดไปครบทุกหน้า",
               [i + 1 for i, t in enumerate(per_page) if "เอกสารลับ" not in t], [])
            ck("ท้ายกระดาษติดไปครบทุกหน้า",
               [i + 1 for i, t in enumerate(per_page) if "ฝ่ายบัญชี" not in t], [])
            # ‼️ เลขหน้าเป็น field ของ Word ถ้าไม่แทนค่าจะได้เลข 1 ซ้ำทุกหน้า
            ck("เลขหน้าอัตโนมัติเดินตามหน้าจริง ไม่ใช่เลขเดิมซ้ำ",
               [i + 1 for i, t in enumerate(per_page) if f"หน้า {i + 1}" not in t], [])
            ck("หัวกระดาษอยู่เหนือเนื้อหา ไม่ทับบรรทัดแรก",
               top3[0].startswith("เอกสารลับ") and "ย่อหน้าที่ 1" in top3[1], True)

            # ── ⑩ รวมไฟล์ Word ที่รูปแบบรายการมีเลขชนกัน ────────────────────────────
            print("\n── ⑩ รูปแบบรายการมีเลขตอนรวมไฟล์ ──")
            pg.goto(f"{base}/#/word-join", wait_until="networkidle")
            pg.wait_for_selector(".dz")
            pg.locator(".dz input[type=file]").first.set_input_files([str(x) for x in num_docs])
            pg.wait_for_timeout(1500)
            press(pg, "รวมไฟล์|รวมเป็น")
            pg.wait_for_selector(".result", timeout=60000)
            pg.wait_for_timeout(1200)
            with pg.expect_download(timeout=30000) as info:
                pg.locator(".result button").last.click()
            out12 = DL / "numbered.docx"
            info.value.save_as(str(out12))
            with zipfile.ZipFile(out12) as z:
                names = z.namelist()
                nbody = z.read("word/document.xml").decode()
                ndef = z.read("word/numbering.xml").decode() if "word/numbering.xml" in names else ""
            fmts = re.findall(r'<w:numFmt w:val="([^"]+)"', ndef)
            defined = re.findall(r'<w:num w:numId="(\d+)"', ndef)
            used = {}
            for m in re.finditer(r'<w:numId w:val="(\d+)"/></w:numPr></w:pPr>'
                                 r'<w:r><w:t>ข้อ \d+ ของ (\w+)</w:t>', nbody):
                used.setdefault(m.group(2), set()).add(m.group(1))
            ck("เก็บรูปแบบรายการของทั้ง 2 ไฟล์ไว้ ไม่บังคับใช้ของไฟล์แรก",
               sorted(set(fmts)), ["decimal", "thaiLetters"])
            ck("สองไฟล์อ้างคนละหมายเลขรายการ ไม่ชนกัน",
               sorted(used.get("AAA", set()) & used.get("BBB", set())), [])
            ck("ทุกหมายเลขรายการที่เนื้อหาอ้างถึง มีนิยามรองรับจริง",
               [x for v in used.values() for x in v if x not in defined], [])

            # ── ⑪ PDF ที่เข้ารหัสไว้ ต้องหยุดแล้วบอก ไม่ใช่ยัดไฟล์เสียให้ ────────────────
            print("\n── ⑪ PDF ที่เข้ารหัสไว้ ──")
            for tool, btn in [("pdf-split", "แยกไฟล์|แยก"),
                              ("pdf-page-numbers", "ใส่เลขหน้า|เพิ่มเลขหน้า")]:
                pg.goto("about:blank")
                pg.goto(f"{base}/#/{tool}", wait_until="networkidle")
                pg.wait_for_selector(".dz")
                pg.locator(".dz input[type=file]").first.set_input_files(str(locked))
                pg.wait_for_timeout(2500)
                press(pg, btn)
                pg.wait_for_timeout(5000)
                says = " ".join(pg.locator(".status").all_inner_texts())
                ck(f"{tool}: ไม่ยัดไฟล์ที่เปิดไม่ขึ้นให้ผู้ใช้",
                   pg.locator(".result button").count(), 0)
                ck(f"{tool}: บอกตรง ๆ ว่าเข้ารหัสไว้ และบอกวิธีแก้",
                   ("เข้ารหัส" in says and "ปลดล็อก" in says), True)
                ck(f"{tool}: ไม่ขึ้นข้อความ Error แบบดิบ ๆ", "Error:" in says, False)

            print("\n── ไม่มี error หลุดออกมา ──")
            ck("ไม่มี console หรือ page error ตลอดทั้งชุด", errs, [])
            ctx.close()
            b.close()
    finally:
        if server:
            server.terminate()
            server.wait(timeout=10)


if __name__ == "__main__":
    try:
        main()
    finally:
        shutil.rmtree(TMP, ignore_errors=True)
    print(f"\nผ่าน {P}, ตก {len(F)}")
    for i, x in enumerate(F, 1):
        print(f"  {i}. {x}")
    sys.exit(1 if F else 0)
