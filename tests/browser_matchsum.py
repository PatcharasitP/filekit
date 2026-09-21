# เครื่องมือ "หายอดที่บวกกันได้เท่านี้" ตรวจของที่ผู้ใช้เอาไปใช้ต่อจริง ไม่ใช่แค่ว่ากดแล้วไม่ error
#
# ‼️ สองข้อที่เคยพังจริงและต้องกันไม่ให้กลับมา (เจอ 15/09/2026 ตอนลองกับไฟล์จริง 14,656 แถว):
#    ① เลขแถวเลื่อน — ตัวอ่านชีตกลางของโปรเจกต์สั่ง blankrows:false แถวที่ทั้งแถวเป็น #N/A
#       จึงหลุดไปทั้งแถว ทำให้ "แถว 14591" กลายเป็น "แถว 14586" งานกระทบยอดต้องเปิดไฟล์ไปดูแถวนั้นต่อ
#       เลขแถวผิดคือผิดทั้งงาน และผิดแบบเงียบ ไม่มีอะไรฟ้อง
#    ② โหมดวางตัวเลขเอง ปุ่มค้นหากดไม่ได้ เพราะกล่องลากไฟล์ปิดปุ่มหลักไว้จนกว่าจะมีไฟล์
#
# ถ้าผิดจะรู้ได้ยังไง: ไฟล์ทดสอบจงใจใส่ #N/A ไว้ก่อนแถวคำตอบ ถ้าใครเผลอกลับไปใช้ตัวอ่านที่ตัดแถวทิ้ง
# ข้อ ① จะแดงทันทีพร้อมเลขแถวที่คลาด
#
# รัน: ../.venv/bin/python tests/browser_matchsum.py
import os
import pathlib
import socket
import subprocess
import sys
import tempfile
import time

import openpyxl
from playwright.sync_api import sync_playwright

sys.path.insert(0, str(pathlib.Path(__file__).resolve().parent))
import fkui  # noqa: E402  ตัวช่วยกลางที่รู้จักโครงหน้าเครื่องมือ v2

ROOT = pathlib.Path(__file__).resolve().parent.parent
TMP = pathlib.Path(tempfile.mkdtemp(prefix="filekit_msum_"))
P, F = 0, []


def ck(name, got, want):
    global P
    ok = got == want
    if ok:
        P += 1
    else:
        F.append(f"{name}\n      ได้    : {got!r}\n      ควรได้ : {want!r}")
    print(f"  {'✅' if ok else '❌'} {name}")


def ok(name, cond, note=""):
    global P
    if cond:
        P += 1
    else:
        F.append(f"{name} {note}")
    print(f"  {'✅' if cond else '❌'} {name}{(' ' + note) if note and not cond else ''}")


def free_port():
    s = socket.socket()
    s.bind(("127.0.0.1", 0))
    port = s.getsockname()[1]
    s.close()
    return port


# คอลัมน์ตัวเลขล้วน ไม่มีหัวตาราง แบบที่ไฟล์กระทบยอดจริงเป็น
# แถว 1-5 ของจริง, แถว 6-7 เป็น #N/A (กับดักเลขแถวเลื่อน), คำตอบซ่อนอยู่หลังจากนั้น
VALUES = [
    100,          # แถว 1
    250.50,       # แถว 2
    "#N/A",       # แถว 3
    7000,         # แถว 4
    49.50,        # แถว 5
    "#N/A",       # แถว 6
    "#N/A",       # แถว 7
    12,           # แถว 8
    88,           # แถว 9
    -30,          # แถว 10  (ค่าติดลบ)
    300,          # แถว 11
    300,          # แถว 12  (ค่าซ้ำกับแถว 11)
    1_500_000,    # แถว 13  (ก้อนใหญ่ ใช้ทดสอบการคัดช่วงค่า)
]
# ชุดที่รวมได้ 400 พอดี (ไม่เกิน 4 รายการ ไม่ใช้ค่าติดลบ):
#   300 + 100, 300 + 88 + 12, 250.50 + 100 + 49.50, 250.50 + 88 + 49.50 + 12
def key_of(*vals):
    return "+".join(str(float(v)) if float(v) % 1 else str(int(v)) for v in sorted(vals))


EXPECT_SETS = {key_of(300, 100), key_of(300, 88, 12), key_of(250.5, 100, 49.5),
               key_of(250.5, 88, 49.5, 12)}


def make_file():
    wb = openpyxl.Workbook()
    ws = wb.active
    for v in VALUES:
        ws.append([v])
    p = TMP / "ยอดกระทบ.xlsx"
    wb.save(p)
    return p


def sets_on_screen(pg):
    """ชุดคำตอบที่แสดงบนจอ คืนเป็น {"ค่า+ค่า": [เลขแถว...]} เรียงค่าแล้วเพื่อเทียบง่าย"""
    return pg.evaluate("""() => {
      const out = {};
      for (const card of document.querySelectorAll('.ms-card')) {
        const trs = [...card.querySelectorAll('.ms-rows tr')];
        const body = trs.slice(0, -1).map(tr => {
          const td = tr.querySelectorAll('td');
          return { v: td[0].textContent.replace(/,/g, ''), where: td[1].textContent };
        });
        body.sort((a, b) => +a.v - +b.v);
        const key = body.map(b => String(+b.v)).join('+');
        out[key] = body.map(b => +(b.where.match(/(\\d+)\\s*$/) || [0, 0])[1]);
      }
      return out;
    }""")


def main():
    base = os.environ.get("FK_BASE")
    server = None
    if not base:
        port = free_port()
        server = subprocess.Popen([sys.executable, "-m", "http.server", str(port)],
                                  cwd=str(ROOT), stdout=subprocess.DEVNULL, stderr=subprocess.DEVNULL)
        base = f"http://localhost:{port}"
        time.sleep(1.5)

    xlsx = make_file()
    try:
        with sync_playwright() as pw:
            b = pw.chromium.launch()
            ctx = b.new_context(viewport={"width": 1360, "height": 1000}, accept_downloads=True,
                                timezone_id="Asia/Bangkok")
            pg = ctx.new_page()
            errs = []
            pg.on("pageerror", lambda e: errs.append(str(e)[:160]))
            pg.on("console", lambda m: errs.append(m.text[:160]) if m.type == "error" else None)

            pg.goto(f"{base}/#/excel-match-sum", wait_until="networkidle")
            pg.wait_for_selector(".dz")
            pg.locator(".dz input[type=file]").first.set_input_files(str(xlsx))
            pg.wait_for_timeout(1800)

            print("\n── ① อ่านไฟล์ที่ไม่มีหัวตาราง ──")
            chips = pg.evaluate("() => [...document.querySelectorAll('.ms-chip')].map(c => c.textContent)")
            nums = [v for v in VALUES if not isinstance(v, str)]
            ok("รู้เองว่าแถวแรกเป็นข้อมูล ไม่ใช่ชื่อคอลัมน์",
               pg.evaluate("() => document.querySelector('.s2-side-bd input[type=checkbox], .s2-stage input[type=checkbox], .ws-left input[type=checkbox]').checked") is False)
            ok(f"อ่านตัวเลขได้ครบ {len(nums)} แถว",
               any(f"{len(nums)}" in c and "อ่านเป็นตัวเลข" in c for c in chips), f"ได้ {chips}")
            ok("นับแถวที่อ่านไม่ออกแยกไว้ 3 แถว",
               any("อ่านไม่ออก" in c and "3" in c for c in chips), f"ได้ {chips}")

            print("\n── ② ชุดที่รวมกันได้ ต้องครบและถูกทุกชุด ──")
            pg.evaluate("""() => {
              document.querySelectorAll('.s2-side-bd input.ms-num, .s2-stage input.ms-num, .ws-right input.ms-num')[0].value = '400';
            }""")
            fkui.act(pg, "ค้นหา")
            pg.wait_for_timeout(1200)
            found = sets_on_screen(pg)
            ck("เจอครบทุกชุดที่เป็นไปได้ ไม่ขาดไม่เกิน", set(found), EXPECT_SETS)

            print("\n── ③ เลขแถวต้องตรงกับไฟล์จริง (แถว #N/A ต้องไม่ทำให้เลื่อน) ──")
            # แผนที่ค่า -> เลขแถวจริงในไฟล์ (นับจาก 1 เหมือนที่ผู้ใช้เห็นใน Excel)
            real = {}
            for i, v in enumerate(VALUES, 1):
                if not isinstance(v, str):
                    real.setdefault(str(float(v)) if float(v) % 1 else str(int(v)), []).append(i)
            bad = []
            for key, rows in found.items():
                for val, row in zip(key.split("+"), rows):
                    if row not in real.get(val, []):
                        bad.append(f"{val} ชี้ไปแถว {row} แต่ของจริงอยู่แถว {real.get(val)}")
            ck("ทุกค่าที่โชว์ ชี้ไปแถวที่มีค่านั้นจริง", bad, [])
            ok("ค่า 300 ที่ซ้ำสองแถว ชี้ไปแถว 11 หรือ 12 เท่านั้น",
               all(r in (11, 12) for k, rows in found.items() for v, r in zip(k.split("+"), rows) if v == "300"))

            print("\n── ④ เงื่อนไขที่ผู้ใช้ปรับได้ ต้องมีผลจริง ──")
            def search(target="400", tol="0", lo="1", hi="4", minv="", maxv="", neg=False):
                # ‼️ ดาวน์โหลดแล้วหน้าเข้าสถานะผลลัพธ์ ปุ่มค้นหาหายไปจากจอ ผู้ใช้ต้องกด "กลับไปแก้" ก่อน
                fkui.back_to_work(pg)
                pg.evaluate("""([t, tol, lo, hi, minv, maxv, neg]) => {
                  const n = document.querySelectorAll('.s2-side-bd input.ms-num, .s2-stage input.ms-num, .ws-right input.ms-num');
                  n[0].value = t; n[1].value = tol; n[2].value = lo; n[3].value = hi;
                  n[4].value = minv; n[5].value = maxv;
                  /* ‼️ หาสวิตช์จากชื่อที่ผู้ใช้เห็น ห้ามนับลำดับ (22/09/2026)
                     v2 รวมแผงซ้ายกับขวาไว้แผงเดียว สวิตช์ตัวแรกจึงกลายเป็น "แถวแรกเป็นหัวตาราง"
                     เทสเลยไปสลับหัวตารางแทนค่าติดลบ แล้วได้คำตอบผิดชุดโดยไม่มีอะไรฟ้อง */
                  const row = [...document.querySelectorAll('.ms-switch-field')]
                    .find((r) => r.textContent.includes('ติดลบ'));
                  const sw = row.querySelector('input[type=checkbox]');
                  if (sw.checked !== neg) sw.click();
                }""", [target, tol, lo, hi, minv, maxv, neg])
                fkui.act(pg, "ค้นหา")
                pg.wait_for_timeout(900)
                return sets_on_screen(pg)

            two = search(hi="2")
            ck("จำกัดไม่เกิน 2 รายการ เหลือเฉพาะคู่", set(two), {key_of(300, 100)})
            three = search(lo="3", hi="3")
            ck("บังคับ 3 รายการพอดี", set(three), {key_of(300, 88, 12), key_of(250.5, 100, 49.5)})
            tol = search(target="401", tol="1", hi="2")
            ok("ยอมคลาดเคลื่อน ±1 ทำให้ 400 เข้าเกณฑ์กับเป้า 401", key_of(300, 100) in tol, f"ได้ {list(tol)}")
            neg = search(target="370", hi="3", neg=True)
            ok("เปิดใช้ค่าติดลบแล้ว -30 ถูกนำมาใช้",
               any("-30" in k for k in neg), f"ได้ {list(neg)}")
            cut = search(target="1500000", hi="1", maxv="1000000")
            ck("คัดรายการเกินหนึ่งล้านออก จึงหาไม่เจอ", set(cut), set())

            print("\n── ⑤ ไม่เจอ ต้องบอกว่าไม่เจอ ──")
            none = search(target="999999", hi="4")
            ck("ไม่มีชุดไหนรวมได้ จึงไม่โชว์ชุดมั่ว", set(none), set())
            msg = pg.evaluate("() => document.querySelector('.ms-banner')?.textContent || ''")
            ok("บอกด้วยว่าค้นครบทุกความเป็นไปได้แล้วหรือยัง",
               ("ค้นครบทุกความเป็นไปได้" in msg) or ("หมดเวลา" in msg), f"ได้ {msg[:80]!r}")

            print("\n── ⑥ ดาวน์โหลดผลเป็น Excel แล้วค่าต้องตรงกับที่โชว์ ──")
            shown = search()
            with pg.expect_download() as dl:
                fkui.dl_button(pg).click()
            path = TMP / "ผลลัพธ์.xlsx"
            dl.value.save_as(str(path))
            wb = openpyxl.load_workbook(path)
            ws = wb.active
            rows = list(ws.iter_rows(min_row=2, values_only=True))
            in_file = {}
            for r in rows:
                in_file.setdefault(r[0], []).append(str(float(r[2])) if float(r[2]) % 1 else str(int(r[2])))
            ck("ไฟล์ที่ดาวน์โหลดมีชุดเท่าที่โชว์บนจอ", len(in_file), len(shown))
            ck("ทุกชุดในไฟล์ตรงกับชุดบนจอ",
               sorted(key_of(*[float(x) for x in v]) for v in in_file.values()), sorted(shown.keys()))

            print("\n── ⑦ ไฟล์ Excel ที่ตั้ง Solver ไว้ให้แล้ว ──")
            # ‼️ จุดตายของฟีเจอร์นี้คือชื่อที่นิยามระดับชีต ถ้าหลุดไปเป็นระดับสมุดงาน
            #    Excel จะเปิด Solver มาเป็นค่าว่าง ผู้ใช้กด Solve แล้วไม่มีอะไรเกิดขึ้น
            #    (พิสูจน์กับ Excel จริง 15/09/2026: ไฟล์ที่มีชื่อครบ กด Solve ได้คำตอบใน 0.82 วินาที)
            shown = search(target="400", hi="4")
            pg.locator("label.seg-item:visible", has_text="ทำเองใน Excel").first.click()
            pg.wait_for_timeout(500)
            with pg.expect_download() as dl2:
                pg.locator(".ms-solverbox button").click()
            spath = TMP / "solver.xlsx"
            dl2.value.save_as(str(spath))
            swb = openpyxl.load_workbook(spath)
            ck("ไฟล์มีสองชีต คือแผ่นงานกับแผ่นวิธีใช้", len(swb.sheetnames), 2)
            sws = swb[swb.sheetnames[0]]
            names = sorted(sws.defined_names.keys())
            need = ["solver_adj", "solver_eng", "solver_lhs1", "solver_num", "solver_opt",
                    "solver_rel1", "solver_rhs1", "solver_typ", "solver_val"]
            ck("การตั้งค่า Solver ถูกฝังเป็นชื่อระดับชีตครบ", [n for n in need if n not in names], [])
            ck("เป้าหมายในไฟล์ตรงกับที่ตั้งบนหน้าเว็บ", float(sws.defined_names["solver_val"].value), 400.0)
            ck("ชนิดโจทย์เป็น Value Of (3)", sws.defined_names["solver_typ"].value.strip(), "3")
            ck("ข้อจำกัดตัวแรกเป็นแบบ binary (5)", sws.defined_names["solver_rel1"].value.strip(), "5")
            nrow = sws.max_row
            ok("จำนวนแถวไม่เกินเพดาน 200 ตัวแปรของ Solver", nrow - 1 <= 200, f"ได้ {nrow - 1} แถว")
            ok("ช่วงตัวแปรครอบทุกแถวที่ใส่มาจริง",
               f"$B$2:$B${nrow}" in sws.defined_names["solver_adj"].value, 
               f"ได้ {sws.defined_names['solver_adj'].value}")
            # แถวที่เป็นคำตอบต้องถูกยกมาไว้ต้น ๆ ไม่งั้นกด Solve แล้วอาจไม่มีคำตอบอยู่ในไฟล์เลย
            head_rows = [sws.cell(row=r, column=3).value for r in range(2, min(12, nrow + 1))]
            answer_rows = {r for rows in shown.values() for r in rows}
            ok("แถวที่เป็นคำตอบถูกยกขึ้นมาไว้ต้นไฟล์",
               bool(answer_rows & set(head_rows)), f"คำตอบอยู่แถว {sorted(answer_rows)} แต่ต้นไฟล์คือ {head_rows}")
            # สูตรต้องเป็นสูตรจริง ไม่ใช่ข้อความ
            ok("ช่องรวมเป็นสูตร SUMPRODUCT จริง",
               str(sws["G1"].value).startswith("=SUMPRODUCT("), f"ได้ {sws['G1'].value!r}")

            print("\n── ⑧ โหมดวางตัวเลขเอง (ไม่มีไฟล์) ──")
            # ‼️ เดินทางเดียวกับผู้ใช้ที่ไม่มีไฟล์ในมือ: เปิดหน้า กดลิงก์บนหน้าเปล่า แล้วพิมพ์
            #    ของเดิมสั่ง click ผ่าน JS ทะลุหน้าเปล่าที่บังอยู่ จึงผ่านทั้งที่ผู้ใช้จริงเข้าโหมดนี้ไม่ได้เลย (22/09/2026)
            # เครื่องมือนี้จำเงื่อนไขล่าสุดไว้ (ขั้น ④ ตั้งไม่เกิน 1 รายการ เพดานหนึ่งล้าน แล้วถูกจำไว้)
            # ขั้นนี้ตรวจทางเข้าแบบไม่มีไฟล์ ไม่ได้ตรวจเรื่องจำค่า จึงเริ่มจากค่าเริ่มต้นเหมือนขั้น ⑨
            pg.evaluate("() => { try { localStorage.removeItem('filekit-state-excel-match-sum'); } catch {} }")
            fkui.open_tool(pg, "excel-match-sum", base)
            ck("เปิดมาไม่มีไฟล์ เริ่มที่หน้าเปล่า", fkui.state(pg), "landing")
            pg.locator(".s2-land button:visible", has_text="วางตัวเลขเอง").first.click()
            pg.wait_for_timeout(600)
            ck("กดวางตัวเลขเองบนหน้าเปล่าแล้วเข้าสถานะทำงาน", fkui.state(pg), "work")
            pg.keyboard.type("100\n250.50\n49.50\n300")
            pg.wait_for_timeout(600)
            ok("พิมพ์ลงช่องวางตัวเลขได้ทันทีโดยไม่ต้องกดหาช่องเอง",
               pg.evaluate("() => document.querySelector('.ms-ta').value").count("\n") == 3)
            ok("ปุ่มค้นหากดได้แม้ไม่มีไฟล์",
               pg.locator("button.s2-cta:visible", has_text="ค้นหา").first.is_enabled())
            pg.locator(".s2-side-bd input.ms-num:visible").first.fill("400")
            fkui.act(pg, "ค้นหา")
            pg.wait_for_timeout(800)
            ck("ค้นจากตัวเลขที่วางเองได้", set(sets_on_screen(pg)),
               {key_of(300, 100), key_of(250.5, 100, 49.5)})

            print("\n── ⑨ โหมดอังกฤษ ──")
            # ‼️ ต้องล้างเงื่อนไขที่จำไว้ก่อนรีโหลด (เพิ่ม 19/09/2026)
            #    ตั้งแต่เครื่องมือนี้จำค่าได้ ขั้นก่อนหน้าตั้งช่วงจำนวนรายการกับช่วงค่าไว้
            #    พอรีโหลดค่าพวกนั้นกลับมา (ซึ่งถูกต้องตามที่ตั้งใจ) แล้วขั้นนี้หาคำตอบไม่เจอ
            #    ขั้นนี้ตรวจเรื่องภาษาอังกฤษ ไม่ได้ตรวจเรื่องจำค่า จึงเริ่มจากค่าเริ่มต้นให้ชัด
            pg.evaluate("() => { try { localStorage.removeItem('filekit-state-excel-match-sum'); } catch {} }")
            pg.evaluate("() => { try { localStorage.setItem('fk-lang','en'); } catch {} }")
            pg.reload(wait_until="networkidle")
            pg.wait_for_timeout(400)
            pg.wait_for_selector(".s2", state="attached")
            pg.locator(".s2-land button:visible", has_text="paste numbers").first.click()
            pg.wait_for_timeout(500)
            pg.keyboard.type("100\n300")
            pg.locator(".s2-side-bd input.ms-num:visible").first.fill("400")
            pg.wait_for_timeout(500)
            fkui.act(pg, "Find")
            pg.wait_for_timeout(800)
            status = pg.evaluate("() => document.querySelector('.status')?.textContent || ''")
            ok("ชุดเดียวต้องเขียนว่า 1 set ไม่ใช่ 1 sets", "1 set " in status and "1 sets" not in status,
               f"ได้ {status!r}")
            # ‼️ เดิมอ่านจาก .ws-body ซึ่งไม่มีแล้วใน v2 ได้ข้อความว่างเสมอ ข้อนี้จึงผ่านแบบตาบอด
            #    ต้องอ่านจากกล่องเครื่องมือทั้งกล่อง และต้องเห็นข้อความจริงก่อน ถึงจะเชื่อว่า "ไม่มีไทย"
            body = pg.evaluate("() => document.querySelector('.s2')?.textContent || ''")
            ok("ตัวตรวจอ่านข้อความของเครื่องมือได้จริง (ไม่ใช่ว่างเปล่า)", "Find" in body and len(body) > 200,
               f"ได้ {len(body)} ตัวอักษร")
            thai = pg.evaluate("() => (document.querySelector('.s2')?.textContent || '')"
                               ".match(/[\\u0E00-\\u0E7F]+/g) || []")
            ck("ไม่มีภาษาไทยตกค้างในโหมด EN", thai, [])
            pg.evaluate("() => { try { localStorage.removeItem('fk-lang'); } catch {} }")

            ok("ไม่มี error ในคอนโซลตลอดการทดสอบ", not errs, f"เจอ {errs[:3]}")
            ctx.close()
            b.close()
    finally:
        if server:
            server.terminate()

    print(f"\nสรุป: ผ่าน {P} ข้อ, ตก {len(F)} ข้อ")
    for f in F:
        print("  ❌ " + f)
    sys.exit(1 if F else 0)


if __name__ == "__main__":
    main()
