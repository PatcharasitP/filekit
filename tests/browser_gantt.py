"""เทสกราฟ Gantt Deneb บนหน้าเว็บจริง (เครื่องมือ pbi-gantt): ตรวจ "สิ่งที่ตาคนใช้เห็น" ไม่ใช่แค่ว่ามีโค้ด

‼️ วัดจากกรอบของแท่งที่เบราว์เซอร์วาดจริง (getBoundingClientRect) แล้วเทียบกับวันที่ในข้อมูล
   ถ้ากราฟโกหกเมื่อไร (แท่งเลื่อนผิดวัน, งานวันเดียวหายเป็นแท่งกว้างศูนย์, ความคืบหน้าไม่ตรงสัดส่วน,
   เส้นวันนี้อยู่ผิดที่) ข้อใดข้อหนึ่งจะแดง
   ‼️ ค่าที่คาดทุกตัวคำนวณจากชุดข้อมูล JSON ตัวเดียวกับที่เว็บอ่าน ไม่ใช่ตัวเลขที่พิมพ์ไว้เอง

รัน: python3 -m http.server 8899 &  แล้ว ../.venv/bin/python tests/browser_gantt.py
"""
import datetime as dt
import json
import pathlib
import re
import sys
import zipfile

from playwright.sync_api import sync_playwright

BASE = __import__("os").environ.get("FK_BASE", "http://localhost:8899")
fails = []


def ck(label, got, want):
    ok = got == want
    print(("  ✅ " if ok else "  ❌ ") + label)
    if not ok:
        print(f"      ได้    : {got}\n      ควรได้ : {want}")
        fails.append(label)


def ck_true(label, cond, detail=""):
    print(("  ✅ " if cond else "  ❌ ") + label)
    if not cond:
        if detail:
            print("      " + detail)
        fails.append(label)


def d(iso):
    return dt.date.fromisoformat(iso)


# อ่านแท่งจาก DOM: ทุกกล่องของชั้นแท่งพื้นหลัง (0) กับชั้นความคืบหน้า (1) เรียงตามแถว
READ = """() => {
  const boxes = [];
  for (const g of document.querySelectorAll('#pbig-chart g[aria-roledescription="rect mark container"]')) {
    const ps = [...g.querySelectorAll('path[aria-label]')].filter(p => /^Task: /.test(p.getAttribute('aria-label')));
    if (!ps.length) continue;
    boxes.push(ps.map(p => {
      const r = p.getBoundingClientRect();
      const name = p.getAttribute('aria-label').match(/^Task: ([^;]+)/)[1];
      return { name, x0: r.left, x1: r.right, y: r.top, w: r.width, h: r.height, bottom: r.bottom, fill: p.getAttribute('fill'),
               opacity: Number(p.getAttribute('opacity') ?? 1) };
    }));
  }
  return boxes;
}"""

# เส้นวันนี้ = เส้นแนวตั้งของ rule mark ที่ยังมองเห็น (opacity > 0) คืนตำแหน่ง x กลางเส้น
READ_TODAY = """() => {
  const out = [];
  for (const g of document.querySelectorAll('#pbig-chart g[aria-roledescription="rule mark container"]'))
    for (const l of g.querySelectorAll('line, path')) {
      const r = l.getBoundingClientRect();
      const op = Number(l.getAttribute('opacity') ?? l.parentElement.getAttribute('opacity') ?? 1);
      if (r.height > 60 && r.width < 6 && op > 0) out.push({ x: r.left + r.width / 2, op });
    }
  return out;
}"""


def open_tool(pg):
    pg.goto("about:blank")
    pg.goto(f"{BASE}/#/pbi-gantt", wait_until="networkidle")
    pg.wait_for_function("() => document.querySelectorAll('#pbig-chart path[aria-label]').length > 0", timeout=20000)
    pg.wait_for_timeout(400)


def pick_dataset(pg, idx):
    pg.locator("input[name='pbig-ds']").nth(idx).check()
    pg.wait_for_timeout(600)


def check_geometry(pg, rows, label):
    """แท่งพื้นหลังต้องเริ่มและกว้างตามวันจริง (จบนับรวมวันนั้น) และเรียงตามวันเริ่ม"""
    layers = pg.evaluate(READ)
    ck_true(f"{label}: วาดครบสองชั้น ชั้นละ {len(rows)} แท่ง", len(layers) >= 2 and len(layers[0]) == len(rows) == len(layers[1]),
            f"ชั้น={[len(x) for x in layers]} ข้อมูล={len(rows)}")
    if not layers or len(layers[0]) != len(rows):
        return None
    base = sorted(layers[0], key=lambda b: b["y"])
    want_order = [r["Task"] for r in sorted(rows, key=lambda r: r["Start"])]   # เรียงเสถียร ตามวันเริ่ม
    ck(f"{label}: เรียงตามวันเริ่ม", [b["name"] for b in base], want_order)
    by = {r["Task"]: r for r in rows}
    # px ต่อวัน วัดจากแท่งที่ยาวที่สุด แล้วเทียบทุกแท่งกับสเกลนี้
    span = lambda r: (d(r["End"]) - d(r["Start"])).days + 1
    top = max(rows, key=span)
    topbox = next(b for b in base if b["name"] == top["Task"])
    ppd = topbox["w"] / span(top)
    origin = min(d(r["Start"]) for r in rows)
    first_x0 = min(b["x0"] for b in base)
    worst_w = worst_x = 0.0
    for b in base:
        r = by[b["name"]]
        worst_w = max(worst_w, abs(b["w"] - span(r) * ppd))
        worst_x = max(worst_x, abs((b["x0"] - first_x0) - (d(r["Start"]) - origin).days * ppd))
    ck_true(f"{label}: ความกว้างแท่งตรงจำนวนวัน (เพี้ยนสูงสุด {worst_w:.2f}px เพดาน 1px)", worst_w <= 1.0)
    ck_true(f"{label}: ตำแหน่งเริ่มแท่งตรงวันเริ่ม (เพี้ยนสูงสุด {worst_x:.2f}px เพดาน 1px)", worst_x <= 1.0)
    # ‼️ ความสูงแท่ง: รอบแรกเทสวัดแค่ความกว้างกับตำแหน่งซ้ายขวา จึงไม่เห็นว่าใส่ข้อมูล 3 งานแล้วแท่งสูง 0px
    #    (พื้นที่พล็อตเหลือศูนย์เพราะหักขอบซ้ำ เห็นจากภาพเว็บจริงเท่านั้น) และไม่เห็นว่าที่ว่างด้านบนกินไปเกินครึ่ง
    box = pg.evaluate("() => { const r = document.querySelector('.pbig-chart-box').getBoundingClientRect(); return { t: r.top, b: r.bottom }; }")
    thin = [b0["name"] for b0 in base if b0["h"] < 8]
    ck_true(f"{label}: ทุกแท่งสูงอย่างน้อย 8px (ต่ำสุด {min(b0['h'] for b0 in base):.1f}px)", not thin, str(thin))
    out_box = [b0["name"] for b0 in base if b0["y"] < box["t"] - 0.5 or b0["bottom"] > box["b"] + 0.5]
    ck_true(f"{label}: ทุกแท่งอยู่ในกรอบกราฟตามแนวตั้ง", not out_box, str(out_box))
    blank = min(b0["y"] for b0 in base) - box["t"]
    ck_true(f"{label}: ที่ว่างเหนือแท่งแรก {blank:.0f}px ไม่เกิน 60% ของกรอบสูง {box['b'] - box['t']:.0f}px",
            blank <= 0.6 * (box["b"] - box["t"]))
    return {"ppd": ppd, "first_x0": first_x0, "origin": origin, "base": base, "prog": layers[1]}


def main():
    ds = json.loads(pathlib.Path("samples/powerbi/gantt-datasets.json").read_text(encoding="utf-8"))
    keys = list(ds)
    with sync_playwright() as P:
        b = P.chromium.launch()
        ctx = b.new_context(viewport={"width": 1440, "height": 950}, accept_downloads=True)
        pg = ctx.new_page()
        errs = []
        pg.on("console", lambda m: errs.append(m.text) if m.type == "error" else None)
        pg.on("pageerror", lambda e: errs.append(str(e)))

        print("\n── ชุดที่ 1: แผนสร้างบ้าน ──")
        open_tool(pg)
        rows = ds[keys[0]]["rows"]
        g = check_geometry(pg, rows, "สร้างบ้าน")
        if g:
            # ความคืบหน้า: ชั้นบนต้องกว้างเป็นสัดส่วน Progress ของแท่งพื้นหลัง
            base = {b0["name"]: b0 for b0 in g["base"]}
            worst = 0.0
            n = 0
            for p in g["prog"]:
                want = next(r for r in rows if r["Task"] == p["name"])["Progress"]
                got = p["w"] / base[p["name"]]["w"]
                worst = max(worst, abs(got - want))
                n += 1
            ck_true(f"สร้างบ้าน: สัดส่วนความคืบหน้าตรงข้อมูลทั้ง {n} แท่ง (เพี้ยนสูงสุด {worst:.4f} เพดาน 0.01)", n > 0 and worst <= 0.01)

            # เส้นวันนี้: ชุดตัวอย่างใช้ demoToday ของชุดนั้น ต้องอยู่ตรงวันนั้นในสเกลเดียวกับแท่ง
            today = pg.evaluate(READ_TODAY)
            ck("สร้างบ้าน: มีเส้นวันนี้เส้นเดียว", len(today), 1)
            if today:
                want_x = g["first_x0"] + (d(ds[keys[0]]["demoToday"]) - g["origin"]).days * g["ppd"]
                ck_true(f"สร้างบ้าน: เส้นวันนี้อยู่ตรงวัน {ds[keys[0]]['demoToday']} (ห่าง {abs(today[0]['x'] - want_x):.2f}px เพดาน 2px)",
                        abs(today[0]["x"] - want_x) <= 2.0)
            note = pg.locator(".pbig-note").inner_text() if pg.locator(".pbig-note:not([hidden])").count() else ""
            ck_true("สร้างบ้าน: บอกผู้ใช้ว่าตัวอย่างนี้ตั้งวันนี้เป็นวันไหน", "ตัวอย่างนี้ตั้ง" in note and "22 มี.ค. 2569" in note, note)

        print("\n── ปิดแถบความคืบหน้า และปิดเส้นวันนี้ ──")
        pg.get_by_label(re.compile("แถบความคืบหน้า")).uncheck()
        pg.wait_for_timeout(500)
        prog = pg.evaluate(READ)[1]
        ck("ปิดแถบความคืบหน้า ชั้นบนมองไม่เห็นทุกแท่ง (opacity 0)", sorted({p["opacity"] for p in prog}), [0.0])
        base_ops = sorted({p["opacity"] for p in pg.evaluate(READ)[0]})
        ck("ปิดแล้วแท่งพื้นหลังทึบเต็ม (opacity 1)", base_ops, [1.0])
        pg.get_by_label(re.compile("แถบความคืบหน้า")).check()
        pg.get_by_label(re.compile("เส้นวันนี้")).uncheck()
        pg.wait_for_timeout(500)
        ck("ปิดเส้นวันนี้ ไม่เหลือเส้นที่มองเห็น", len(pg.evaluate(READ_TODAY)), 0)
        pg.get_by_label(re.compile("เส้นวันนี้")).check()

        print("\n── กำหนดวันนี้เองนอกช่วงแผน: ต้องไม่ขีดเส้น และต้องบอกผู้ใช้ ──")
        pg.get_by_label(re.compile("กำหนดวันนี้เอง")).fill("2030-01-01")
        pg.wait_for_timeout(600)
        ck("วันนี้ 2030 อยู่นอกแผน ไม่มีเส้นวันนี้", len(pg.evaluate(READ_TODAY)), 0)
        note = pg.locator(".pbig-note").inner_text() if pg.locator(".pbig-note:not([hidden])").count() else ""
        ck_true("มีข้อความบอกว่าวันนี้อยู่นอกช่วงของแผน", "อยู่นอกช่วงของแผนนี้" in note, note)
        pg.get_by_label(re.compile("กำหนดวันนี้เอง")).fill("ไม่ใช่วันที่")
        pg.wait_for_timeout(500)
        note = pg.locator(".pbig-note").inner_text() if pg.locator(".pbig-note:not([hidden])").count() else ""
        ck_true("พิมพ์วันที่มั่ว มีข้อความบอกรูปแบบที่ถูก", "ปปปป-ดด-วว" in note, note)
        pg.get_by_label(re.compile("กำหนดวันนี้เอง")).fill("")

        print("\n── ชุดที่ 3: งานวันเดียวต้องเป็นแท่งกว้าง 1 วันเต็ม ──")
        pick_dataset(pg, 2)
        rows3 = ds[keys[2]]["rows"]
        g3 = check_geometry(pg, rows3, "ซ่อมบำรุง")
        if g3:
            one = [b0 for b0 in g3["base"] if next(r for r in rows3 if r["Task"] == b0["name"])["Start"] ==
                   next(r for r in rows3 if r["Task"] == b0["name"])["End"]]
            ck_true("ในชุดนี้มีงานวันเดียวจริง (ประชากรไม่ว่าง)", len(one) >= 2, f"พบ {len(one)}")
            ck_true("งานวันเดียวกว้างเท่า 1 วัน ไม่ใช่ศูนย์", all(abs(b0["w"] - g3["ppd"]) <= 1.0 for b0 in one),
                    str([(b0["name"], round(b0["w"], 2), round(g3["ppd"], 2)) for b0 in one]))
            two = [b0 for b0 in g3["base"] if (d(next(r for r in rows3 if r["Task"] == b0["name"])["End"]) -
                                              d(next(r for r in rows3 if r["Task"] == b0["name"])["Start"])).days == 1]
            ck_true("งาน 2 วันกว้างเป็นสองเท่าของงานวันเดียว", bool(two) and abs(two[0]["w"] / one[0]["w"] - 2.0) <= 0.05,
                    str([(round(two[0]["w"], 2), round(one[0]["w"], 2))] if two else "ไม่มีงาน 2 วัน"))

        print("\n── ชุดที่ 2: ไม่มี Progress เลย สเปกต้องไม่พัง ──")
        pick_dataset(pg, 1)
        rows2 = ds[keys[1]]["rows"]
        g2 = check_geometry(pg, rows2, "ระบบไอที")
        if g2:
            ck("ไม่มี Progress ชั้นบนมองไม่เห็นทุกแท่ง", sorted({p["opacity"] for p in g2["prog"]}), [0.0])
            ck("ไม่มี Progress แท่งพื้นหลังทึบเต็ม", sorted({p["opacity"] for p in g2["base"]}), [1.0])

        print("\n── ส่งออก: สเปกที่คัดลอกต้องสะอาด ไม่ปนค่าของตัวอย่าง ──")
        pick_dataset(pg, 0)
        with pg.expect_download(timeout=30000) as dl:
            pg.locator("button", has_text="ดาวน์โหลด .json").first.click()
        spec = json.loads(pathlib.Path(dl.value.path()).read_text(encoding="utf-8"))
        ck("ชื่อไฟล์สเปก", dl.value.suggested_filename, "gantt-spec.json")
        pv = {p["name"]: p.get("value") for p in spec["params"]}
        ck("ในสเปกที่ส่งออก todayDate ยังว่าง (demoToday ใช้แค่พรีวิว)", pv.get("todayDate"), "")
        ck("ในสเปกที่ส่งออก ข้อมูลยังเป็น placeholder ของ Deneb", spec["data"], {"name": "dataset"})
        with pg.expect_download(timeout=30000) as dl:
            pg.locator("button", has_text="ดาวน์โหลดข้อมูลชุดนี้ .csv").first.click()
        raw = pathlib.Path(dl.value.path()).read_bytes()
        ck_true("ไฟล์ csv มี BOM (ไม่งั้น Excel อ่านไทยเพี้ยน)", raw.startswith(b"\xef\xbb\xbf"))
        lines = raw.decode("utf-8-sig").splitlines()
        ck("หัว csv", lines[0], "Task,Start,End,Group,Progress")
        ck("csv ครบทุกแถวข้อมูล", len(lines) - 1, len(rows))

        print("\n── ไฟล์ Excel ตัวอย่างโหลดได้ ค่าตรงชุดข้อมูล แล้วใส่กลับเข้าเครื่องมือได้ ──")
        with pg.expect_download(timeout=60000) as dl:
            pg.locator("button", has_text="ดาวน์โหลด .xlsx ตัวอย่าง").first.click()
        raw_x = pathlib.Path(dl.value.path())
        xlsx = raw_x.with_suffix(".xlsx")
        xlsx.write_bytes(raw_x.read_bytes())
        ck("ชื่อไฟล์ที่ผู้ใช้ได้", dl.value.suggested_filename, "pbi-gantt-samples.xlsx")
        ck_true("เป็นไฟล์ xlsx จริง", zipfile.is_zipfile(xlsx) and xlsx.stat().st_size > 3000)
        import openpyxl
        wb = openpyxl.load_workbook(xlsx)
        ck("ชีตครบทุกชุดข้อมูล บวกชีตข้อมูลมีจุดผิด", wb.sheetnames, [k[:31] for k in keys] + ["ลองข้อมูลมีจุดผิด"])
        wrong = []
        for k in keys:
            got = list(wb[k[:31]].iter_rows(min_row=2, values_only=True))
            if len(got) != len(ds[k]["rows"]):
                wrong.append((k, "จำนวนแถว", len(ds[k]["rows"]), len(got)))
            for a, row in zip(ds[k]["rows"], got):
                have = {"Task": row[0], "Start": row[1].date().isoformat(), "End": row[2].date().isoformat(),
                        "Group": row[3], "Progress": row[4]}
                for col in ("Task", "Start", "End", "Group", "Progress"):
                    if a.get(col) != have[col]:
                        wrong.append((k, col, a.get(col), have[col]))
        ck("ค่าทุกแถวในไฟล์ตรงกับชุดข้อมูลของหน้าเว็บ", wrong[:4], [])

        # ปิดวง: ใส่ชีตวันที่จริงของ Excel กลับเข้าไป ต้องได้กราฟเท่ากับชุดตัวอย่างเป๊ะ
        pg.locator(".pbig-own input[type=file]").set_input_files(str(xlsx))
        pg.wait_for_selector(".pbig-own-pick:not([hidden])", timeout=20000)
        sheet_sel = pg.locator(".pbig-own-pick select").first
        sheet_sel.select_option("house_build")
        pg.wait_for_timeout(500)
        picked = pg.locator(".pbig-own-pick select").evaluate_all("els => els.map(e => e.selectedOptions[0].textContent)")
        ck("เดาคอลัมน์จากหัวไทยได้ครบห้าช่อง", picked[1:], ["ชื่องาน", "วันเริ่ม", "วันจบ", "กลุ่ม", "ความคืบหน้า"])
        pg.locator(".pbig-own-pick button", has_text="ใช้ข้อมูลนี้").click()
        pg.wait_for_timeout(900)
        g_own = check_geometry(pg, rows, "ไฟล์ Excel วันที่จริง")
        if g_own:
            # ต้องเป็นกราฟหน้าตาเดียวกับชุดตัวอย่างแรก: สัดส่วนแท่งเท่ากันทุกแท่ง
            ck_true("ไฟล์ Excel วันที่จริง: ความคืบหน้าตรงข้อมูล",
                    all(abs(p["w"] / next(x for x in g_own["base"] if x["name"] == p["name"])["w"] -
                            next(r for r in rows if r["Task"] == p["name"])["Progress"]) <= 0.01 for p in g_own["prog"]))

        # ชีตข้อมูลมีจุดผิด: วันที่ พ.ศ. แบบข้อความ + แถวเสีย 3 แบบ ต้องข้ามและบอกเลขบรรทัด
        sheet_sel.select_option("ลองข้อมูลมีจุดผิด")
        pg.wait_for_timeout(500)
        pg.locator(".pbig-own-pick button", has_text="ใช้ข้อมูลนี้").click()
        pg.wait_for_timeout(900)
        layers = pg.evaluate(READ)
        names = sorted([b0["name"] for b0 in layers[0]], key=lambda n: n)
        ck("ข้อมูลมีจุดผิด: เหลือเฉพาะงานที่ใช้ได้ 3 งาน", names, sorted(["เตรียมพื้นที่", "เทฐานราก", "ก่ออิฐผนัง"]))
        # ‼️ ข้อมูลน้อย (3 งาน) กล่องกราฟเตี้ยสุด เคยได้แท่งสูง 0px ทั้งที่ DOM มีแท่งครบ ต้องวัดความสูงจริง
        ck_true("ข้อมูลมีจุดผิด (3 งาน กล่องเตี้ยสุด): ทุกแท่งสูงอย่างน้อย 8px", all(b0["h"] >= 8 for b0 in layers[0]),
                str([(b0["name"], round(b0["h"], 1)) for b0 in layers[0]]))
        note = pg.locator(".pbig-note").inner_text() if pg.locator(".pbig-note:not([hidden])").count() else ""
        ck_true("บอกว่าข้าม 3 แถว", "ข้าม 3 แถว" in note, note)
        for line, why in [(4, "วันจบมาก่อนวันเริ่ม"), (5, "ไม่มีชื่องาน"), (6, "วันเริ่มอ่านไม่ได้")]:
            ck_true(f"บอกบรรทัด {line}: {why}", f"บรรทัด {line} {why}" in note, note)
        # วันที่ พ.ศ. ต้องอ่านเป็น ค.ศ. ถูกปี: เริ่ม 01/07/2569 ห่างเทฐานราก 06/07/2569 ห้าวัน
        base = {b0["name"]: b0 for b0 in layers[0]}
        ppd = (base["เตรียมพื้นที่"]["w"]) / 5   # เตรียมพื้นที่ 01 ถึง 05 ก.ค. = 5 วัน
        dx = base["เทฐานราก"]["x0"] - base["เตรียมพื้นที่"]["x0"]
        ck_true(f"วันที่ พ.ศ. แบบข้อความอ่านถูก (เทฐานรากเริ่มหลังเตรียมพื้นที่ 5 วัน วัดได้ {dx / ppd:.2f} วัน)", abs(dx / ppd - 5) <= 0.05)

        print("\n── โหมดอังกฤษ ──")
        ctx2 = b.new_context(viewport={"width": 1440, "height": 950})
        ctx2.add_init_script("try{localStorage.setItem('fk-lang','en')}catch(e){}")
        pe = ctx2.new_page()
        pe.goto(f"{BASE}/#/pbi-gantt", wait_until="networkidle")
        pe.wait_for_function("() => document.querySelectorAll('#pbig-chart path[aria-label]').length > 0", timeout=20000)
        pe.wait_for_timeout(500)
        ck("หน้าเป็นภาษาอังกฤษจริง", pe.evaluate("document.documentElement.lang"), "en")
        en_names = [b0["name"] for b0 in pe.evaluate(READ)[0]]
        ck_true("ชื่องานเป็นอังกฤษล้วน ไม่มีตัวอักษรไทยหลงอยู่", en_names and not any(any("฀" <= c <= "๿" for c in n) for n in en_names),
                str(en_names))
        chart_txt = pe.evaluate("() => document.querySelector('#pbig-chart svg').textContent")
        ck_true("ตัวอักษรในกราฟ (ป้ายวันนี้ คำอธิบายสี) ไม่มีภาษาไทยในโหมดอังกฤษ",
                not any("฀" <= c <= "๿" for c in chart_txt), chart_txt[:200])
        aria = pe.evaluate("() => [...document.querySelectorAll('#pbig-chart path[aria-label]')].map(p => p.getAttribute('aria-label')).join(' | ')")
        ck_true("หัวข้อ tooltip ในโหมดอังกฤษไม่มีภาษาไทย", not any("฀" <= c <= "๿" for c in aria), aria[:200])
        ctx2.close()

        print("\n── จอแคบ: คำอธิบายสีต้องเห็นครบทุกกลุ่ม ไม่ถูกตัดที่ขอบ ──")
        # ‼️ ที่มา 30/09/2026: รอบแรก legend แถวเดียวถูกตัดที่ 390px เหลือ 3 จาก 5 กลุ่ม (เห็นจากภาพจริง)
        #    แก้รอบสองแล้วยังตัดที่ 320px เพราะสูตรอ่านความกว้างผิดตัว (width หักขอบแกนชื่องานไปแล้ว)
        groups = sorted({r["Group"] for r in ds[keys[0]]["rows"]})
        READ_LEGEND = """(names) => {
          const box = document.querySelector('.pbig-chart-box').getBoundingClientRect();
          const texts = [...document.querySelectorAll('#pbig-chart svg text')];
          const find = (t) => texts.find(e => e.textContent.trim() === t);
          const rect = (e) => { const r = e.getBoundingClientRect(); return { l: r.left, r: r.right, t: r.top, b: r.bottom }; };
          const today = find('วันนี้');
          const axis = texts.filter(e => /^\\d{2} [A-Z][a-z]{2}( \\d{4})?$/.test(e.textContent.trim()));
          return { box: { l: box.left, r: box.right }, found: names.map(n => { const e = find(n); return e ? { n, ...rect(e) } : { n, missing: true }; }),
                   plus: texts.some(e => /^\\+\\d+$/.test(e.textContent.trim())), today: today ? rect(today) : null,
                   axisTop: axis.length ? Math.min(...axis.map(e => e.getBoundingClientRect().top)) : null };
        }"""
        for width in (320, 390, 720):
            cw = b.new_context(viewport={"width": width, "height": 900}, is_mobile=(width < 500), has_touch=(width < 500))
            pw = cw.new_page()
            open_tool(pw)
            lg = pw.evaluate(READ_LEGEND, groups)
            ck_true(f"{width}px: พบชื่อกลุ่มในกราฟครบ {len(groups)} กลุ่ม (ประชากรไม่ว่าง)", len(groups) == 5 and not any(f.get("missing") for f in lg["found"]),
                    str([f["n"] for f in lg["found"] if f.get("missing")]))
            ck_true(f"{width}px: ไม่มี +N (ที่พอโชว์ครบไม่ควรตัด)", not lg["plus"])
            inside = [f for f in lg["found"] if not f.get("missing") and f["l"] >= lg["box"]["l"] - 0.5 and f["r"] <= lg["box"]["r"] + 0.5]
            ck_true(f"{width}px: ชื่อกลุ่มอยู่ในกรอบกราฟทุกกลุ่ม (ไม่ล้นขอบ)", len(inside) == len(groups),
                    str([(f["n"], f.get("missing") or (round(f["l"]), round(f["r"]), round(lg["box"]["r"]))) for f in lg["found"] if f not in inside]))
            if lg["today"] and lg["axisTop"] is not None:
                low = max(f["b"] for f in lg["found"] if not f.get("missing"))
                ck_true(f"{width}px: คำอธิบายสีไม่ทับป้ายวันนี้ และป้ายวันนี้ไม่ทับแกนวันที่",
                        low <= lg["today"]["t"] + 1 and lg["today"]["b"] <= lg["axisTop"] + 1,
                        f"legend ล่าง={low:.0f} วันนี้={lg['today']['t']:.0f}-{lg['today']['b']:.0f} แกนบน={lg['axisTop']:.0f}")
            cw.close()

        real_errs = [e for e in errs if "favicon" not in e.lower()]
        ck("ไม่มี error ใน console", real_errs, [])
        b.close()

    print("\n" + "━" * 54)
    print(f"ตก {len(fails)} ข้อ" if fails else "ผ่านทุกข้อ")
    sys.exit(1 if fails else 0)


main()
