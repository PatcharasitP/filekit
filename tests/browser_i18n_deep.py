# -*- coding: utf-8 -*-
"""
tests/browser_i18n_deep.py — จับ "ไทยตกค้าง/แปลผิด" ในโหมดอังกฤษของ FileKit
ในจังหวะที่ tests/browser_lang.py (สแกนแค่ตอนเปิดหน้าเปล่า) ยังไม่เคยเห็น
─────────────────────────────────────────────────────────────────────────────
ครอบคลุม (ตามโจทย์):
  ① ข้อความหลัง "ลงมือทำจริง" — อัปโหลดไฟล์ตัวอย่าง → กดปุ่ม → ผลลัพธ์/สรุป
  ② ข้อความตอนไฟล์เสีย — ไฟล์เสียที่สร้างเองในโฟลเดอร์ชั่วคราว (ไม่แตะ samples/)
  ③ กล่องโต้ตอบ — ตัวดูรูปเต็มจอ (lightbox), กล่องเตือนไฟล์ผิดชนิด, แถบออฟไลน์ทุกสถานะ
  ④ attribute อ่านจอ (aria-label/title/placeholder) — สแกนซ้ำ "หลัง" การกระทำ ไม่ใช่แค่ตอนโหลดหน้า
  ⑤ ข้อความที่ JS สร้างตอนรันไทม์ — "N ไฟล์, รวม X MB", document.title, สรุปผลต่าง ๆ
  ⑥ สลับภาษาหลังมีผลลัพธ์ค้างอยู่ — ต้องไม่ทิ้งไทยค้างจอ (ดีไซน์จริง: สลับ = reload)
  ⑦ ตรวจย้อนกลับ — โหมดไทยหลังลงมือทำจริง มีอังกฤษทั้งประโยคหลุดมาไหม (ตัวแทน 10 เครื่องมือ
     ข้ามทุกตระกูลโค้ด — ที่เหลือคัดกรองด้วย static grep ที่พบว่า i18n ครบ 100% ในไฟล์เหล่านั้น)

‼️ ไม่แก้ src/**, index.html, เทสอื่น — สร้าง/แก้ได้เฉพาะไฟล์นี้ (+ไฟล์ชั่วคราวนอกโปรเจกต์)
‼️ กติกา ALLOW: คำไทยที่ "ควรเป็นไทยแม้โหมดอังกฤษ" ต้องอนุญาตเป็นรายคำ อธิบายได้ว่าทำไม
    เนื้อหาที่มาจากไฟล์ตัวอย่างที่เราอัปโหลดเอง (ชื่อไฟล์ Thai, ข้อมูลในตาราง Excel/Word)
    ก็เป็นแบบเดียวกัน — จึงอนุญาตเป็น "extra_allow" เฉพาะรอบทดสอบเครื่องมือนั้น ไม่ใช่ยกเว้นทั้งหน้า
"""
import os, re, sys, json, subprocess, pathlib, tempfile, shutil, zipfile

from playwright.sync_api import sync_playwright, TimeoutError as PWTimeout

ROOT = pathlib.Path(__file__).resolve().parent.parent
SAMPLES = ROOT / "samples"
BASE = os.environ.get("FK_BASE", "http://localhost:8923")
SIF_TIMEOUT = 10_000

WORKDIR = pathlib.Path(tempfile.mkdtemp(prefix="filekit_i18n_deep_"))
BROKEN_DIR = WORKDIR / "broken"
BROKEN_DIR.mkdir(parents=True, exist_ok=True)

THAI = re.compile(r"[฀-๿]")
# ‼️ ตัวจับ "อังกฤษทั้งประโยคหลุดมา" ตอนโหมดไทย — ต้องยาว >=3 คำอังกฤษติดกัน กันชนคำย่อ/ชื่อเทคนิคสั้น ๆ
EN_RUN = re.compile(r"[A-Za-z][A-Za-z.'-]*(?:[  ][A-Za-z][A-Za-z.'-]*){2,}")


def log(msg=""):
    print(msg, flush=True)


P, F = 0, []


def ck(name, ok, detail=""):
    global P
    if ok:
        P += 1
    else:
        F.append(f"{name}{detail}")
    log(f"  {'✅' if ok else '❌'} {name}")


# ── ทะเบียนเครื่องมือจากของจริง (src/registry.js) — ห้ามฮาร์ดโค้ดรายชื่อ ──────────
def _tool_ids():
    out = subprocess.run(
        ["node", "--input-type=module", "-e",
         'import {TOOLS} from "./src/registry.js"; console.log(JSON.stringify(TOOLS.map(t=>t.id)))'],
        cwd=str(ROOT), capture_output=True, text=True, check=True).stdout
    return json.loads(out)


ALL_TOOL_IDS = set(_tool_ids())

# ─────────────────────────────────────────────────────────────────────────
# ALLOW — คำไทยที่ "ควรเป็นไทยแม้โหมดอังกฤษ" ทั่วทั้งเว็บ (เหมือน tests/browser_lang.py
# คัดลอกมาเพราะ browser_lang.py รัน main() ทันทีตอน import — import ไม่ได้)
# ─────────────────────────────────────────────────────────────────────────
ALLOW = {
    "หนึ่งแสนสองหมื่นแปดพันสี่ร้อยบาทถ้วน", "บาทถ้วน", "สตางค์", "๑๒๓", "๐๑๒๓๔๕๖๗๘๙",
    "ม.ค.", "๑๕/๐๑/๒๕๖๙", "15 ม.ค. 2569", "พ.ศ.", "ค.ศ.",
    "เธชเธงเธฑ", "à¸ªà¸§",
    "ไทย",
    "ใช่", "ไม่ใช่", "มี", "ไม่มี",
    "ไม่",
    "ต./ตำบล/แขวง", "อ./อำเภอ/เขต", "จ./จังหวัด", "กรุงเทพฯ/กทม.", "ต./อ./จ.",
    "นางสาวสมหญิง ใจดี", "ณ อยุธยา",
    "฿",
}

# ‼️ container ที่เนื้อหาข้างในเป็น "ข้อมูลจากไฟล์ที่ผู้ใช้ป้อน" ไม่ใช่ข้อความของเว็บ —
#    ตรวจยืนยันจากโค้ดจริงแล้วว่าไม่ใช่จุดที่ต้องแปล (ดูเหตุผลในรายงานท้ายไฟล์เทส):
#    - table.xt td.old/td.new  (sheetpick.js บรรทัด 195/200/204) = ค่าเดิม/ค่าแปลงจากสเปรดชีตผู้ใช้
#    - .preview-text           (thai-encoding.js บรรทัด 76-77)   = เนื้อไฟล์ CSV จริงของผู้ใช้
#    - .mm-key / .mm-val / .stack (word-mailmerge.js)            = ชื่อตัวยึด/ค่าตัวอย่างจากเทมเพลต+Excel ผู้ใช้
DATA_CONTAINER_EXCLUDE = ["td.old", "td.new", ".preview-text", ".mm-key", ".mm-val", ".stack"]

# ดึงข้อความที่ "ตาเห็นจริง" เหมือน browser_lang.py แต่:
#  - รับ extra selector มายกเว้นเพิ่มได้ (เนื้อหาข้อมูลผู้ใช้ที่โผล่หลังลงมือทำ)
#  - ข้าม <option> ทั้งหมด (ตัวเลือกในดรอปดาวน์คอลัมน์/ชีทที่มาจากไฟล์ผู้ใช้ล้วน ๆ หลังอัปโหลด
#    ส่วนข้อความตัวเลือกคงที่ เช่น "— ไม่ใช้ —" ตรวจยืนยันด้วยโค้ดแล้วว่าแปลครบทุกจุด — ดูรายงานท้ายไฟล์)
VISIBLE_TEXT_JS = """(extra) => {
  const out = [];
  const skip = (node) => extra.some(sel => { try { return node.closest(sel); } catch(e) { return false; } });
  const walk = document.createTreeWalker(document.body, NodeFilter.SHOW_TEXT);
  for (let n = walk.nextNode(); n; n = walk.nextNode()) {
    const t = (n.nodeValue || "").trim();
    if (!t) continue;
    const p = n.parentElement;
    if (!p || /^(SCRIPT|STYLE|NOSCRIPT|OPTION)$/.test(p.tagName)) continue;
    if (p.closest("[hidden]")) continue;
    if (p.closest("#lang")) continue;
    if (skip(p)) continue;
    if (p.checkVisibility && !p.checkVisibility()) continue;
    out.push(t);
  }
  for (const el of document.querySelectorAll("[placeholder],[aria-label],[title]")) {
    if (el.checkVisibility && !el.checkVisibility()) continue;
    if (el.closest("#lang")) continue;
    if (skip(el)) continue;
    for (const a of ["placeholder","aria-label","title"]) {
      const v = el.getAttribute(a); if (v) out.push(v);
    }
  }
  return out;
}"""


def capture_texts(pg, extra_exclude=None):
    return pg.evaluate(VISIBLE_TEXT_JS, extra_exclude or DATA_CONTAINER_EXCLUDE)


def thai_leftovers(texts, extra_allow=()):
    allow = sorted(set(ALLOW) | set(extra_allow), key=len, reverse=True)
    bad = []
    for t in texts:
        s = t
        for a in allow:
            if a:
                s = s.replace(a, "")
        if THAI.search(s):
            bad.append(t)
    return bad


def english_leak(texts, extra_allow=()):
    """ย้อนกลับ: โหมดไทย มีประโยคอังกฤษทั้งท่อนหลุดมาไหม (>=3 คำอังกฤษติดกัน)
    allow เฉพาะวลีสั้น ๆ ที่รู้จักกันว่าใช้ปนในไทยได้ตามธรรมชาติ (ชื่อแบรนด์/มาตรฐานไฟล์)"""
    allow_en = {"FileKit", "UTF-8", "TIS-620", "Windows-874", "GitHub"} | set(extra_allow)
    bad = []
    for t in texts:
        for m in EN_RUN.finditer(t):
            frag = m.group(0)
            if any(a.lower() in frag.lower() for a in allow_en):
                continue
            bad.append((t, frag))
    return bad


# ─────────────────────────────────────────────────────────────────────────
# ไฟล์ตัวอย่างชั่วคราว: ที่อยู่ไทย (thai-address ไม่มีไฟล์ตัวอย่างที่มีคอลัมน์ที่อยู่ใน samples/)
# ─────────────────────────────────────────────────────────────────────────
def make_address_xlsx():
    import openpyxl
    p = WORKDIR / "ที่อยู่ตัวอย่าง-ทดสอบ.xlsx"
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.append(["ชื่อ", "ที่อยู่"])
    ws.append(["สมชาย ใจดี", "123 หมู่ 4 ต.บางพลี อ.บางพลี จ.สมุทรปราการ 10540"])
    ws.append(["สุดารัตน์ รักงาน", "45/6 ถ.สุขุมวิท แขวงคลองตันเหนือ เขตวัฒนา กรุงเทพมหานคร 10110"])
    ws.append(["ปัณณ์ วงศ์สุวรรณ", "เลขที่ 9 หมู่บ้านสวนทอง ตำบลหนองปรือ อำเภอบางละมุง จังหวัดชลบุรี 20150"])
    wb.save(p)
    return p


# ─────────────────────────────────────────────────────────────────────────
# ไฟล์เสีย 4 ชนิด (นอกโปรเจกต์) — ใช้เช็คข้อความ error ในโหมดอังกฤษ
# ─────────────────────────────────────────────────────────────────────────
def make_broken_files():
    files = {}
    p = BROKEN_DIR / "broken-fake.pdf"
    p.write_bytes(b"not a real pdf file, just garbage bytes 1234567890 !!!")
    files["garbage_pdf"] = p

    p = BROKEN_DIR / "empty.xlsx"
    p.write_bytes(b"")
    files["empty_xlsx"] = p

    p = BROKEN_DIR / "broken-header.jpg"
    p.write_bytes(bytes([0xFF, 0xD8, 0xFF]) + os.urandom(500))
    files["broken_jpg"] = p

    p = BROKEN_DIR / "fake-empty-zip.docx"
    with zipfile.ZipFile(p, "w") as z:
        z.writestr("hello.txt", "not a real Word file — just an empty zip renamed")
    files["fake_docx"] = p
    return files


# ─────────────────────────────────────────────────────────────────────────
# นำทางไปเครื่องมือ (SPA hash-router — ผ่าน about:blank ก่อนเสมอกันไม่ mount ใหม่)
# ─────────────────────────────────────────────────────────────────────────
def goto(pg, tool_id):
    assert tool_id in ALL_TOOL_IDS, f"ไม่พบเครื่องมือ id={tool_id} ใน src/registry.js"
    pg.goto("about:blank")
    pg.goto(f"{BASE}/#/{tool_id}", wait_until="networkidle")


def goto_home(pg):
    pg.goto("about:blank")
    pg.goto(f"{BASE}/#/", wait_until="networkidle")


def stems(paths):
    return {pathlib.Path(p).stem for p in paths}


def xlsx_headers(path):
    """หัวคอลัมน์จริงของไฟล์ — sheetpick.js โชว์ 'ชื่อคอลัมน์เดิม (original)' ใน <th> เสมอ
    (บรรทัด 186 ของ src/sheetpick.js) เป็น "ข้อมูลจากไฟล์ผู้ใช้" ไม่ใช่ข้อความเว็บ — allow เฉพาะรอบนี้"""
    import openpyxl
    wb = openpyxl.load_workbook(path, data_only=True)
    ws = wb.active
    row = next(ws.iter_rows(min_row=1, max_row=1, values_only=True), ())
    return {str(c).strip() for c in row if c and str(c).strip()}


# ─────────────────────────────────────────────────────────────────────────
# ① เครื่องมือ "อัปโหลด → กดปุ่ม → รอผล" — 19 ตัว
#    (ปุ่ม EN/TH ยืนยันตรงจาก src/tools/*.js: button(tr("ไทย","EN")) — ไม่เดา ใช้ได้ทั้ง 2 ภาษา
#    เพราะ do_tool_action ถูกเรียกทั้งใน context อังกฤษ (forward) และไทย (reverse §⑦))
# ─────────────────────────────────────────────────────────────────────────
SIMPLE = {
    "pdf-pages":    (["ตัวอย่าง-รายงานประจำเดือน.pdf"], "Save", "บันทึก"),
    "pdf-merge":    (["ตัวอย่าง-รายงานประจำเดือน.pdf", "ตัวอย่าง-ใบปะหน้าเอกสาร.pdf"], "Merge", "รวมไฟล์"),
    "pdf-split":    (["ตัวอย่าง-รายงานประจำเดือน.pdf"], "Split file", "แยกไฟล์"),
    "pdf-compress": (["ตัวอย่าง-รายงานประจำเดือน.pdf"], "Compress file", "บีบอัดไฟล์"),
    "pdf-watermark":(["ตัวอย่าง-รายงานประจำเดือน.pdf"], "Add watermark", "ใส่ลายน้ำ"),
    "pdf-to-images":(["ตัวอย่าง-รายงานประจำเดือน.pdf"], "Convert to images", "แปลงเป็นรูป"),
    "pdf-to-text":  (["ตัวอย่าง-รายงานประจำเดือน.pdf"], "Extract text", "ดึงข้อความ"),
    "pdf-to-word":  (["ตัวอย่าง-รายงานประจำเดือน.pdf"], "Convert to Word", "แปลงเป็น Word"),
    "pdf-to-excel": (["ตัวอย่าง-รายงานประจำเดือน.pdf"], "Convert to Excel", "แปลงเป็น Excel"),
    "word-to-pdf":  (["ตัวอย่าง-ใบเสนอราคา.docx"], "Convert to PDF", "แปลงเป็น PDF"),
    "excel-to-pdf": (["ตัวอย่าง-ข้อมูลใบเสนอราคา.xlsx"], "Convert to PDF", "แปลงเป็น PDF"),
    "images-to-pdf":(["ตัวอย่าง-รูปภาพ-1.jpg", "ตัวอย่าง-รูปภาพ-2.png"], "Create PDF", "สร้างไฟล์ PDF"),
    "image-convert":(["ตัวอย่าง-รูปภาพ-2.png"], "Convert files", "แปลงไฟล์"),
    "image-resize": (["ตัวอย่าง-รูปภาพ-1.jpg"], "Resize and compress", "ย่อและบีบอัด"),
    "word-join":    (["ตัวอย่าง-ใบเสนอราคา.docx", "ตัวอย่าง-หนังสือแจ้งผลประเมิน.docx"], "Merge files", "รวมไฟล์"),
    "word-clean":   (["ตัวอย่าง-ใบเสนอราคา.docx"], "Clean", "ล้าง"),
    "powerpoint-to-word": (["ตัวอย่าง-นำเสนอบริษัท.pptx"], "Convert to Word", "แปลงเป็น Word"),
    "powerpoint-to-pdf":  (["ตัวอย่าง-นำเสนอบริษัท.pptx"], "Create PDF", "สร้างไฟล์ PDF"),
    "excel-csv":    (["ตัวอย่าง-ข้อมูลใบเสนอราคา.xlsx"], "Convert file", "แปลงไฟล์"),
}

# ② เครื่องมือแบบ sheetpick — อัปโหลดแล้วประมวลผลอัตโนมัติ (ไม่มีปุ่ม "แปลง") แล้วกดดาวน์โหลด
SHEETPICK = {
    "thai-date":   ["ตัวอย่าง-วันที่ไทยและเลขบัตร.xlsx"],
    "thai-id":     ["ตัวอย่าง-วันที่ไทยและเลขบัตร.xlsx"],
    "thai-name":   ["ตัวอย่าง-ข้อมูลใบเสนอราคา.xlsx"],
    "thai-number": ["ตัวอย่าง-ข้อมูลใบเสนอราคา.xlsx"],
    # thai-address ใช้ไฟล์ชั่วคราว (สร้างตอนรัน) — ใส่ path เต็มตอนเรียกจริง
}
SHEETPICK_DL_TEXT = {"en": "Download as Excel", "th": "ดาวน์โหลดเป็น Excel"}


def run_simple(pg, tool_id, files_rel):
    file_paths = [str(SAMPLES / f) for f in files_rel]
    goto(pg, tool_id)
    pg.set_input_files("input[type=file]", file_paths, timeout=SIF_TIMEOUT)
    pg.wait_for_selector(".file-row", timeout=15_000)
    return file_paths


def click_go(pg, text, timeout=25_000):
    pg.get_by_role("button", name=text, exact=True).click()
    pg.wait_for_selector(".status-wrap .status.show", timeout=timeout)
    pg.wait_for_timeout(450)


def run_sheetpick(pg, tool_id, file_path, download_text=None, extra_wait=0):
    goto(pg, tool_id)
    pg.set_input_files("input[type=file]", str(file_path), timeout=SIF_TIMEOUT)
    pg.wait_for_selector(".stats .stat", timeout=20_000)
    if extra_wait:
        pg.wait_for_timeout(extra_wait)
    if download_text:
        with pg.expect_download():
            pg.get_by_role("button", name=download_text, exact=True).click()
        pg.wait_for_timeout(300)


# ── word-mailmerge: 2 dropzone (เทมเพลต + ข้อมูล) ───────────────────────────
def run_mailmerge(pg, click_text):
    goto(pg, "word-mailmerge")
    inputs = pg.locator("input[type=file]")
    inputs.nth(0).set_input_files(str(SAMPLES / "ตัวอย่าง-หนังสือแจ้งผลประเมิน.docx"), timeout=SIF_TIMEOUT)
    pg.wait_for_timeout(700)
    inputs.nth(1).set_input_files(str(SAMPLES / "ตัวอย่าง-ข้อมูลพนักงาน.xlsx"), timeout=SIF_TIMEOUT)
    pg.wait_for_timeout(1200)
    pg.get_by_role("button", name=click_text, exact=True).click()
    pg.wait_for_selector(".status-wrap .status.show", timeout=25_000)
    pg.wait_for_timeout(500)


# ── word-replace: ต้องกรอกคู่ค้นหา–แทนที่ก่อนกด ─────────────────────────────
def run_word_replace(pg, click_text):
    goto(pg, "word-replace")
    pg.set_input_files("input[type=file]", str(SAMPLES / "ตัวอย่าง-ใบเสนอราคา.docx"), timeout=SIF_TIMEOUT)
    pg.wait_for_selector(".file-row", timeout=15_000)
    row = pg.locator(".rep-row").first
    row.locator("input").nth(0).fill("บริษัท")
    row.locator("input").nth(1).fill("Company Ltd.")
    pg.get_by_role("button", name=click_text, exact=True).click()
    pg.wait_for_selector(".status-wrap .status.show", timeout=20_000)
    pg.wait_for_timeout(400)


# ── pdf-sign: อัปโหลด PDF + รูปลายเซ็น → คลิกวางบนหน้า → บันทึก ─────────────
def run_pdf_sign(pg, click_text):
    goto(pg, "pdf-sign")
    inputs = pg.locator("input[type=file]")
    inputs.nth(0).set_input_files(str(SAMPLES / "ตัวอย่าง-รายงานประจำเดือน.pdf"), timeout=SIF_TIMEOUT)
    pg.wait_for_selector(".sign-page", timeout=15_000)
    pg.wait_for_timeout(600)
    # อัปโหลดรูปเป็นลายเซ็น (input ที่ 2 = upInput ซ่อนอยู่ accept image/*)
    inputs.nth(1).set_input_files(str(SAMPLES / "ตัวอย่าง-รูปภาพ-1.jpg"), timeout=SIF_TIMEOUT)
    pg.wait_for_timeout(500)
    # คลิกกลางหน้าเพื่อวางลายเซ็น
    pg.locator(".sign-stage").click(position={"x": 200, "y": 200})
    pg.wait_for_timeout(300)
    pg.get_by_role("button", name=click_text, exact=True).click()
    pg.wait_for_selector(".status-wrap .status.show", timeout=20_000)
    pg.wait_for_timeout(400)


# ── thai-encoding: อัปโหลด → การ์ดขึ้นอัตโนมัติ → กด "Save as UTF-8"/"บันทึกเป็น UTF-8" ──
def run_thai_encoding(pg, click_text):
    goto(pg, "thai-encoding")
    pg.set_input_files("input[type=file]", str(SAMPLES / "ตัวอย่าง-ไทยเพี้ยน-แบบ TIS620.csv"), timeout=SIF_TIMEOUT)
    pg.wait_for_selector(".enc-card", timeout=15_000)
    with pg.expect_download():
        pg.get_by_role("button", name=click_text, exact=True).click()
    pg.wait_for_timeout(400)


# ‼️ ปุ่ม/ข้อความยึด 2 ภาษา สำหรับ custom-runner (ต้องคู่กับ src/tools/*.js บรรทัดจริง —
#    ตรวจแล้วจากโค้ด ไม่ใช่เดา) — dict คีย์ "en"/"th"
CUSTOM_BTN = {
    "word-mailmerge": {"en": "Generate all documents", "th": "สร้างเอกสารทั้งชุด"},
    "word-replace":   {"en": "Replace", "th": "แทนที่"},
    "pdf-sign":       {"en": "Save signed PDF", "th": "บันทึกไฟล์เซ็นแล้ว"},
    "thai-encoding":  {"en": "Save as UTF-8", "th": "บันทึกเป็น UTF-8"},
}


# ─────────────────────────────────────────────────────────────────────────
# runner กลาง: ทำ action ของเครื่องมือหนึ่งตัว ต้องรู้ "lang" ปัจจุบันของ context
# เพื่อกดปุ่มด้วยข้อความภาษาที่ตรงกับที่หน้าจอกำลังโชว์อยู่จริง (en หรือ th)
# คืน (texts, extra_allow_thai_stems)
# ─────────────────────────────────────────────────────────────────────────
def do_tool_action(pg, tool_id, address_xlsx_path, lang="en"):
    extra_allow = set()
    if tool_id in SIMPLE:
        files_rel, btn_en, btn_th = SIMPLE[tool_id]
        btn = btn_en if lang == "en" else btn_th
        paths = run_simple(pg, tool_id, files_rel)
        extra_allow |= stems(paths)
        if tool_id == "pdf-split":
            # ตั้งโหมด "ทุก N หน้า" ให้ได้ผลลัพธ์หลายไฟล์ + ปุ่ม ZIP (จังหวะที่เทสเดิมไม่เคยเห็น)
            pg.locator('input[type=radio][value="every"]').check()
            pg.locator("input[type=number]").fill("1")
            pg.wait_for_timeout(300)
        if tool_id == "image-convert":
            pg.locator("select").select_option("webp")
        click_go(pg, btn)
    elif tool_id in SHEETPICK:
        file_rel = SHEETPICK[tool_id][0]
        path = SAMPLES / file_rel
        extra_allow |= stems([path]) | xlsx_headers(path)
        extra_wait = 800 if tool_id == "thai-address" else 0
        run_sheetpick(pg, tool_id, path, download_text=SHEETPICK_DL_TEXT[lang], extra_wait=extra_wait)
    elif tool_id == "thai-address":
        extra_allow |= stems([address_xlsx_path]) | xlsx_headers(address_xlsx_path)
        run_sheetpick(pg, "thai-address", address_xlsx_path, download_text=SHEETPICK_DL_TEXT[lang], extra_wait=900)
    elif tool_id == "word-mailmerge":
        extra_allow |= stems([SAMPLES / "ตัวอย่าง-หนังสือแจ้งผลประเมิน.docx", SAMPLES / "ตัวอย่าง-ข้อมูลพนักงาน.xlsx"])
        run_mailmerge(pg, CUSTOM_BTN["word-mailmerge"][lang])
    elif tool_id == "word-replace":
        extra_allow |= stems([SAMPLES / "ตัวอย่าง-ใบเสนอราคา.docx"])
        run_word_replace(pg, CUSTOM_BTN["word-replace"][lang])
    elif tool_id == "pdf-sign":
        extra_allow |= stems([SAMPLES / "ตัวอย่าง-รายงานประจำเดือน.pdf", SAMPLES / "ตัวอย่าง-รูปภาพ-1.jpg"])
        run_pdf_sign(pg, CUSTOM_BTN["pdf-sign"][lang])
    elif tool_id == "thai-encoding":
        extra_allow |= stems([SAMPLES / "ตัวอย่าง-ไทยเพี้ยน-แบบ TIS620.csv"])
        run_thai_encoding(pg, CUSTOM_BTN["thai-encoding"][lang])
    else:
        raise AssertionError(f"ไม่มี config การทดสอบสำหรับ {tool_id}")

    texts = capture_texts(pg)
    return texts, extra_allow


# เครื่องมือทั้งหมดที่มี action-runner จริง (สแกนได้) — pdf-ocr ไม่รวม (ดูเหตุผลท้ายไฟล์)
ACTIONABLE = set(SIMPLE) | set(SHEETPICK) | {"thai-address", "word-mailmerge", "word-replace", "pdf-sign", "thai-encoding"}
SKIPPED_TOOLS = ALL_TOOL_IDS - ACTIONABLE

# ตัวแทนสำหรับ "ตรวจย้อนกลับ" (โหมดไทย หลังลงมือทำจริง) — ครบทุกตระกูลโค้ดหลัก
REVERSE_SAMPLE = [
    "pdf-merge", "pdf-to-text", "pdf-sign", "word-clean", "word-mailmerge",
    "word-replace", "image-resize", "excel-csv", "thai-date", "thai-address",
]


# ─────────────────────────────────────────────────────────────────────────
# ③ กล่องโต้ตอบ: lightbox / เตือนไฟล์ผิดชนิด / แถบออฟไลน์ (ทำได้ทั้ง 2 ภาษา)
# ─────────────────────────────────────────────────────────────────────────
def check_lightbox_image(pg, lang):
    goto(pg, "images-to-pdf")
    src = SAMPLES / "ตัวอย่าง-รูปภาพ-1.jpg"
    pg.set_input_files("input[type=file]", str(src), timeout=SIF_TIMEOUT)
    pg.wait_for_selector(".file-row button.thumb", timeout=15_000)
    pg.locator(".file-row button.thumb").first.click()
    pg.wait_for_selector("dialog.pv[open]", timeout=10_000)
    pg.wait_for_timeout(500)
    texts = capture_texts(pg, extra_exclude=[])  # lightbox caption มีชื่อไฟล์ (Thai stem) ปนอยู่ตั้งใจ
    pg.keyboard.press("Escape")
    pg.wait_for_timeout(200)
    return texts, stems([src])


def check_lightbox_pdf(pg, lang):
    # ‼️ ต้องใช้เครื่องมือที่มี "pdfjs" ใน libs (registry.js) — canView() ใน preview.js
    #    เปิดดู PDF ได้ก็ต่อเมื่อ window.pdfjsLib โหลดแล้วเท่านั้น (pdf-merge ใช้แค่ pdflib ไม่มี pdfjs
    #    เลยกดรูปย่อ PDF แล้วไม่มีอะไรเกิดขึ้นเลย — จับได้จริงตอนรันสคริปต์นี้ ไม่ใช่เดา)
    goto(pg, "pdf-to-images")
    src = SAMPLES / "ตัวอย่าง-รายงานประจำเดือน.pdf"
    pg.set_input_files("input[type=file]", str(src), timeout=SIF_TIMEOUT)
    pg.wait_for_selector(".file-row button.thumb", timeout=15_000)
    pg.locator(".file-row button.thumb").first.click()
    pg.wait_for_selector("dialog.pv[open]", timeout=10_000)
    pg.wait_for_timeout(900)  # รอ pdf.js วาดหน้าแรก + caption "page 1 of N"
    texts = capture_texts(pg, extra_exclude=[])
    pg.keyboard.press("Escape")
    pg.wait_for_timeout(200)
    return texts, stems([src])


def check_wrong_type(pg, lang):
    """อัปโหลด PPTX (ผิดชนิด) เข้าเครื่องมือที่รับเฉพาะ PDF — เช็คกล่องเตือน + ปุ่มพาไปเครื่องมือที่ถูก"""
    goto(pg, "pdf-merge")
    src = SAMPLES / "ตัวอย่าง-นำเสนอบริษัท.pptx"
    pg.set_input_files("input[type=file]", str(src), timeout=SIF_TIMEOUT)
    pg.wait_for_selector(".wrong-type", timeout=10_000)
    pg.wait_for_timeout(300)
    texts = capture_texts(pg, extra_exclude=[])
    return texts, stems([src])


def check_offline_bar(pg, lang):
    goto_home(pg)
    pg.wait_for_selector(".ob-btn", timeout=15_000)
    pg.click(".ob-btn")
    pg.wait_for_timeout(600)
    mid_texts = capture_texts(pg, extra_exclude=[])  # จับสถานะ "กำลังเตรียมไฟล์..." กลางทาง
    pg.wait_for_selector(".ob-ready, .ob-fail", timeout=90_000)
    pg.wait_for_timeout(300)
    end_texts = capture_texts(pg, extra_exclude=[])
    return mid_texts + end_texts, set()


# ─────────────────────────────────────────────────────────────────────────
# ② ไฟล์เสีย — โหมดอังกฤษ ต้องขึ้น error อ่านรู้เรื่องเป็นอังกฤษล้วน
# ─────────────────────────────────────────────────────────────────────────
def check_broken(pg, tool_id, path, wait_selector, click_en=None, wait_has_text=None):
    goto(pg, tool_id)
    pg.set_input_files("input[type=file]", str(path), timeout=SIF_TIMEOUT)
    pg.wait_for_selector(".file-row", timeout=15_000)
    if click_en:
        pg.get_by_role("button", name=click_en, exact=True).click()
    if wait_has_text:
        # ‼️ ".ws-empty" ใช้ซ้ำ 2 จุดในหน้า image-resize (กล่อง "ยังไม่มีไฟล์" ของ workspace ที่ถูก
        #    ซ่อนอยู่ + กล่องข้อความพรีวิวพัง) — ต้องเจาะจงด้วย has_text เหมือน tests/browser_files.py
        #    ไม่งั้น wait_for_selector ไปจับตัวที่ hidden แล้วรอตลอด (จับได้จริงตอนรันสคริปต์นี้)
        pg.locator(wait_selector, has_text=wait_has_text).first.wait_for(state="visible", timeout=15_000)
    else:
        pg.wait_for_selector(wait_selector, timeout=15_000)
    pg.wait_for_timeout(300)
    return capture_texts(pg, extra_exclude=[])


BROKEN_CASES = [
    # (tool_id, ไฟล์เสีย key, selector รอผล, ปุ่มที่ต้องกดก่อน (None = ประมวลผลอัตโนมัติตอนอัปโหลด),
    #  has_text ถ้า selector ใช้ซ้ำหลายจุดในหน้า)
    ("pdf-to-text", "garbage_pdf", ".status-wrap .status.show.err", "Extract text", None),
    ("excel-csv", "empty_xlsx", ".status-wrap .status.show.err", "Convert file", None),
    ("image-resize", "broken_jpg", ".ws-empty", None, "Could not create preview"),
    ("word-clean", "fake_docx", ".status-wrap .status.show.err", None, None),
]


# ─────────────────────────────────────────────────────────────────────────
# ⑤/⑥ document.title + สลับภาษาหลังมีผลลัพธ์ค้างอยู่
# ─────────────────────────────────────────────────────────────────────────
def check_title_and_switch(pg):
    """‼️ ต้องไม่ preset localStorage['fk-lang'] ด้วย ctx.add_init_script ในเทสนี้ — init script
    รันซ้ำทุกครั้งที่มีการโหลดเอกสารใหม่ "รวมถึง location.reload() ที่ setLang() เรียกตอนสลับภาษา"
    ถ้า preset ไว้ init script จะเขียนทับค่ากลับเป็นเดิมทันทีหลัง reload ทำให้ดูเหมือนสลับภาษาไม่ได้
    ทั้งที่จริงแอปทำงานถูกต้อง (จับได้จริงตอนเขียนเทสนี้ — เสียเวลาไล่บั๊กปลอมอยู่พักหนึ่งก่อนเจอสาเหตุจริง
    ใน Playwright เอง ไม่ใช่บั๊กของเว็บ) — เทสนี้จึงเริ่มจากค่าเริ่มต้นจริง (ไทย) แล้วกดสลับผ่าน UI ล้วน ๆ
    เหมือนผู้ใช้จริงเป๊ะ: เปิดเว็บ (ไทย) → สลับอังกฤษ → ใช้เครื่องมือ → สลับกลับไทย"""
    results = []

    goto_home(pg)
    title_th = pg.evaluate("document.title")
    results.append(("ค่าเริ่มต้นจริง (ไม่ preset) ต้องเป็นไทย", pg.evaluate("document.documentElement.lang") == "th",
                     f" (ได้ {pg.evaluate('document.documentElement.lang')!r})"))

    pg.locator('#lang .langopt[data-lang="en"]').click()
    pg.wait_for_timeout(900)
    title = pg.evaluate("document.title")
    results.append(("หลังกด EN: homepage document.title เป็นอังกฤษ ไม่มีไทย", not THAI.search(title), f" (ได้ {title!r})"))
    results.append(("หลังกด EN: homepage document.title มี 'FileKit'", "FileKit" in title, f" (ได้ {title!r})"))

    goto(pg, "pdf-merge")
    title2 = pg.evaluate("document.title")
    results.append(("pdf-merge document.title เป็นอังกฤษ ไม่มีไทย", not THAI.search(title2), f" (ได้ {title2!r})"))
    results.append(("pdf-merge document.title ลงท้าย '— FileKit'", title2.endswith("— FileKit"), f" (ได้ {title2!r})"))

    # ── สลับภาษาหลังมีผลลัพธ์ค้างอยู่: ทำ pdf-merge จนเสร็จ (โหมดอังกฤษ) แล้วกด TH ──
    src = [str(SAMPLES / "ตัวอย่าง-รายงานประจำเดือน.pdf"), str(SAMPLES / "ตัวอย่าง-ใบปะหน้าเอกสาร.pdf")]
    pg.set_input_files("input[type=file]", src, timeout=SIF_TIMEOUT)
    pg.wait_for_selector(".file-row", timeout=15_000)
    pg.get_by_role("button", name="Merge", exact=True).click()
    pg.wait_for_selector(".status-wrap .status.show", timeout=20_000)
    results.append(("ก่อนสลับภาษา: มีผลลัพธ์ปรากฏแล้ว (.results .result)", pg.locator(".results .result").count() > 0, ""))

    pg.locator('#lang .langopt[data-lang="th"]').click()
    pg.wait_for_timeout(900)
    lang_after = pg.evaluate("document.documentElement.lang")
    results.append(("สลับกลับ TH: documentElement.lang เปลี่ยนเป็น th จริง", lang_after == "th", f" (ได้ {lang_after!r})"))
    # ดีไซน์จริง: สลับภาษา = reload หน้า (src/i18n.js setLang) → ผลลัพธ์เก่าต้องถูกล้าง ไม่ใช่ค้างเป็นอังกฤษ/ไทยผสม
    stray_en_result = pg.locator(".results .result").count()
    results.append(("หลังสลับภาษา (reload) ผลลัพธ์เก่าถูกล้าง ไม่ค้างพันภาษา", stray_en_result == 0,
                     f" (เจอ .results .result ค้างอยู่ {stray_en_result} ชิ้น)"))
    left_texts = capture_texts(pg)
    leaks = english_leak(left_texts, extra_allow=stems(src))
    results.append(("หน้าเครื่องมือหลังสลับ→ไทย ไม่มีอังกฤษเป็นก้อนความหมายหลุดมา", not leaks,
                     f" เจอ: {[frag for _, frag in leaks[:5]]}" if leaks else ""))
    return results, []


# ─────────────────────────────────────────────────────────────────────────
# main sweep — เรียกได้ 2 รอบตาม DoD (ต้องไม่ flaky)
# ─────────────────────────────────────────────────────────────────────────
def run_full_suite(round_no, broken_files, address_xlsx):
    findings = []   # (kind, tool_id, detail) รายการปัญหาที่เจอจริง
    local_pass, local_fail = 0, 0

    def record(name, ok, detail=""):
        nonlocal local_pass, local_fail
        if ok:
            local_pass += 1
        else:
            local_fail += 1
        ck(name, ok, detail)

    with sync_playwright() as p:
        browser = p.chromium.launch()

        # ═══ ① EN forward: ทุกเครื่องมือที่ทำ action ได้ (28/29 — pdf-ocr ข้าม ดูท้ายไฟล์) ═══
        log(f"\n━━ รอบ {round_no} · ① โหมดอังกฤษ หลังลงมือทำจริง — {len(ACTIONABLE)} เครื่องมือ ━━")
        ctx = browser.new_context(viewport={"width": 1280, "height": 1000})
        ctx.add_init_script("try{localStorage.setItem('fk-lang','en')}catch(e){}")
        pg = ctx.new_page()
        for tool_id in sorted(ACTIONABLE):
            try:
                texts, extra_allow = do_tool_action(pg, tool_id, address_xlsx, lang="en")
                bad = thai_leftovers(texts, extra_allow=extra_allow)
                if bad:
                    findings.append(("thai-leftover-after-action", tool_id, bad[:10]))
                record(f"{tool_id} — ไม่มีไทยตกค้างหลังลงมือทำจริง (EN)", not bad,
                       "\n      เจอ: " + " | ".join(bad[:6]) if bad else "")
            except (PWTimeout, Exception) as e:
                findings.append(("runner-error", tool_id, str(e)[:300]))
                record(f"{tool_id} — รัน action ได้ไม่มี exception", False, f"\n      {type(e).__name__}: {str(e)[:200]}")
        ctx.close()

        # ═══ ② ไฟล์เสีย โหมดอังกฤษ ═══
        log(f"\n━━ รอบ {round_no} · ② ไฟล์เสีย โหมดอังกฤษ ━━")
        ctx = browser.new_context(viewport={"width": 1280, "height": 1000})
        ctx.add_init_script("try{localStorage.setItem('fk-lang','en')}catch(e){}")
        pg = ctx.new_page()
        for tool_id, key, sel, click_en, has_text in BROKEN_CASES:
            try:
                texts = check_broken(pg, tool_id, broken_files[key], sel, click_en=click_en, wait_has_text=has_text)
                bad = thai_leftovers(texts)
                if bad:
                    findings.append(("thai-leftover-broken-file", tool_id, bad[:10]))
                record(f"ไฟล์เสีย {tool_id} ({key}) — ข้อความ error ไม่มีไทยตกค้าง (EN)", not bad,
                       "\n      เจอ: " + " | ".join(bad[:6]) if bad else "")
            except Exception as e:
                findings.append(("runner-error-broken", f"{tool_id}/{key}", str(e)[:300]))
                record(f"ไฟล์เสีย {tool_id} ({key}) — รันได้ไม่มี exception", False, f"\n      {type(e).__name__}: {str(e)[:200]}")
        ctx.close()

        # ═══ ③ กล่องโต้ตอบ โหมดอังกฤษ ═══
        log(f"\n━━ รอบ {round_no} · ③ กล่องโต้ตอบ โหมดอังกฤษ ━━")
        ctx = browser.new_context(viewport={"width": 1280, "height": 1000})
        ctx.add_init_script("try{localStorage.setItem('fk-lang','en')}catch(e){}")
        pg = ctx.new_page()
        for name, fn in [("lightbox รูปภาพ", check_lightbox_image),
                          ("lightbox PDF (caption 'page 1 of N')", check_lightbox_pdf),
                          ("กล่องเตือนไฟล์ผิดชนิด (.wrong-type)", check_wrong_type)]:
            try:
                texts, extra_allow = fn(pg, "en")
                bad = thai_leftovers(texts, extra_allow=extra_allow)
                if bad:
                    findings.append(("thai-leftover-dialog", name, bad[:10]))
                record(f"{name} — ไม่มีไทยตกค้าง (EN)", not bad,
                       "\n      เจอ: " + " | ".join(bad[:6]) if bad else "")
            except Exception as e:
                findings.append(("runner-error-dialog", name, str(e)[:300]))
                record(f"{name} — รันได้ไม่มี exception", False, f"\n      {type(e).__name__}: {str(e)[:200]}")

        try:
            texts, _ = check_offline_bar(pg, "en")
            bad = thai_leftovers(texts)
            if bad:
                findings.append(("thai-leftover-offline", "offline-bar", bad[:10]))
            record("แถบออฟไลน์ทุกสถานะ (กำลังโหลด/เสร็จ) — ไม่มีไทยตกค้าง (EN)", not bad,
                   "\n      เจอ: " + " | ".join(bad[:6]) if bad else "")
        except Exception as e:
            findings.append(("runner-error-offline", "offline-bar", str(e)[:300]))
            record("แถบออฟไลน์ — รันได้ไม่มี exception", False, f"\n      {type(e).__name__}: {str(e)[:200]}")
        ctx.close()

        # ═══ ⑤/⑥ document.title + สลับภาษาหลังมีผลลัพธ์ ═══
        log(f"\n━━ รอบ {round_no} · ⑤⑥ document.title + สลับภาษาหลังมีผลลัพธ์ค้างอยู่ ━━")
        # ‼️ ห้าม preset localStorage['fk-lang'] ที่นี่ — ดูเหตุผลใน docstring ของ check_title_and_switch
        ctx = browser.new_context(viewport={"width": 1280, "height": 1000})
        pg = ctx.new_page()
        try:
            results, leftover = check_title_and_switch(pg)
            for name, ok, detail in results:
                record(name, ok, detail)
            if leftover:
                findings.append(("thai-leftover-after-lang-switch", "pdf-merge", leftover[:10]))
        except Exception as e:
            findings.append(("runner-error-switch", "lang-switch", str(e)[:300]))
            record("document.title/สลับภาษา — รันได้ไม่มี exception", False, f"\n      {type(e).__name__}: {str(e)[:200]}")
        ctx.close()

        # ═══ ⑦ ตรวจย้อนกลับ: โหมดไทย หลังลงมือทำจริง (ตัวแทน 10 เครื่องมือ) ═══
        log(f"\n━━ รอบ {round_no} · ⑦ ตรวจย้อนกลับ โหมดไทย หลังลงมือทำจริง — ตัวแทน {len(REVERSE_SAMPLE)} เครื่องมือ ━━")
        ctx = browser.new_context(viewport={"width": 1280, "height": 1000})
        ctx.add_init_script("try{localStorage.setItem('fk-lang','th')}catch(e){}")
        pg = ctx.new_page()
        for tool_id in REVERSE_SAMPLE:
            try:
                texts, extra_allow = do_tool_action(pg, tool_id, address_xlsx, lang="th")
                leaks = english_leak(texts, extra_allow=extra_allow)
                if leaks:
                    findings.append(("english-leak-th-mode", tool_id, leaks[:10]))
                record(f"{tool_id} — โหมดไทยหลังลงมือทำจริง ไม่มีอังกฤษทั้งประโยคหลุดมา", not leaks,
                       "\n      เจอ: " + " | ".join(f"{frag!r}" for _, frag in leaks[:6]) if leaks else "")
            except Exception as e:
                findings.append(("runner-error-reverse", tool_id, str(e)[:300]))
                record(f"{tool_id} — ตรวจย้อนกลับ รันได้ไม่มี exception", False, f"\n      {type(e).__name__}: {str(e)[:200]}")
        ctx.close()

        browser.close()

    return local_pass, local_fail, findings


def main():
    log(f"เว็บที่ทดสอบ: {BASE}\n")
    log(f"เครื่องมือทั้งหมดในทะเบียน: {len(ALL_TOOL_IDS)} ตัว")
    log(f"ทดสอบ action จริงได้: {len(ACTIONABLE)} ตัว")
    if SKIPPED_TOOLS:
        log(f"ข้าม (เหตุผลดูท้ายรายงาน): {sorted(SKIPPED_TOOLS)}")

    broken_files = make_broken_files()
    address_xlsx = make_address_xlsx()

    all_findings = []
    round_results = []
    try:
        for round_no in (1, 2):
            log("\n" + "#" * 70)
            log(f"# รอบที่ {round_no}/2 (เช็คว่าไม่ flaky ตาม DoD)")
            log("#" * 70)
            passed, failed, findings = run_full_suite(round_no, broken_files, address_xlsx)
            round_results.append((round_no, passed, failed))
            all_findings.extend((round_no, *fnd) for fnd in findings)
    finally:
        shutil.rmtree(WORKDIR, ignore_errors=True)

    log("\n" + "═" * 70)
    log("สรุปผลรัน 2 รอบ")
    log("═" * 70)
    for round_no, passed, failed in round_results:
        log(f"  รอบ {round_no}: ผ่าน {passed} · ตก {failed}")

    log("\n" + "═" * 70)
    log("ตารางข้อความที่ยังไม่ได้แปล / น่าสงสัย (ถ้ามี)")
    log("═" * 70)
    if not all_findings:
        log("  (ไม่พบ — ทุกจังหวะที่ไล่ตรวจ แปลครบทั้ง EN forward และ TH reverse)")
    else:
        seen = set()
        for round_no, kind, tool_id, detail in all_findings:
            key = (kind, tool_id, str(detail))
            if key in seen:
                continue
            seen.add(key)
            log(f"  [{kind}] {tool_id}")
            if isinstance(detail, list):
                for d in detail:
                    log(f"      - {d}")
            else:
                log(f"      - {detail}")

    log("\n" + "═" * 70)
    log("ขอบเขตที่ไล่ตรวจจริง")
    log("═" * 70)
    log(f"  · EN forward (หลังลงมือทำจริง): {len(ACTIONABLE)}/{len(ALL_TOOL_IDS)} เครื่องมือ, "
        f"ทั้งข้อความสำเร็จ/สถานะ/ผลลัพธ์/aria-label/title/placeholder ที่ยังเห็นบนจอ")
    log(f"  · ไฟล์เสีย (EN): {len(BROKEN_CASES)} เคส ครอบคลุม pdf.js / xlsx / canvas-image / docx-zip")
    log(f"  · กล่องโต้ตอบ (EN): lightbox รูป, lightbox PDF (caption หน้า), เตือนไฟล์ผิดชนิด, แถบออฟไลน์ทุกสถานะ")
    log(f"  · document.title: หน้าแรก + หน้าเครื่องมือ")
    log(f"  · สลับภาษาหลังมีผลลัพธ์ค้างอยู่: pdf-merge (produce EN result → switch → reload ล้างสถานะ)")
    log(f"  · ตรวจย้อนกลับ (TH หลังลงมือทำจริง): {len(REVERSE_SAMPLE)} เครื่องมือตัวแทนทุกตระกูลโค้ด "
        f"({', '.join(REVERSE_SAMPLE)})")
    log(f"  · ไม่ได้รัน action จริง (เหตุผล): pdf-ocr — ต้องโหลด Tesseract language pack ~10-30MB "
        f"จาก CDN ครั้งแรก เสี่ยง flaky/ช้าเกินจะพิสูจน์ 'ไม่ flaky'; ตรวจ static แล้วพบทุกข้อความ "
        f"(รวม progress ระหว่างอ่าน) ผ่าน tr() ครบใน src/tools/pdf-ocr.js, src/ocr.js ไม่ import i18n "
        f"เพราะไม่มีข้อความ user-facing (มีแต่ logic ล้วน)")

    n_fail_total = sum(f for _, _, f in round_results)
    n_findings = len({(k, t, str(d)) for _, k, t, d in all_findings if k.startswith("thai-leftover") or k.startswith("english-leak")})
    sys.exit(1 if (n_fail_total or n_findings) else 0)


if __name__ == "__main__":
    main()
