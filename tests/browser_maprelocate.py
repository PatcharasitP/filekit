# เครื่องมือ "แผนที่การย้ายที่ตั้ง" ตรวจของที่ผู้ใช้เอาไปใช้ต่อจริง
#
# ‼️ ที่เทสนี้จับ
#    ① เดาคอลัมน์พิกัดถูกเอง และระยะทางบนจอตรงกับค่าที่ DAX คำนวณไว้ในงานจริง
#    ② แผนที่วาดของจริงลงบน canvas (ตรวจด้วยการนับพิกเซลที่เปลี่ยนสี ไม่ใช่แค่ว่ามี element)
#    ③ คลิกคู่ในตารางแล้วแผนที่ซูมจริง (ภาพเปลี่ยน) และแถบบอกคู่ที่เลือกขึ้นถูกคู่
#    ④ ไฟล์ที่ดาวน์โหลดมีชีตครบและค่า WKT ถูกต้อง (ลองจิจูดมาก่อนละติจูด) พร้อมใส่ Icon Map Pro
#    ⑤ แถวที่พิกัดเสียต้องถูกข้ามและรายงาน ไม่ใช่กลืนเงียบหรือวาดจุดผิดที่
#
# รัน: ../.venv/bin/python tests/browser_maprelocate.py  (หรือ tests/run.sh browser_maprelocate)
import math
import os
import pathlib
import socket
import subprocess
import sys
import tempfile
import time

import openpyxl
from playwright.sync_api import sync_playwright

ROOT = pathlib.Path(__file__).resolve().parent.parent
TMP = pathlib.Path(tempfile.mkdtemp(prefix="filekit_maprel_"))
P, F = 0, []


def ck(name, got, want):
    global P
    ok_ = got == want
    if ok_:
        P += 1
    else:
        F.append(f"{name}\n      ได้    : {got!r}\n      ควรได้ : {want!r}")
    print(f"  {'✅' if ok_ else '❌'} {name}")


def ok(name, cond, note=""):
    global P
    if cond:
        P += 1
    else:
        F.append(f"{name} {note}")
    print(f"  {'✅' if cond else '❌'} {name}{(' ' + note) if note and not cond else ''}")


def free_port():
    s = socket.socket()
    s.bind(("127.0.0.1", 0))
    port = s.getsockname()[1]
    s.close()
    return port


# ข้อมูลจริง 8 คู่แรกจากงานย้ายสถานี พร้อมระยะทางที่ DAX คำนวณไว้
REAL = [
    ("BKK1000", 13.56586, 100.503058, "BKK9000", 13.503432, 100.498311, "กรุงเทพมหานคร", 6.961),
    ("CMI1001", 18.589618, 98.846942, "CMI9001", 18.556227, 98.82117, "เชียงใหม่", 4.601),
    ("PKT1002", 8.022087, 98.416206, "PKT9002", 8.015232, 98.36025, "ภูเก็ต", 6.208),
    ("KKN1003", 16.567413, 102.87926, "KKN9003", 16.570288, 102.926302, "ขอนแก่น", 5.024),
    ("NMA1004", 14.785268, 102.256181, "NMA9004", 14.728768, 102.27518, "นครราชสีมา", 6.606),
    ("HYI1005", 7.101189, 100.564394, "HYI9005", 7.106192, 100.584013, "สงขลา", 2.235),
    ("CBI1006", 13.601056, 101.199359, "CBI9006", 13.625614, 101.129836, "ชลบุรี", 7.994),
    ("AYA1007", 14.359994, 100.35706, "AYA9007", 14.337653, 100.392958, "พระนครศรีอยุธยา", 4.596),
]
BAD = [
    ("BAD00001", "ไม่ทราบ", 100.5, "BAD9001", 13.6, 100.6, "กรุงเทพมหานคร", None),   # พิกัดเดิมอ่านไม่ออก
    ("BAD00002", 0, 0, "BAD9002", 13.6, 100.6, "ชลบุรี", None),                      # 0,0 คือค่าว่างที่ปลอมเป็นพิกัด
]


def make_file():
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "RELOCATE"
    ws.append(["SITE_CODE_OLD", "LAT_OLD", "LON_OLD", "SITE_CODE_NEW", "LAT_NEW", "LON_NEW", "PROVINCE"])
    for r in REAL + BAD:
        ws.append(list(r[:7]))
    p = TMP / "ย้ายที่ตั้งสถานี.xlsx"
    wb.save(p)
    return p


def canvas_ink(pg):
    """นับพิกเซลที่ไม่ใช่สีพื้น เพื่อยืนยันว่าวาดของจริง ไม่ใช่ผืนว่าง"""
    return pg.evaluate("""() => {
      const c = document.querySelector('.mr-canvas');
      const ctx = c.getContext('2d');
      const d = ctx.getImageData(0, 0, c.width, c.height).data;
      const first = [d[0], d[1], d[2]];
      let diff = 0;
      for (let i = 0; i < d.length; i += 4 * 37) {          // สุ่มอ่านทุก 37 พิกเซลก็พอ
        if (Math.abs(d[i] - first[0]) + Math.abs(d[i+1] - first[1]) + Math.abs(d[i+2] - first[2]) > 24) diff++;
      }
      return { diff, w: c.width, h: c.height };
    }""")


def dot_pixels(pg):
    """นับพิกเซลที่เป็นสีจุดเดิม (ส้ม) และจุดใหม่ (น้ำเงิน) บนภาพจริง

    ‼️ เกิดจากบั๊กจริง 16/09/2026: กล่องพื้นของป้ายระยะทางวางทับจุดกึ่งกลางพอดี
       จุดทั้งสองของทุกคู่จึงถูกบังมิด เหลือแต่ตัวเลขลอยอยู่บนแผนที่ เทสเดิมที่ดูแค่
       ว่ามีหมึกบนผืนภาพผ่านฉลุย เพราะป้ายก็เป็นหมึกเหมือนกัน
    """
    return pg.evaluate("""() => {
      const c = document.querySelector('.mr-canvas');
      const d = c.getContext('2d').getImageData(0, 0, c.width, c.height).data;
      let orange = 0, blue = 0;
      for (let i = 0; i < d.length; i += 4) {
        const r = d[i], g = d[i+1], b = d[i+2];
        if (r > 180 && g > 90 && g < 170 && b < 110) orange++;
        else if (b > 130 && r < 110 && g > 80 && g < 150) blue++;
      }
      return { orange, blue };
    }""")


def canvas_hash(pg):
    return pg.evaluate("""() => {
      const c = document.querySelector('.mr-canvas');
      const d = c.getContext('2d').getImageData(0, 0, c.width, c.height).data;
      let h = 0;
      for (let i = 0; i < d.length; i += 4 * 101) h = (h * 31 + d[i] + d[i+1] * 3 + d[i+2] * 7) >>> 0;
      return h;
    }""")


def table_rows(pg):
    return pg.evaluate("""() => [...document.querySelectorAll('.mr-pairs tbody tr')].map(tr => {
      const td = tr.querySelectorAll('td');
      return { old: td[0].textContent.trim(), neu: td[1].textContent.trim(),
               prov: td[2].textContent.trim(), km: td[3].textContent.trim(), dir: td[5].textContent.trim() };
    })""")


def main():
    base = os.environ.get("FK_BASE")
    server = None
    if not base:
        port = free_port()
        server = subprocess.Popen([sys.executable, "-m", "http.server", str(port)],
                                  cwd=str(ROOT), stdout=subprocess.DEVNULL, stderr=subprocess.DEVNULL)
        base = f"http://localhost:{port}"
        time.sleep(1.5)

    xlsx = make_file()
    try:
        with sync_playwright() as pw:
            b = pw.chromium.launch()
            ctx = b.new_context(viewport={"width": 1400, "height": 1000}, accept_downloads=True,
                                timezone_id="Asia/Bangkok")
            pg = ctx.new_page()
            errs = []
            pg.on("pageerror", lambda e: errs.append(str(e)[:200]))
            pg.on("console", lambda m: errs.append(m.text[:200]) if m.type == "error" else None)

            pg.goto(f"{base}/#/map-relocate", wait_until="networkidle")
            pg.wait_for_selector(".dz")
            pg.locator(".dz input[type=file]").first.set_input_files(str(xlsx))
            pg.wait_for_timeout(2200)

            print("\n── ① เดาคอลัมน์พิกัดเองได้ ──")
            picked = pg.evaluate("""() => [...document.querySelectorAll('.ws-left select')]
              .map(s => s.options[s.selectedIndex]?.textContent || '')""")
            for want in ["LAT_OLD", "LON_OLD", "LAT_NEW", "LON_NEW", "PROVINCE"]:
                ok(f"จับคู่ {want} ได้เอง", want in picked, f"ได้ {picked}")

            print("\n── ② อ่านคู่ย้ายได้ครบ และรายงานแถวที่พิกัดเสีย ──")
            chips = pg.evaluate("() => [...document.querySelectorAll('.mr-chip')].map(c => c.textContent)")
            ok(f"บอกว่ามีคู่ย้าย {len(REAL)} คู่",
               any(f"{len(REAL)}" in c and "คู่ย้าย" in c for c in chips), f"ได้ {chips}")
            ok("บอกว่าข้ามแถวที่พิกัดไม่ครบ 2 แถว",
               any("ข้าม" in c and "2" in c for c in chips), f"ได้ {chips}")
            far = max(REAL, key=lambda r: r[7])
            ok(f"บอกระยะไกลสุด {far[7]:.2f} กม.",
               any(f"{far[7]:.2f}" in c for c in chips), f"ได้ {chips}")

            print("\n── ③ แผนที่ต้องวาดของจริงลงผืนภาพ ──")
            ink = canvas_ink(pg)
            ok(f"ผืนภาพมีขนาดจริง {ink['w']}x{ink['h']}", ink["w"] > 300 and ink["h"] > 300, f"ได้ {ink}")
            ok("มีเนื้อภาพวาดอยู่จริง ไม่ใช่ผืนว่าง", ink["diff"] > 40, f"พิกเซลที่ต่างจากพื้น {ink['diff']}")
            ok("กรอบภาพสูงกว่ากว้าง เพราะประเทศไทยเป็นแนวตั้ง",
               ink["h"] >= ink["w"] * 0.9, f"ได้ {ink['w']}x{ink['h']}")
            dots = dot_pixels(pg)
            ok(f"เห็นจุดสถานีเดิมสีส้มจริงบนภาพ ({dots['orange']} พิกเซล)", dots["orange"] > 60, f"ได้ {dots}")
            ok(f"เห็นจุดสถานีใหม่สีน้ำเงินจริงบนภาพ ({dots['blue']} พิกเซล)", dots["blue"] > 60, f"ได้ {dots}")
            ok("ป้ายระยะทางไม่บังจุดจนหาย (จำนวนจุดสองสีใกล้เคียงกัน)",
               min(dots["orange"], dots["blue"]) > max(dots["orange"], dots["blue"]) * 0.45, f"ได้ {dots}")

            print("\n── ④ ตารางคู่ย้าย ระยะทางต้องตรงกับที่ DAX คำนวณไว้ ──")
            pg.locator(".seg-item", has_text="ตาราง").first.click()
            pg.wait_for_timeout(500)
            rows = table_rows(pg)
            ck("จำนวนแถวเท่ากับคู่ที่อ่านได้", len(rows), len(REAL))
            want_km = {r[0]: f"{r[7]:.2f}" for r in REAL}
            got_km = {r["old"]: r["km"] for r in rows}
            ck("ระยะทางทุกคู่ตรงกับค่าที่ DAX คำนวณไว้", got_km, want_km)
            ck("เรียงจากไกลไปใกล้ให้ตั้งแต่แรก", rows[0]["old"], far[0])
            ck("จังหวัดติดมาถูกคู่", {r["old"]: r["prov"] for r in rows}[far[0]], far[6])
            bkk = [r for r in rows if r["old"] == "BKK1000"][0]
            ck("ทิศทางคำนวณถูก (คู่นี้ย้ายลงใต้)", bkk["dir"], "ใต้")

            print("\n── ⑤ คลิกคู่ในตารางแล้วแผนที่ต้องซูมจริง ──")
            pg.locator(".seg-item", has_text="แผนที่").first.click()
            pg.wait_for_timeout(400)
            before = canvas_hash(pg)
            pg.locator(".seg-item", has_text="ตาราง").first.click()
            pg.wait_for_timeout(300)
            pg.locator(".mr-pairs tbody tr", has_text=far[0]).first.click()
            pg.wait_for_timeout(700)
            after = canvas_hash(pg)
            ok("ภาพแผนที่เปลี่ยนไปจริงหลังเลือกคู่ (ไม่ใช่แค่ไฮไลต์ตาราง)", before != after, f"เดิม {before} ใหม่ {after}")
            selbar = pg.evaluate("() => document.querySelector('.mr-sel')?.textContent || ''")
            ok("แถบบนบอกคู่ที่เลือกถูกคู่", far[0] in selbar and f"{far[7]:.2f}" in selbar, f"ได้ {selbar!r}")
            pg.locator(".mr-sel button").first.click()
            pg.wait_for_timeout(600)
            ok("กดดูทั้งประเทศแล้วกลับมาเหมือนเดิม", canvas_hash(pg) == before)

            print("\n── ⑤ข ซูมพอดีข้อมูล และป้ายต้องไม่ทับกัน ──")
            pg.locator(".seg-item", has_text="แผนที่").first.click()
            pg.wait_for_timeout(400)
            country = canvas_hash(pg)
            pg.evaluate("""() => {
              const s = [...document.querySelectorAll('.ws-right select')]
                .find(x => [...x.options].some(o => o.textContent.includes('ซูมพอดี')));
              s.value = 'data'; s.dispatchEvent(new Event('change', { bubbles: true }));
            }""")
            pg.wait_for_timeout(700)
            ok("สลับเป็นซูมพอดีข้อมูลแล้วภาพเปลี่ยนจริง", canvas_hash(pg) != country)
            pg.evaluate("""() => {
              const s = [...document.querySelectorAll('.ws-right select')]
                .find(x => [...x.options].some(o => o.textContent.includes('ทุกคู่')));
              s.value = 'all'; s.dispatchEvent(new Event('change', { bubbles: true }));
            }""")
            pg.wait_for_timeout(700)
            legend = pg.evaluate("() => document.querySelector('.mr-legend')?.textContent || ''")
            ok("เปิดป้ายทุกคู่แล้วบอกตรง ๆ ว่าซ่อนป้ายที่ทับกันไปกี่ป้าย",
               "ซ่อนป้าย" in legend, f"ได้ {legend!r}")
            pg.evaluate("""() => {
              const s = [...document.querySelectorAll('.ws-right select')]
                .find(x => [...x.options].some(o => o.textContent.includes('ทุกคู่')));
              s.value = 'top'; s.dispatchEvent(new Event('change', { bubbles: true }));
              const e = [...document.querySelectorAll('.ws-right select')]
                .find(x => [...x.options].some(o => o.textContent.includes('ซูมพอดี')));
              e.value = 'country'; e.dispatchEvent(new Event('change', { bubbles: true }));
            }""")
            pg.wait_for_timeout(600)

            print("\n── ⑥ สรุปรายจังหวัด ──")
            pg.locator(".seg-item", has_text="จังหวัด").first.click()
            pg.wait_for_timeout(400)
            provs = pg.evaluate("""() => [...document.querySelectorAll('.mr-provs tbody tr')]
              .map(tr => tr.querySelectorAll('td')[0].textContent.trim())""")
            ck("มีครบทุกจังหวัดที่อยู่ในไฟล์", sorted(provs), sorted({r[6] for r in REAL}))

            print("\n── ⑦ กรองระยะขั้นต่ำ ต้องมีผลจริง ──")
            pg.evaluate("""() => {
              const i = [...document.querySelectorAll('.ws-right input.mr-num')]
                .find(x => x.placeholder && (x.placeholder.includes('ไม่กรอง') || x.placeholder.includes('no filter')));
              i.value = '6'; i.dispatchEvent(new Event('change', { bubbles: true }));
            }""")
            pg.wait_for_timeout(600)
            pg.locator(".seg-item", has_text="ตาราง").first.click()
            pg.wait_for_timeout(400)
            left = [r["old"] for r in table_rows(pg)]
            ck("เหลือเฉพาะคู่ที่ย้ายไกลกว่า 6 กม.", sorted(left), sorted([r[0] for r in REAL if r[7] >= 6]))
            pg.evaluate("""() => {
              const i = [...document.querySelectorAll('.ws-right input.mr-num')]
                .find(x => x.placeholder && (x.placeholder.includes('ไม่กรอง') || x.placeholder.includes('no filter')));
              i.value = ''; i.dispatchEvent(new Event('change', { bubbles: true }));
            }""")
            pg.wait_for_timeout(500)

            print("\n── ⑧ ไฟล์สำหรับ Power BI ต้องพร้อมใส่ Icon Map Pro ──")
            with pg.expect_download() as dl:
                pg.locator(".ws-footer button", has_text="Power BI").first.click()
            out = TMP / "pbi.xlsx"
            dl.value.save_as(str(out))
            wb2 = openpyxl.load_workbook(out)
            ck("มีชีตครบตามที่ Power BI ต้องใช้", wb2.sheetnames, ["MAPSHAPES", "PAIR", "PATH_UNPIVOT"])
            ms = wb2["MAPSHAPES"]
            head = [c.value for c in ms[1]]
            ok("MAPSHAPES มีคอลัมน์ KIND และ WKT", "KIND" in head and "WKT" in head, f"ได้ {head}")
            rows_ms = list(ms.iter_rows(min_row=2, values_only=True))
            ck("หนึ่งคู่ให้ 4 แถว (จุดเดิม จุดใหม่ เส้น ป้าย)", len(rows_ms), len(REAL) * 4)
            ki, wi, la, lo = head.index("KIND"), head.index("WKT"), head.index("LATITUDE"), head.index("LONGITUDE")
            line = [r for r in rows_ms if r[ki] == "LINE"][0]
            ok("แถวเส้นมี WKT LINESTRING", str(line[wi]).startswith("LINESTRING ("), f"ได้ {line[wi]!r}")
            nums = [float(x) for x in str(line[wi]).replace("LINESTRING (", "").replace(")", "").replace(",", " ").split()]
            ok("ลองจิจูดมาก่อนละติจูดใน WKT (ถ้าสลับ วิชวลจะไม่ขึ้นโดยไม่มี error)",
               nums[0] > 90 and nums[1] < 25, f"ได้ {nums[:2]}")
            ck("แถวเส้นต้องไม่มีพิกัดจุด ไม่งั้นจะมีหมุดเกินมาบนแผนที่", (line[la], line[lo]), (None, None))
            # ระยะทางในชีตต้องตรงกับค่าที่ DAX คำนวณไว้ทุกคู่
            pair = wb2["PAIR"]
            ph = [c.value for c in pair[1]]
            di, ci = ph.index("DISTANCE_KM"), ph.index("SITE_CODE_OLD")
            got = {r[ci]: round(float(r[di]), 2) for r in pair.iter_rows(min_row=2, values_only=True)}
            ck("ระยะทางในไฟล์ตรงกับค่าที่ DAX คำนวณไว้ทุกคู่", got, {r[0]: round(r[7], 2) for r in REAL})

            print("\n── ⑨ บันทึกเป็นภาพได้จริง ──")
            pg.locator(".seg-item", has_text="แผนที่").first.click()
            pg.wait_for_timeout(400)
            with pg.expect_download() as dl2:
                pg.locator(".ws-footer button", has_text="PNG").first.click()
            png = TMP / "map.png"
            dl2.value.save_as(str(png))
            size = png.stat().st_size
            ok(f"ได้ไฟล์ภาพขนาด {size:,} ไบต์", size > 20000, f"ได้ {size}")
            head8 = png.read_bytes()[:8]
            ck("เป็นไฟล์ PNG จริง", head8, b"\x89PNG\r\n\x1a\n")

            print("\n── ⑨ข สลับโหมดมืดแล้วแผนที่ต้องวาดใหม่ตามธีม ──")
            pg.locator(".seg-item", has_text="แผนที่").first.click()
            pg.wait_for_timeout(400)
            light = canvas_hash(pg)
            light_px = pg.evaluate("""() => {
              const c = document.querySelector('.mr-canvas');
              const d = c.getContext('2d').getImageData(4, 4, 1, 1).data;
              return d[0] + d[1] + d[2];
            }""")
            pg.evaluate("() => document.documentElement.setAttribute('data-theme', 'dark')")
            pg.wait_for_timeout(900)
            dark_px = pg.evaluate("""() => {
              const c = document.querySelector('.mr-canvas');
              const d = c.getContext('2d').getImageData(4, 4, 1, 1).data;
              return d[0] + d[1] + d[2];
            }""")
            ok("ภาพแผนที่เปลี่ยนตามธีมจริง (ไม่ใช่ผืนสว่างค้างอยู่บนหน้ามืด)",
               canvas_hash(pg) != light and dark_px < light_px, f"สว่าง {light_px} มืด {dark_px}")
            pg.evaluate("() => document.documentElement.setAttribute('data-theme', 'light')")
            pg.wait_for_timeout(700)

            print("\n── ⑩ ไม่มี error ในคอนโซล ──")
            real_errs = [e for e in errs if "favicon" not in e.lower()]
            ck("เงียบสนิทตลอดการใช้งาน", real_errs, [])

            b.close()
    finally:
        if server:
            server.terminate()

    print(f"\nสรุป: ผ่าน {P} ข้อ, ตก {len(F)} ข้อ")
    for f in F:
        print("  ❌ " + f)
    sys.exit(1 if F else 0)


if __name__ == "__main__":
    main()
