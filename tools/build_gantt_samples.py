"""สร้าง samples/powerbi/pbi-gantt-samples.xlsx จาก samples/powerbi/gantt-datasets.json

ทุกชุดข้อมูลเป็นชีตละชุด (วันที่เป็นวันที่จริงของ Excel ไม่ใช่ข้อความ) แล้วเพิ่มอีกหนึ่งชีตชื่อ "ลองข้อมูลมีจุดผิด"
ที่ใส่ของเสียไว้ตั้งใจ: วันที่ พ.ศ. แบบข้อความ, วันจบก่อนวันเริ่ม, ไม่มีชื่องาน, วันที่ไม่มีจริง
เอาไว้ให้ผู้ใช้ลองช่อง "ใช้ข้อมูลของคุณเอง" แล้วเห็นว่าเครื่องมือบอกแถวที่ข้ามพร้อมเลขบรรทัด

รัน: python3 tools/build_gantt_samples.py   (จากโฟลเดอร์ FileKit/)
เทส tests/browser_gantt.py เทียบค่าในไฟล์กับ JSON ทุกแถว ถ้าแก้ JSON ต้องรันไฟล์นี้ใหม่
"""
import datetime as dt
import json
import pathlib

from openpyxl import Workbook
from openpyxl.styles import Font
from openpyxl.utils import get_column_letter

ROOT = pathlib.Path(__file__).resolve().parent.parent
SRC = ROOT / "samples/powerbi/gantt-datasets.json"
OUT = ROOT / "samples/powerbi/pbi-gantt-samples.xlsx"
HEAD = ["ชื่องาน", "วันเริ่ม", "วันจบ", "กลุ่ม", "ความคืบหน้า"]
BAD_SHEET = "ลองข้อมูลมีจุดผิด"


def day(iso):
    return dt.datetime.strptime(iso, "%Y-%m-%d")


def put_header(ws):
    ws.append(HEAD)
    for c in ws[1]:
        c.font = Font(bold=True)
    for i, w in enumerate([34, 13, 13, 20, 13], 1):
        ws.column_dimensions[get_column_letter(i)].width = w


def main():
    data = json.loads(SRC.read_text(encoding="utf-8"))
    wb = Workbook()
    wb.remove(wb.active)
    for key, info in data.items():
        ws = wb.create_sheet(key[:31])
        put_header(ws)
        for r in info["rows"]:
            ws.append([r["Task"], day(r["Start"]), day(r["End"]), r.get("Group"), r.get("Progress")])
        for row in ws.iter_rows(min_row=2, min_col=2, max_col=3):
            for c in row:
                c.number_format = "dd/mm/yyyy"
        for c in ws["E"][1:]:
            c.number_format = "0%"

    bad = wb.create_sheet(BAD_SHEET)
    put_header(bad)
    # วันที่เป็นข้อความปี พ.ศ. (พิมพ์มือแบบที่ทีมงานไทยทำจริง) และของเสีย 3 แบบ อยู่บรรทัด 4, 5, 6 ของชีต (หัวตาราง = บรรทัด 1)
    bad.append(["เตรียมพื้นที่", "01/07/2569", "05/07/2569", "เตรียมงาน", "100%"])
    bad.append(["เทฐานราก", "06/07/2569", "20/07/2569", "โครงสร้าง", "60%"])
    bad.append(["งานที่วันจบมาก่อนวันเริ่ม", "25/07/2569", "10/07/2569", "โครงสร้าง", "0%"])
    bad.append([None, "01/08/2569", "10/08/2569", "งานระบบ", "0%"])
    bad.append(["งานที่วันที่ไม่มีจริง", "31/04/2569", "10/05/2569", "งานระบบ", "0%"])
    bad.append(["ก่ออิฐผนัง", "21/07/2569", "15/08/2569", "งานผนัง", "20%"])
    wb.save(OUT)
    print(f"เขียน {OUT.relative_to(ROOT)} ชีต {wb.sheetnames}")


if __name__ == "__main__":
    main()
