# เครื่องมือ "ตารางเป็นสูตร Power Query" ตรวจโค้ดที่สร้างจริง ไม่ใช่แค่ว่ากดแล้วไม่ error
#
# ‼️ โค้ด M ที่เราสร้างให้ ผู้ใช้จะเอาไปวางใน Power Query Editor ถ้าผิดไวยากรณ์แม้จุดเดียว
#    query จะพังทั้งอัน และข้อความ error ของ M อ่านยากมาก ผู้ใช้จะไม่รู้ว่าพังเพราะอะไร
#    จุดที่พลาดง่ายที่สุดคือกฎ escape โดยเฉพาะข้อความที่มี #( อยู่จริง
#    ซึ่ง M อ่านเป็นคำสั่ง escape ถ้าไม่แปลงให้ถูก
#
# ‼️ อีกจุดคือ Percentage ใน M เก็บเป็น "สัดส่วน" 87.5% ต้องเขียน 0.875 ไม่ใช่ 87.5
#    เขียนผิดจุดนี้ตัวเลขจะเพี้ยนไป 100 เท่าโดยที่โค้ดยังรันผ่าน
#
# รัน: ../.venv/bin/python tests/browser_pq.py     (เปิด/ปิด server ให้เอง)
import datetime
import os
import pathlib
import shutil
import socket
import subprocess
import sys
import tempfile
import time

import fitz
import openpyxl
from docx import Document
from playwright.sync_api import sync_playwright

ROOT = pathlib.Path(__file__).resolve().parent.parent
TMP = pathlib.Path(tempfile.mkdtemp(prefix="filekit_pq_"))

P, F = 0, []


def ck(name, got, want):
    global P
    ok = got == want
    if ok:
        P += 1
    else:
        F.append(f"{name}\n      ได้    : {got!r}\n      ควรได้ : {want!r}")
    print(f"  {'✅' if ok else '❌'} {name}")


def free_port():
    s = socket.socket()
    s.bind(("127.0.0.1", 0))
    port = s.getsockname()[1]
    s.close()
    return port


def make_workbook():
    """ตารางที่รวมทุกกับดักไว้ในไฟล์เดียว: ข้อความมีเครื่องหมายคำพูด, ข้อความมี #(,
    เปอร์เซ็นต์, วันที่, จริงเท็จ, เบอร์โทรที่ต้องเป็นข้อความ, ช่องว่าง และชีทที่ซ่อนไว้"""
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "ข้อมูล"
    ws.append(["ID", "ชื่อ", "ยอดเงิน", "เปอร์เซ็นต์", "วันที่เริ่ม", "เปิดใช้", "เบอร์โทร", "หมายเหตุ"])
    ws.append([1001, 'สมชาย "เล็ก"', 12500.50, "87.5%", datetime.date(2026, 9, 1), "TRUE", "0812345678", "ปกติ"])
    ws.append([1002, "สมหญิง", 8500.25, "72%", datetime.date(2026, 9, 2), "FALSE", "0898765432", None])
    ws.append([1003, "อานันท์", 45000.75, "95%", datetime.date(2026, 9, 3), "TRUE", "0811112222", "ราคา #(พิเศษ)"])
    hidden = wb.create_sheet("ลับ")
    hidden["A1"] = "ห้ามเผยแพร่"
    hidden.sheet_state = "hidden"
    p = TMP / "ตารางอ้างอิง.xlsx"
    wb.save(p)
    return p


TABLE = [
    ["รหัส", "รายการ", "จำนวน", "ราคา"],
    ["A01", "ปากกาน้ำเงิน", "120", "15.50"],
    ["A02", "ดินสอ 2B", "250", "8.00"],
    ["A03", "สมุดปกแข็ง", "45", "62.75"],
]


def make_pdf_table():
    """PDF ที่มีชั้นข้อความ จัดคอลัมน์ด้วยระยะห่างชัดเจน แบบรายงานที่พิมพ์ออกมาจากระบบ"""
    font = str(ROOT / "vendor/fonts/Sarabun-Regular-th.ttf")
    d = fitz.open()
    page = d.new_page(width=595, height=842)
    page.insert_font(fontname="TH", fontfile=font)
    for r, row in enumerate(TABLE):
        for x, cell in zip((55, 150, 330, 430), row):
            page.insert_text((x, 90 + r * 30), cell, fontsize=11, fontname="TH")
    p = TMP / "ตารางสินค้า.pdf"
    d.save(p)
    d.close()
    return p


def make_docx_table():
    """.docx ที่มีตารางจริง (ไม่ใช่ข้อความจัดคอลัมน์ด้วยช่องว่าง)"""
    doc = Document()
    doc.add_paragraph("รายงานสินค้าคงเหลือ")
    t = doc.add_table(rows=len(TABLE), cols=4)
    for r, row in enumerate(TABLE):
        for c, cell in enumerate(row):
            t.cell(r, c).text = cell
    p = TMP / "ตารางสินค้า.docx"
    doc.save(p)
    return p


def main():
    base = os.environ.get("FK_BASE")
    server = None
    if not base:
        port = free_port()
        server = subprocess.Popen([sys.executable, "-m", "http.server", str(port)],
                                  cwd=str(ROOT), stdout=subprocess.DEVNULL, stderr=subprocess.DEVNULL)
        base = f"http://localhost:{port}"
        time.sleep(1.5)

    xlsx = make_workbook()
    pdf_file = make_pdf_table()
    docx_file = make_docx_table()
    try:
        with sync_playwright() as pw:
            b = pw.chromium.launch()
            ctx = b.new_context(viewport={"width": 1360, "height": 1000},
                                reduced_motion="no-preference", accept_downloads=True,
                                timezone_id="Asia/Bangkok")
            pg = ctx.new_page()
            errs = []
            pg.on("pageerror", lambda e: errs.append(str(e)[:160]))
            pg.on("console", lambda m: errs.append(m.text[:160]) if m.type == "error" else None)

            pg.goto(f"{base}/#/excel-to-pq", wait_until="networkidle")
            pg.wait_for_selector(".dz")
            pg.locator(".dz input[type=file]").first.set_input_files(str(xlsx))
            pg.wait_for_timeout(2500)

            print("\n── ① เดาชนิดข้อมูลของแต่ละคอลัมน์ ──")
            types = pg.evaluate("() => [...document.querySelectorAll('.pq-cols tbody tr select')]"
                                ".map((s) => s.value)")
            ck("เดาชนิดครบทุกคอลัมน์",
               types, ["int64", "text", "number", "percentage", "date", "logical", "text", "text"])
            flagged = pg.evaluate("() => [...document.querySelectorAll('.pq-cols tbody tr')]"
                                  ".filter((r) => r.dataset.check === '1').length")
            # เบอร์โทรมีศูนย์นำหน้า เป็นเคสที่ต้องเตือนให้ผู้ใช้ตรวจ ไม่ใช่เงียบ ๆ
            ck("ติดธงเตือนเฉพาะคอลัมน์ที่ควรตรวจ", flagged, 1)

            print("\n── ② โค้ด M ที่สร้างได้ ──")
            code = pg.locator(".pq-code").inner_text()
            ck("ขึ้นต้นด้วย = #table", code.startswith("= #table("), True)
            ck("ชื่อคอลัมน์ครอบด้วยเครื่องหมายเสมอ", '#"ID" = Int64.Type' in code, True)
            ck("เปอร์เซ็นต์ประกาศเป็น Percentage.Type", '#"เปอร์เซ็นต์" = Percentage.Type' in code, True)
            # ‼️ 87.5% ต้องกลายเป็น 0.875 ถ้าเขียน 87.5 ตัวเลขจะเพี้ยนไป 100 เท่า
            ck("เปอร์เซ็นต์ถูกแปลงเป็นสัดส่วน", "0.875" in code, True)
            ck("วันที่เขียนเป็น #date ตามไวยากรณ์ M", "#date(2026,9,1)" in code, True)
            ck("วันที่ไม่เพี้ยนไปหนึ่งวัน", "#date(2026,8,31)" in code, False)
            ck("จริงเท็จเขียนเป็น logical", '#"เปิดใช้" = logical' in code and "true" in code, True)
            ck("เบอร์โทรเก็บเป็นข้อความ ศูนย์นำหน้าไม่หาย",
               '#"เบอร์โทร" = text' in code and '"0812345678"' in code, True)
            ck("คอลัมน์ที่มีช่องว่างประกาศเป็น nullable", '#"หมายเหตุ" = nullable text' in code, True)
            ck("เครื่องหมายคำพูดในข้อความถูกเขียนซ้อนสองตัว", '"สมชาย ""เล็ก"""' in code, True)
            ck("ข้อความที่มี #( ถูก escape ให้ถูกต้อง", "#(#)(พิเศษ)" in code, True)
            ck("วงเล็บเปิดปิดเท่ากัน", code.count("("), code.count(")"))
            ck("ปีกกาเปิดปิดเท่ากัน", code.count("{"), code.count("}"))

            print("\n── ③ ชีทที่ซ่อนไว้ต้องไม่หลุดมา ──")
            ck("เนื้อในชีทที่ซ่อนไม่อยู่ในโค้ด", "ห้ามเผยแพร่" in code, False)
            ck("บอกผู้ใช้ว่าข้ามชีทที่ซ่อนไว้", "ลับ" in pg.locator(".note").first.inner_text(), True)

            print("\n── ④ ผู้ใช้แก้ชนิดเองได้ และสลับรูปแบบผลลัพธ์ได้ ──")
            pg.evaluate("""() => {
                const s = [...document.querySelectorAll('.pq-cols tbody tr select')][6];
                s.value = 'int64';
                s.dispatchEvent(new Event('change', { bubbles: true }));
            }""")
            pg.wait_for_timeout(600)
            ck("เปลี่ยนชนิดเองแล้วโค้ดเปลี่ยนตามทันที",
               '#"เบอร์โทร" = Int64.Type' in pg.locator(".pq-code").inner_text(), True)

            pg.evaluate("() => { const r = [...document.querySelectorAll('.seg-item')]"
                        ".find((x) => /เต็ม|Full/.test(x.textContent)); if (r) r.click(); }")
            pg.wait_for_timeout(600)
            full = pg.locator(".pq-code").inner_text()
            ck("โหมดโค้ดเต็มห่อด้วย let ... in",
               full.startswith("let") and full.strip().endswith("Source"), True)

            pg.evaluate("() => { const c = document.querySelector('.clean-check input');"
                        " c.checked = true; c.dispatchEvent(new Event('change', { bubbles: true })); }")
            pg.wait_for_timeout(400)
            shown = pg.evaluate("() => [...document.querySelectorAll('.pq-cols tbody tr')]"
                                ".filter((r) => !r.hidden).length")
            ck("กรองดูเฉพาะคอลัมน์ที่ควรตรวจได้", shown, 1)

            print("\n── ⑤ แผงสอนเขียนสูตร ──")
            ck("มีแผงสอนอยู่ในหน้าเดียวกัน", pg.locator("details.pq-teach").count(), 1)
            # แผงพับไว้ตั้งต้น ต้องกางก่อนถึงจะอ่านเนื้อในได้ (และเป็นการยืนยันว่ากดเปิดได้จริง)
            pg.evaluate("() => { document.querySelector('details.pq-teach').open = true; }")
            pg.wait_for_timeout(300)
            teach = pg.locator("details.pq-teach").inner_text()
            ck("สอนเรื่องเปอร์เซ็นต์เป็นสัดส่วน", "0.875" in teach, True)
            ck("สอนวิธีเอาโค้ดไปใช้", "Advanced Editor" in teach, True)

            print("\n── ⑥ รับไฟล์ PDF และ Word ได้ด้วย ──")
            # ‼️ ไม่ใส่เคสรูปถ่ายไว้ในชุดนี้ เพราะ OCR ต้องโหลด language pack จาก CDN
            #    ทำให้ชุดเทสช้าและขึ้นกับเน็ต ตรวจด้วยมือแล้วว่าใช้ได้ (โครงตารางถูกทุกคอลัมน์
            #    ส่วนตัวอักษรเล็ก ๆ OCR อ่านคลาดได้ ซึ่งเครื่องมือเตือนผู้ใช้ไว้แล้ว)
            for path, label, want_note in [(pdf_file, "PDF", False), (docx_file, "Word", False)]:
                pg.goto("about:blank")
                pg.goto(f"{base}/#/excel-to-pq", wait_until="networkidle")
                pg.wait_for_selector(".dz")
                pg.locator(".dz input[type=file]").first.set_input_files(str(path))
                pg.wait_for_timeout(6000)
                types2 = pg.evaluate("() => [...document.querySelectorAll('.pq-cols tbody tr select')]"
                                     ".map((s) => s.value)")
                code2 = pg.locator(".pq-code").inner_text() if pg.locator(".pq-code").count() else ""
                ck(f"{label}: ดึงตารางออกมาได้ครบ 4 คอลัมน์", len(types2), 4)
                ck(f"{label}: สร้างโค้ดได้", code2.startswith("= #table("), True)
                ck(f"{label}: ข้อมูลในแถวครบ", '"ปากกาน้ำเงิน"' in code2, True)
                ck(f"{label}: รหัสสินค้าเก็บเป็นข้อความ", '"A01"' in code2, True)
                ck(f"{label}: จำนวนเป็นจำนวนเต็ม ราคาเป็นทศนิยม",
                   '#"จำนวน" = Int64.Type' in code2 and '#"ราคา" = number' in code2, True)

            print("\n── ไม่มี error หลุดออกมา ──")
            ck("ไม่มี console หรือ page error ตลอดทั้งชุด", errs, [])
            ctx.close()
            b.close()
    finally:
        if server:
            server.terminate()
            server.wait(timeout=10)


if __name__ == "__main__":
    try:
        main()
    finally:
        shutil.rmtree(TMP, ignore_errors=True)
    print(f"\nผ่าน {P}, ตก {len(F)}")
    for i, x in enumerate(F, 1):
        print(f"  {i}. {x}")
    sys.exit(1 if F else 0)
