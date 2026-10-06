# เครื่องมือ "ดึงข้อมูลจากอีกไฟล์มาเติม" (excel-lookup) ตรวจของที่ผู้ใช้เอาไปใช้ต่อจริง
#
# ‼️ คำตอบที่ถูกไม่ได้เดาเอง: คำนวณเองอีกทางด้วย openpyxl จากไฟล์ทดสอบชุดเดียวกัน (ไม่ผ่านโค้ดของเครื่องมือ)
#    แล้วเทียบกับที่หน้าเว็บโชว์ และกับไฟล์ที่หน้าเว็บสร้างจริง
# ‼️ ไฟล์หลัก tests/fixtures/lookup-main.xlsx ถูก Excel จริงเขียน (มีสูตร shared, ฟิลเตอร์, แถวซ่อน, แช่แข็งหัวตาราง)
#    ข้อที่ต้องไม่กลับมาพัง: สูตร Mapping ต้องยังเป็นสูตร ฟิลเตอร์ต้องยังอยู่ ช่องสูตรที่ดูว่างต้องไม่ถูกทับ
# ‼️ การเขียนกลับทับไฟล์เดิมทดสอบด้วยมือจับไฟล์จำลองบน OPFS (headless เปิดหน้าต่างเลือกไฟล์จริงไม่ได้)
#    ส่วนที่ Excel ล็อกไฟล์พิสูจน์แยกกับ Excel จริงไว้แล้ว (ดู src/fshandle.js) ที่นี่จำลองด้วยมือจับที่ close() ล้ม
#
# ถ้าผิดจะรู้ได้ยังไง: ทุกข้อเทียบกับค่าที่คำนวณอิสระ ข้อ ⑥ ⑦ ⑧ มีเคสของเสียที่ต้องแดงถ้ากันไม่ได้
#
# รัน: tests/run.sh browser_lookup
import datetime
import hashlib
import os
import pathlib
import sys
import tempfile

import openpyxl
from playwright.sync_api import sync_playwright

sys.path.insert(0, str(pathlib.Path(__file__).resolve().parent))
import fkui  # noqa: E402

ROOT = pathlib.Path(__file__).resolve().parent.parent
FIX = ROOT / "tests" / "fixtures"
MAIN = FIX / "lookup-main.xlsx"
SRC = FIX / "lookup-source.xlsx"
TMP = pathlib.Path(tempfile.mkdtemp(prefix="filekit_lookup_"))
P, F = 0, []


def ck(name, got, want):
    global P
    ok = got == want
    if ok:
        P += 1
    else:
        F.append(f"{name}\n      ได้    : {got!r}\n      ควรได้ : {want!r}")
    print(f"  {'✅' if ok else '❌'} {name}")


# ── คำตอบที่ถูก คำนวณอิสระจากไฟล์ ─────────────────────────────────────────────
def expected():
    wb = openpyxl.load_workbook(MAIN, data_only=True)
    ws = wb["2026"]
    main = []   # (excel_row, key)
    for r in range(3, ws.max_row + 1):
        a, f = ws.cell(r, 1).value, ws.cell(r, 6).value
        if a is None and f is None:
            continue
        main.append((r, f"{a}{f}"))
    sw = openpyxl.load_workbook(SRC, data_only=True)["TaxTeam"]
    src = {}
    for r in range(2, sw.max_row + 1):
        k = sw.cell(r, 1).value
        if k is None or str(k).startswith("#"):
            continue
        k = str(int(k)) if isinstance(k, (int, float)) else str(k).strip()
        src.setdefault(k, []).append((r, sw.cell(r, 2).value, sw.cell(r, 3).value, sw.cell(r, 4).value))
    per = {}
    for r, k in main:
        hits = src.get(k, [])
        same = len({h[1:] for h in hits}) <= 1
        per[r] = ("none" if not hits else "one" if len(hits) == 1 else "dupSame" if same else "dupDiff", hits)
    return main, src, per


MAIN_ROWS, SRC_IDX, PER = expected()


def sec_only(path):
    """แถวของไฟล์รองที่คีย์ไม่อยู่ในไฟล์หลัก (src_row, คีย์, วันที่, เลขใบกำกับ, เลขส่งทีมภาษี) คำนวณอิสระ"""
    keys = {k for _, k in MAIN_ROWS}
    sw = openpyxl.load_workbook(path, data_only=True).worksheets[0]
    out = []
    for r in range(2, sw.max_row + 1):
        k = sw.cell(r, 1).value
        if k is None or str(k).startswith("#"):
            continue
        k = str(int(k)) if isinstance(k, (int, float)) else str(k).strip()
        if k not in keys:
            out.append((r, k, sw.cell(r, 2).value, sw.cell(r, 3).value, sw.cell(r, 4).value))
    return out


SEC_ONLY = sec_only(SRC)
# ไฟล์รองที่มีแถวใหม่เพิ่ม 2 แถว (คีย์เป็นเลขหนึ่งแถว เป็นข้อความไม่มีวันที่อีกแถว) ไว้ทดสอบ Addon หลายแถว
_w = openpyxl.load_workbook(SRC)
_s = _w.worksheets[0]
_s.append([10682619999901, datetime.datetime(2026, 6, 1), "TX-ADD1", "TT-A1", None])
_s.append(["10682619999902", None, "TX-ADD2", None, None])
SRC_ADD = TMP / "lookup-source-addon.xlsx"
_w.save(SRC_ADD)
SEC_ONLY_ADD = sec_only(SRC_ADD)
CNT = {s: sum(1 for v in PER.values() if v[0] == s) for s in ("none", "one", "dupSame", "dupDiff")}


def chips(pg):
    return pg.evaluate("() => [...document.querySelectorAll('.stats .stat')].map(c => c.textContent)")


def set_files(pg, main=MAIN, src=SRC):
    pg.locator(".dz input[type=file]").nth(0).set_input_files(str(main))
    pg.wait_for_timeout(1500)
    pg.locator(".dz input[type=file]").nth(1).set_input_files(str(src))
    pg.wait_for_timeout(1800)


def click_btn(pg, text):
    loc = pg.locator("button:visible").filter(has_text=text)
    loc.first.click()


def download(pg, text):
    with pg.expect_download(timeout=60000) as d:
        click_btn(pg, text)
    path = TMP / d.value.suggested_filename
    d.value.save_as(str(path))
    return path


def main():
    base = os.environ.get("FK_BASE") or "http://127.0.0.1:8899"
    errs = []
    with sync_playwright() as pw:
        b = pw.chromium.launch()
        ctx = b.new_context(viewport={"width": 1360, "height": 1000}, accept_downloads=True, timezone_id="Asia/Bangkok")
        pg = ctx.new_page()
        pg.on("pageerror", lambda e: errs.append(str(e)[:200]))
        pg.on("console", lambda m: errs.append(m.text[:200]) if m.type == "error" else None)
        fkui.open_tool(pg, "excel-lookup", base)
        set_files(pg)

        print("\n── ① เดาค่าเริ่มต้นได้ตรงกับไฟล์จริง ──")
        head_vals = pg.evaluate("() => [...document.querySelectorAll('input[type=number]')].map(i => i.value)")
        ck("หัวตารางไฟล์หลักอยู่แถว 2 (แถว 1 เป็นยอดรวม) และไฟล์รองแถว 1", head_vals, ["2", "1"])
        keys = pg.evaluate("""() => [...document.querySelectorAll('label.field select')]
            .map(s => s.options[s.selectedIndex] && s.options[s.selectedIndex].textContent)
            .filter(t => t && /Mapping/.test(t))""")
        ck("คีย์เดาเป็น Mapping ทั้งสองไฟล์ (G ฝั่งหลัก A ฝั่งรอง)", keys, ["G: Mapping", "A: Mapping"])
        pulled = pg.evaluate("() => [...document.querySelectorAll('.lk-pr.on .nm')].map(n => n.textContent)")
        ck("ติ๊กให้เฉพาะ 3 คอลัมน์ที่ไฟล์หลักว่างรอเติม (ไม่ติ๊ก Remark)", pulled,
           ["B: Tax Inv. Date (SM)", "C: Tax Inv. No (SM)", "D: Send to Tax Team No."])

        print("\n── ② ตัวเลขสรุปตรงกับที่คำนวณอิสระ ──")
        ch = chips(pg)
        found = CNT["one"] + CNT["dupSame"] + CNT["dupDiff"]
        ck("จำนวนแถวไฟล์หลัก", ch[0], f"ไฟล์หลัก {len(MAIN_ROWS)} แถว")
        ck("เจอ", ch[1], f"เจอ {found} แถว")
        ck("ไม่เจอ", next((c for c in ch if c.startswith("ไม่เจอ")), None), f"ไม่เจอ {CNT['none']} แถว")
        ck("เจอซ้ำค่าเหมือนกัน", next((c for c in ch if "ค่าเหมือนกัน" in c), None), f"เจอซ้ำค่าเหมือนกัน {CNT['dupSame']} แถว")
        ck("เจอซ้ำค่าต่างกัน", next((c for c in ch if "ค่าต่างกัน" in c), None), f"เจอซ้ำค่าต่างกัน {CNT['dupDiff']} แถว")
        banner = pg.evaluate("() => document.querySelector('.lk-banner')?.textContent || ''")
        dup_keys = sum(1 for v in SRC_IDX.values() if len(v) > 1)
        ck("ไฟล์รองมีคีย์ซ้ำกี่ค่า", f"คีย์ซ้ำ {dup_keys} ค่า" in banner.replace("ไฟล์รองมี", ""), True)
        ck("บอกจำนวนแถวไฟล์รองที่คีย์ #N/A (ข้าม 2 แถว)", "2 แถว" in banner and "ข้าม" in banner, True)

        print("\n── ③ ไฟล์ที่ดาวน์โหลด: สิ่งที่ต้องอยู่ครบและสิ่งที่ต้องถูกเติม ──")
        path = download(pg, "ดาวน์โหลดไฟล์ที่เติมแล้ว")
        wf = openpyxl.load_workbook(path)["2026"]
        wv = openpyxl.load_workbook(path, data_only=True)["2026"]
        ck("คอลัมน์ Mapping ยังเป็นสูตร (ไม่กลายเป็นตัวเลขตาย)", [str(wf["G3"].value), str(wf["G4"].value), str(wf["G20"].value)],
           ["=A3&F3", "=A4&F4", "=A20&F20"])
        ck("ยอดรวมแถว 1 ยังเป็นสูตร", str(wf["K1"].value), "=SUBTOTAL(9,K3:K20)")
        ck("ฟิลเตอร์ยังอยู่ และครอบคอลัมน์ผลที่ต่อท้าย (R ถึง U) ให้กรองได้", wf.auto_filter.ref, "A2:U20")
        ck("แช่แข็งหัวตารางยังอยู่", wf.freeze_panes, "C3")
        ck("แถวที่ซ่อนตามฟิลเตอร์ยังซ่อน", [wf.row_dimensions[r].hidden for r in (4, 6, 9, 13)], [True] * 4)
        ck("ช่องสูตร ='' (P5) ไม่ถูกทับ", str(wf["P5"].value), '=""')
        want_fill = {}
        for r, (st, hits) in PER.items():
            if hits:
                want_fill[r] = hits[0][1:]
        bad = []
        for r, (d, no, send) in want_fill.items():
            if r == 5:
                continue     # P5 เป็นสูตร ต้องไม่ถูกทับ (เช็คแยกข้างบน)
            got = (wv.cell(r, 14).value, wv.cell(r, 15).value, wv.cell(r, 16).value)
            exp_no = "OLD-001" if r == 12 else no       # แถว 12 มีของเดิมอยู่ เติมเฉพาะช่องว่างจึงไม่ทับ
            if got[1] != exp_no or got[2] != send or (got[0] and d and got[0].date() != d.date()) or (bool(got[0]) != bool(d)):
                bad.append((r, got, (d, exp_no, send)))
        ck("ทุกแถวที่เจอ ถูกเติมค่าแถวแรกของไฟล์รอง (ทั้งวันที่ เลขใบกำกับ เลขส่งทีมภาษี)", bad, [])
        ck("แถว 12 คงค่าเดิม OLD-001 (โหมดเติมเฉพาะช่องว่าง)", wv["O12"].value, "OLD-001")
        ck("ช่องที่ไม่เจอเว้นว่าง (ยกเว้นแถว 12 ที่มีของเดิม)", [wv.cell(r, 15).value for r, v in PER.items() if v[0] == "none"],
           ["OLD-001" if r == 12 else None for r, v in PER.items() if v[0] == "none"])
        head = [wv.cell(2, c).value for c in (18, 19, 20, 21)]
        ck("มีคอลัมน์ผลการหาท้ายตาราง 4 คอลัมน์", head, ["ผลการหา", "เจอในไฟล์รองกี่แถว", "เติมแล้ว", "ค่ามาจากแถวในไฟล์รอง"])
        # ‼️ 06/10/2026 พี่ปอนด์ขอ Yes/No ว่าแถวไหนถูกเติม และค่ามาจากแถวไหนของไฟล์รอง (โหมดเอาแถวแรก = แถวแรกที่เจอ)
        ck("เติมแล้ว Yes ทุกแถวที่เจอ No ทุกแถวที่ไม่เจอ", [wv.cell(r, 20).value for r in sorted(PER)],
           ["No" if PER[r][0] == "none" else "Yes" for r in sorted(PER)])
        ck("ค่ามาจากแถวในไฟล์รอง = แถวแรกที่เจอ (ไม่เจอเว้นว่าง)", [wv.cell(r, 21).value for r in sorted(PER)],
           [str(PER[r][1][0][0]) if PER[r][1] else None for r in sorted(PER)])
        ck("จำนวนที่เจอตรงกับไฟล์รองทุกแถว", [wv.cell(r, 19).value for r in sorted(PER)], [len(PER[r][1]) for r in sorted(PER)])
        ck("ข้อความผลของแถวที่ซ้ำค่าต่างกัน", [wv.cell(r, 18).value for r, v in PER.items() if v[0] == "dupDiff"],
           [f"เจอ {len(v[1])} แถว (ค่าต่างกัน)" for r, v in PER.items() if v[0] == "dupDiff"])

        print("\n── ④ เปลี่ยนตัวเลือกแล้วผลต้องเปลี่ยนตาม ──")
        pg.evaluate("""() => { const s = [...document.querySelectorAll('select')].find(x => [...x.options].some(o => o.value === 'blank'));
          s.value = 'blank'; s.dispatchEvent(new Event('change', { bubbles: true })); }""")
        pg.wait_for_timeout(600)
        path2 = download(pg, "ดาวน์โหลดไฟล์ที่เติมแล้ว")
        w2 = openpyxl.load_workbook(path2, data_only=True)["2026"]
        dd = [r for r, v in PER.items() if v[0] == "dupDiff"]
        ck("โหมด ‘ไม่เติมถ้าแถวซ้ำมีค่าต่างกัน’ แถวซ้ำค่าต่างกันเว้นว่าง", [w2.cell(r, 15).value for r in dd], [None] * len(dd))
        ck("โหมดไม่เติม: แถวซ้ำค่าต่างกันได้ เติมแล้ว No และไม่มีเลขแถวไฟล์รอง", [(w2.cell(r, 20).value, w2.cell(r, 21).value) for r in dd], [("No", None)] * len(dd))
        ds = [r for r, v in PER.items() if v[0] == "dupSame"]
        ck("แต่แถวซ้ำที่ค่าเหมือนกันยังเติม (ไม่กำกวม)", [w2.cell(r, 15).value for r in ds], [PER[r][1][0][2] for r in ds])
        ck("ข้อความผลบอกว่าไม่ได้เติม", w2.cell(dd[0], 18).value, f"เจอ {len(PER[dd[0]][1])} แถว (ค่าต่างกัน ไม่ได้เติม)")
        pg.evaluate("""() => { const s = [...document.querySelectorAll('select')].find(x => [...x.options].some(o => o.value === 'blank'));
          s.value = 'first'; s.dispatchEvent(new Event('change', { bubbles: true })); }""")
        pg.locator(".seg-item").filter(has_text="ทับของเดิมด้วย").click()
        pg.wait_for_timeout(600)
        path3 = download(pg, "ดาวน์โหลดไฟล์ที่เติมแล้ว")
        w3 = openpyxl.load_workbook(path3, data_only=True)["2026"]
        r12 = PER[12]
        ck("โหมด ‘ทับของเดิมด้วย’ แถว 12 ถูกทับด้วยค่าจากไฟล์รอง", w3["O12"].value, r12[1][0][2] if r12[1] else None)
        pg.locator(".seg-item").filter(has_text="เติมเฉพาะช่องว่าง").click()
        pg.wait_for_timeout(500)

        print("\n── ⑤ แท็บผลแยกกลุ่มพร้อมจำนวน แท็บแรกเป็นแถวที่เจอ เลขแถวตรง Excel ──")
        # ‼️ 06/10/2026 พี่ปอนด์: ตัวอย่างผลควรโชว์ที่เจอ (ไฟล์จริงเจอ 7 จาก 23,205 เดิมเห็นแต่ “ไม่เจอ”) และ “ต้องตรวจ” ไม่สื่อ
        tabs = pg.evaluate("() => [...document.querySelectorAll('.lk-tabs .seg-item')].map(x => x.textContent)")
        n_multi = CNT["dupSame"] + CNT["dupDiff"]
        ck("แท็บพร้อมจำนวน ไม่มีคำว่าต้องตรวจ", tabs,
           [f"เจอ {found}", f"ไม่เจอ {CNT['none']}", f"เจอหลายแถว {n_multi}", f"มีแต่ในไฟล์รอง {len(SEC_ONLY)}", f"คีย์ซ้ำในไฟล์รอง {dup_keys}"])
        table_rows = lambda: pg.evaluate("""() => [...document.querySelectorAll('.xt-wrap table tbody tr')].map(tr => [...tr.children].map(td => td.textContent))""")
        pg.locator(".lk-tabs .seg-item").filter(has_text="เจอ ").first.click()
        pg.wait_for_timeout(400)
        ck("แท็บเจอ (ค่าเริ่มต้น) มีเฉพาะแถวที่เจอ ครบทุกแถว", sorted(int(r[0]) for r in table_rows()), sorted(r for r, v in PER.items() if v[0] != "none"))
        pg.locator(".lk-tabs .seg-item").filter(has_text="ไม่เจอ").click()
        pg.wait_for_timeout(400)
        ck("แท็บไม่เจอ ครบทุกแถว", sorted(int(r[0]) for r in table_rows()), sorted(r for r, v in PER.items() if v[0] == "none"))
        pg.locator(".lk-tabs .seg-item").filter(has_text="เจอหลายแถว").click()
        pg.wait_for_timeout(400)
        rows = table_rows()
        ck("แท็บเจอหลายแถว ครบทุกแถว", sorted(int(r[0]) for r in rows), sorted(r for r, v in PER.items() if v[0] in ("dupSame", "dupDiff")))
        dr = next(r for r in rows if int(r[0]) in dd)
        ck("แถวซ้ำบอกเลขแถวของไฟล์รองที่เจอ", dr[3], ", ".join(str(h[0]) for h in PER[int(dr[0])][1]))
        pg.locator(".lk-tabs .seg-item").filter(has_text="มีแต่ในไฟล์รอง").click()
        pg.wait_for_timeout(400)
        ck("แท็บมีแต่ในไฟล์รอง บอกเลขแถวของไฟล์รอง", [int(r[0]) for r in table_rows()], [s[0] for s in SEC_ONLY])
        want_rows = sorted(r for r, v in PER.items() if v[0] != "one")
        rep = download(pg, "ดาวน์โหลดรายงาน")
        rw = openpyxl.load_workbook(rep)
        ck("รายงานมีชีตสรุป ไม่เจอและเจอหลายแถว คีย์ซ้ำ มีแต่ในไฟล์รอง", rw.sheetnames, ["สรุป", "ไม่เจอและเจอหลายแถว", "คีย์ซ้ำในไฟล์รอง", "มีแต่ในไฟล์รอง"])
        ck("ชีตไม่เจอและเจอหลายแถว มีครบทุกรายการ", rw["ไม่เจอและเจอหลายแถว"].max_row - 1, len(want_rows))

        print("\n── ⑪ เพิ่มแถวที่ไฟล์หลักไม่มีเป็นแถวใหม่ (Addon) ──")
        pg.locator(".dz input[type=file]").nth(1).set_input_files(str(SRC_ADD))
        pg.wait_for_timeout(1800)
        ck("ก่อนเปิด: บอกว่ามีแต่ในไฟล์รองกี่แถว ยังไม่เพิ่ม", next((c for c in chips(pg) if "ไฟล์รอง" in c and "แถว" in c), None), f"มีแต่ในไฟล์รอง {len(SEC_ONLY_ADD)} แถว")
        pg.locator(".lk-switch").filter(has_text="Addon").locator("input").check()
        pg.locator(".lk-addon button").filter(has_text="วันนี้").click()
        pg.wait_for_timeout(700)
        ck("เปิดแล้ว: จะเพิ่มแถวใหม่ตามจำนวนที่คำนวณอิสระ", next((c for c in chips(pg) if "จะเพิ่มแถวใหม่" in c), None), f"จะเพิ่มแถวใหม่ {len(SEC_ONLY_ADD)} แถว")
        tag = "Addon " + datetime.datetime.now(datetime.timezone(datetime.timedelta(hours=7))).strftime("%d/%m/%Y")
        ck("ปุ่มวันนี้ใส่ป้ายวันที่วันนี้", pg.locator(".lk-addon input[type=text]").input_value(), tag)
        pa = download(pg, "ดาวน์โหลดไฟล์ที่เติมแล้ว")
        af = openpyxl.load_workbook(pa)["2026"]
        av = openpyxl.load_workbook(pa, data_only=True)["2026"]
        last_main = max(PER)
        head2 = [av.cell(2, c).value for c in range(1, af.max_column + 1)]
        tcol = head2.index("Addon") + 1 if "Addon" in head2 else None
        scol = head2.index("ผลการหา") + 1 if "ผลการหา" in head2 else None
        ck("มีคอลัมน์ป้าย Addon และคอลัมน์ผล", [tcol is not None, scol is not None], [True, True])
        got = []
        for i, (sr, key, d, no, send) in enumerate(SEC_ONLY_ADD):
            r = last_main + 1 + i
            dv = av.cell(r, 14).value
            got.append((str(av.cell(r, 7).value).strip(), dv.date() if dv else None, av.cell(r, 15).value, av.cell(r, 16).value,
                        av.cell(r, tcol).value if tcol else None, av.cell(r, scol).value if scol else None))
        want = [(key, d.date() if d else None, no, send, tag, f"เพิ่มจากไฟล์รอง แถว {sr}") for sr, key, d, no, send in SEC_ONLY_ADD]
        ck("แถวใหม่ต่อท้ายตาราง: คีย์ วันที่ เลขใบกำกับ เลขส่งทีมภาษี ป้าย ผล ตรงกับไฟล์รองทุกแถว", got, want)
        ck("ไม่มีแถวเกินมา", av.max_row, last_main + len(SEC_ONLY_ADD))
        ck("แถว Addon: เติมแล้ว Yes และค่ามาจากแถวของไฟล์รองที่ถูกต้อง",
           [(av.cell(last_main + 1 + i, 20).value, av.cell(last_main + 1 + i, 21).value) for i in range(len(SEC_ONLY_ADD))],
           [("Yes", str(sr)) for sr, *_ in SEC_ONLY_ADD])
        ck("ฟิลเตอร์ขยายครอบแถวใหม่และคอลัมน์ผลกับป้าย Addon", af.auto_filter.ref, f"A2:V{last_main + len(SEC_ONLY_ADD)}")
        ck("ช่องวันที่ของแถวใหม่ใช้รูปแบบเดียวกับแถวบน", af.cell(last_main + 1, 14).number_format, af.cell(last_main, 14).number_format)
        ck("สูตร Mapping ของแถวเดิมยังเป็นสูตร", str(af["G3"].value), "=A3&F3")
        pg.locator(".lk-switch").filter(has_text="Addon").locator("input").uncheck()
        pg.locator(".dz input[type=file]").nth(1).set_input_files(str(SRC))
        pg.wait_for_timeout(1800)

        print("\n── ⑫ คีย์ต่อกันได้หลายคอลัมน์ (เดิมตายตัว 2) ──")
        add = pg.locator(".lk-keys").first.locator("button").filter(has_text="ต่ออีกคอลัมน์")
        add.click(); pg.wait_for_timeout(300)
        add = pg.locator(".lk-keys").first.locator("button").filter(has_text="ต่ออีกคอลัมน์")
        add.click(); pg.wait_for_timeout(300)
        ck("กดต่ออีกคอลัมน์สองครั้ง ได้ 3 ช่องคีย์", pg.locator(".lk-keys").first.locator("select").count(), 3)
        sels = pg.locator(".lk-keys").first.locator("select")
        sels.nth(0).select_option("0"); sels.nth(1).select_option("5")     # Company Code + Doc. No.
        pg.locator(".lk-keys").first.locator(".lk-key button").last.click()   # เอาช่องที่ 3 ออก
        pg.wait_for_timeout(700)
        ck("รหัสบริษัท+เลขเอกสาร (A+F) เจอเท่ากับใช้ Mapping", chips(pg)[1], f"เจอ {found} แถว")
        pg.locator(".lk-keys").first.locator(".lk-key button").last.click()
        pg.locator(".lk-keys").first.locator("select").first.select_option("5")   # Doc. No. อย่างเดียว เทียบกับ Mapping ไม่มีทางเจอ
        pg.wait_for_timeout(700)
        bn = pg.evaluate("() => document.querySelector('.lk-banner')?.textContent || ''")
        ck("เจอน้อยผิดปกติ: โชว์หน้าตาคีย์สองฝั่งให้เทียบเอง", ["เจอแค่ 0" in bn, "2619062802" in bn, "10682619062802" in bn], [True, True, True])
        pg.locator(".lk-keys").first.locator("select").first.select_option("6")
        pg.wait_for_timeout(700)
        ck("กลับมาใช้ Mapping แล้วไม่มีคำเตือนเจอน้อย", "เจอแค่" in pg.evaluate("() => document.querySelector('.lk-banner')?.textContent || ''"), False)

        print("\n── ⑥ เขียนกลับทับไฟล์เดิม (มือจับจำลองบน OPFS) ──")
        b.close()
        b = pw.chromium.launch()
        ctx = b.new_context(viewport={"width": 1360, "height": 1000}, accept_downloads=True, timezone_id="Asia/Bangkok")
        pg = ctx.new_page()
        pg.on("pageerror", lambda e: errs.append(str(e)[:200]))
        main_bytes = list(MAIN.read_bytes())
        pg.add_init_script("""
          window.__opfs = null;
          window.__failClose = false;
          window.__putMain = async (arr) => {
            const root = await navigator.storage.getDirectory();
            const fh = await root.getFileHandle('Pending VAT 2026.xlsx', { create: true });
            const w = await fh.createWritable(); await w.write(new Uint8Array(arr)); await w.close();
            window.__opfs = fh; return fh;
          };
          // มือจับที่ครอบของจริง: ทำให้ close() ล้มได้ เหมือน Excel ล็อกไฟล์ตอนสลับไฟล์พักทับ
          window.showOpenFilePicker = async () => {
            const real = window.__opfs;
            return [new Proxy(real, { get(t, k) {
              if (k === 'createWritable') return async () => {
                const w = await t.createWritable();
                if (!window.__failClose) return w;
                return { write: (d) => w.write(d), close: async () => { await w.abort(); const e = new Error('The requested file could not be read, typically due to permission problems'); e.name = 'NoModificationAllowedError'; throw e; }, abort: () => w.abort().catch(() => {}) };
              };
              if (k === 'queryPermission' || k === 'requestPermission') return async () => 'granted';
              const v = t[k]; return typeof v === 'function' ? v.bind(t) : v;
            } })];
          };
        """)
        fkui.open_tool(pg, "excel-lookup", base)
        pg.evaluate("(a) => window.__putMain(a)", main_bytes)
        sha0 = hashlib.sha256(bytes(main_bytes)).hexdigest()

        def opfs_bytes():
            return bytes(pg.evaluate("async () => { const f = await window.__opfs.getFile(); return [...new Uint8Array(await f.arrayBuffer())]; }"))

        # โหลดไฟล์รองก่อน แล้วเปิดไฟล์หลักผ่านปุ่มแก้ตรง
        pg.locator(".dz input[type=file]").nth(1).set_input_files(str(SRC))
        pg.wait_for_timeout(1200)
        # ‼️ 06/10/2026 เลือกไฟล์หลักรอบเดียว: กดกล่องไฟล์หลักตามปกติก็ได้มือจับ ไม่มีปุ่มแยกให้เลือกซ้ำแล้ว
        ck("ไม่มีปุ่มเลือกไฟล์หลักซ้ำอีกรอบ", pg.locator("button").filter(has_text="เปิดไฟล์หลักแบบแก้ตรง").count(), 0)
        pg.locator(".dz-wrap").first.locator(".dz").click()      # กล่องย่อเป็นแถบเดียวในสถานะทำงาน คนจริงกดแถบนี้
        pg.wait_for_timeout(2200)
        note = pg.evaluate("() => document.querySelector('.lk-inplace')?.textContent || ''")
        ck("รู้ว่าแก้ตรงในไฟล์ได้ และเตือนเรื่อง Excel ล็อกไฟล์", "แก้ตรงในไฟล์นี้ได้" in note and "Excel" in note, True)
        ck("ตัวเลขสรุปเหมือนตอนอัปโหลดปกติ", chips(pg)[1], f"เจอ {found} แถว")

        # ลองก่อนแบบ close ล้ม = Excel ล็อก
        pg.evaluate("() => { window.__failClose = true; }")
        click_btn(pg, "บันทึกลงไฟล์หลักเดิม")
        ck("กดครั้งแรกแค่ขอยืนยัน ยังไม่เขียน", hashlib.sha256(opfs_bytes()).hexdigest(), sha0)
        btn = pg.locator("button:visible").filter(has_text="กดอีกครั้งเพื่อเขียนทับ").first.text_content()
        ck("ปุ่มบอกชื่อไฟล์กับจำนวนช่องที่จะเติม", "Pending VAT 2026.xlsx" in btn and "ช่อง" in btn, True)
        pg.locator("button:visible").filter(has_text="กดอีกครั้งเพื่อเขียนทับ").first.click()
        pg.wait_for_timeout(2500)
        msg = pg.evaluate("() => document.querySelector('.status')?.textContent || ''")
        ck("ไฟล์ถูกล็อก: บอกให้ปิดไฟล์ใน Excel และบอกว่าไฟล์เดิมไม่เสีย", "Excel" in msg and "ไม่ได้เสียหาย" in msg, True)
        ck("ไฟล์เดิมไม่ถูกแตะเลยแม้เขียนล้ม (sha256 เท่าเดิม)", hashlib.sha256(opfs_bytes()).hexdigest(), sha0)

        # ปล่อยล็อก แล้วบันทึกจริง
        pg.evaluate("() => { window.__failClose = false; }")
        click_btn(pg, "บันทึกลงไฟล์หลักเดิม")
        pg.locator("button:visible").filter(has_text="กดอีกครั้งเพื่อเขียนทับ").first.click()
        pg.wait_for_timeout(3500)
        msg = pg.evaluate("() => document.querySelector('.status')?.textContent || ''")
        ck("บันทึกสำเร็จและอ่านกลับจากไฟล์แล้วตรงทุกช่อง", "อ่านกลับจากไฟล์แล้วตรงทุกช่อง" in msg or msg, True)
        saved = TMP / "saved-in-place.xlsx"
        saved.write_bytes(opfs_bytes())
        sf = openpyxl.load_workbook(saved)["2026"]
        sv = openpyxl.load_workbook(saved, data_only=True)["2026"]
        ck("ไฟล์บนดิสก์: Mapping ยังเป็นสูตร ฟิลเตอร์ยังอยู่", [str(sf["G3"].value), sf.auto_filter.ref], ["=A3&F3", "A2:U20"])
        ck("ไฟล์บนดิสก์: ถูกเติมตรงกับที่คำนวณอิสระ", [sv.cell(r, 15).value for r in sorted(PER)],
           [("OLD-001" if r == 12 else (PER[r][1][0][2] if PER[r][1] else None)) for r in sorted(PER)])
        ck("ไฟล์บนดิสก์ต่างจากเดิม (เขียนจริง)", hashlib.sha256(opfs_bytes()).hexdigest() != sha0, True)
        # โหลดสถานะใหม่หลังบันทึก: ทำซ้ำต้องไม่ต่อคอลัมน์ซ้ำ
        ch2 = chips(pg)
        ck("หลังบันทึกโหลดไฟล์ใหม่ ช่องที่เติมแล้วนับเป็น ‘ไม่ทับของเดิม’", any("ไม่ทับของเดิม" in c for c in ch2), True)

        print("\n── ⑦ ย้อนกลับ ──")
        click_btn(pg, "ย้อนกลับ")
        pg.wait_for_timeout(2500)
        ck("ย้อนกลับแล้วไฟล์เหมือนเดิมทุกไบต์ (sha256)", hashlib.sha256(opfs_bytes()).hexdigest(), sha0)

        print("\n── ⑧ กันเขียนทับงานที่ใครแก้ไปหลังเปิดเข้ามา ──")
        click_btn(pg, "บันทึกลงไฟล์หลักเดิม")
        # มีคนแก้ไฟล์ระหว่างที่หน้านี้เปิดอยู่ (ขนาดไม่เท่าเดิม)
        pg.evaluate("async () => { const w = await window.__opfs.createWritable(); await w.write(new Uint8Array(await (await window.__opfs.getFile()).arrayBuffer())); await w.write(new Uint8Array([0])); await w.close(); }")
        sha_changed = hashlib.sha256(opfs_bytes()).hexdigest()
        pg.locator("button:visible").filter(has_text="กดอีกครั้งเพื่อเขียนทับ").first.click()
        pg.wait_for_timeout(2000)
        msg = pg.evaluate("() => document.querySelector('.status')?.textContent || ''")
        ck("ไฟล์ถูกเปลี่ยนหลังเปิด: ไม่เขียนทับและบอกเหตุผล", "ถูกเปลี่ยนหลังจากที่เปิดเข้ามา" in msg, True)
        ck("ไฟล์ที่คนอื่นแก้ไว้ไม่ถูกแตะ", hashlib.sha256(opfs_bytes()).hexdigest(), sha_changed)

        print("\n── ⑨ ลากไฟล์มาวาง: จับมือจับได้ตอน drop ──")
        pg.evaluate("(a) => window.__putMain(a)", main_bytes)
        fkui.open_tool(pg, "excel-lookup", base)
        pg.evaluate("(a) => window.__putMain(a)", main_bytes)
        pg.locator(".dz input[type=file]").nth(1).set_input_files(str(SRC))
        pg.wait_for_timeout(1200)
        pg.evaluate("""async () => {
          const fh = window.__opfs; const file = await fh.getFile();
          // ‼️ เบราว์เซอร์จริงให้ getAsFileSystemHandle() บน DataTransferItem จำลองสิ่งเดียวกัน
          DataTransferItem.prototype.getAsFileSystemHandle = function () { return Promise.resolve(fh); };
          const dt = new DataTransfer(); dt.items.add(file);
          const zone = document.querySelectorAll('.dz')[0];
          zone.dispatchEvent(new DragEvent('drop', { bubbles: true, cancelable: true, dataTransfer: dt }));
        }""")
        pg.wait_for_timeout(2500)
        note = pg.evaluate("() => document.querySelector('.lk-inplace')?.textContent || ''")
        ck("ไฟล์ที่ลากมาวางได้มือจับ พร้อมแก้ตรง", "แก้ตรงในไฟล์นี้ได้" in note, True)

        print("\n── ⑩ คีย์เป็นสูตรที่ไม่มีค่าเก็บไว้ ต้องบอกสาเหตุ ไม่ใช่ให้เดา ──")
        # ไฟล์ที่สร้างจากสคริปต์ไม่เคยถูก Excel บันทึก สูตร Mapping จึงไม่มีค่าที่คำนวณไว้ ทุกแถวกลายเป็น “คีย์ว่าง”
        nm = openpyxl.Workbook(); w = nm.active; w.title = "S"
        w.append(["Company Code", "Doc. No.", "Mapping", "Tax Inv. No (SM)"])
        for i in range(1, 6):
            w.append([1068, 2600 + i, f"=A{i+1}&B{i+1}", None])
        nm.save(TMP / "nocache-main.xlsx")
        ns = openpyxl.Workbook(); w = ns.active; w.title = "T"
        w.append(["Mapping", "Tax Inv. No (SM)"]); w.append(["10682601", "TX-1"])
        ns.save(TMP / "nocache-src.xlsx")
        ctx2 = b.new_context(viewport={"width": 1360, "height": 1000}, timezone_id="Asia/Bangkok")
        p2 = ctx2.new_page()
        fkui.open_tool(p2, "excel-lookup", base)
        p2.locator(".dz input[type=file]").nth(0).set_input_files(str(TMP / "nocache-main.xlsx"))
        p2.wait_for_timeout(1200)
        p2.locator(".dz input[type=file]").nth(1).set_input_files(str(TMP / "nocache-src.xlsx"))
        p2.wait_for_timeout(1500)
        bn = p2.evaluate("() => document.querySelector('.lk-banner')?.textContent || ''")
        ck("คีย์ว่างเกินครึ่ง: บอกให้เปิดใน Excel บันทึกก่อนถ้าเป็นสูตร", "ว่างเกินครึ่ง" in bn and "เปิดไฟล์ใน Excel แล้วกดบันทึก" in bn, True)
        ck("ไฟล์ปกติที่ Excel บันทึกแล้วไม่ขึ้นคำเตือนนี้", "ว่างเกินครึ่ง" in banner, False)

        print("\n── ⑬ ไฟล์รองคอลัมน์เยอะ: ค้นหาคอลัมน์ได้ และชื่อหัวไม่ตรงเป๊ะก็จับคู่ปลายทางให้ ──")
        # ‼️ 06/10/2026 พี่ปอนด์: ควรมีให้ค้นหาไหม และถ้าสองไฟล์เขียนชื่อต่างกันจะรู้ได้ไงว่าลงคอลัมน์ไหน
        wide = openpyxl.Workbook(); w = wide.active; w.title = "W"
        w.append(["Mapping", "tax inv date sm", "Tax.Inv.No (SM)"] + [f"Extra {i}" for i in range(1, 7)])
        w.append(["10682619062802", datetime.datetime(2026, 3, 27), "TX-W1"] + [i for i in range(1, 7)])
        wide.save(TMP / "wide-src.xlsx")
        p3 = ctx2.new_page()
        fkui.open_tool(p3, "excel-lookup", base)
        p3.locator(".dz input[type=file]").nth(0).set_input_files(str(MAIN))
        p3.wait_for_timeout(1500)
        p3.locator(".dz input[type=file]").nth(1).set_input_files(str(TMP / "wide-src.xlsx"))
        p3.wait_for_timeout(1500)
        rows3 = p3.evaluate("""() => [...document.querySelectorAll('.lk-pr')].map(r => ({ nm: r.querySelector('.nm').textContent,
            on: r.classList.contains('on'), dest: r.querySelector('select').value, hint: r.querySelector('.lk-hint').textContent }))""")
        ck("ชื่อหลวม: tax inv date sm ลง N และ Tax.Inv.No (SM) ลง O ติ๊กให้แล้ว",
           [(r["nm"], r["on"], r["dest"]) for r in rows3[:2]], [("B: tax inv date sm", True, "13"), ("C: Tax.Inv.No (SM)", True, "14")])
        ck("ชื่อหลวมบอกให้ตรวจอีกที", "ชื่อใกล้กัน" in rows3[0]["hint"], True)
        ck("คอลัมน์ที่ไม่มีชื่อตรง บอกว่าจะเป็นคอลัมน์ใหม่เมื่อติ๊ก", (rows3[2]["dest"], "ไม่มีชื่อตรง" in rows3[2]["hint"]), ("new", True))
        ck("คอลัมน์เกิน 6 มีช่องค้นหา", p3.locator(".lk-find").is_visible(), True)
        p3.locator(".lk-find").fill("extra 4")
        p3.wait_for_timeout(300)
        vis = p3.evaluate("() => [...document.querySelectorAll('.lk-pr')].filter(r => !r.hidden).map(r => r.querySelector('.nm').textContent)")
        ck("ค้นหา extra 4 เหลือแถวเดียว", vis, ["G: Extra 4"])
        p3.locator("button").filter(has_text="ติ๊กทั้งหมด").click()
        p3.wait_for_timeout(300)
        on = p3.evaluate("() => [...document.querySelectorAll('.lk-pr.on .nm')].map(n => n.textContent)")
        ck("ติ๊กทั้งหมดระหว่างค้นหา ติ๊กเฉพาะที่เห็น (ของเดิมยังติ๊ก)", on, ["B: tax inv date sm", "C: Tax.Inv.No (SM)", "G: Extra 4"])
        p3.locator(".lk-find").fill("ไม่มีแน่นอน")
        p3.wait_for_timeout(300)
        ck("ค้นไม่เจอบอกตรง ๆ", "ไม่มีคอลัมน์ที่ชื่อมี" in p3.evaluate("() => [...document.querySelectorAll('.lk-note')].filter(n => !n.hidden).map(n => n.textContent).join(' ')"), True)
        b.close()

    ck("ไม่มี error ใน console", errs, [])
    print(f"\n{'✅' if not F else '❌'} ผ่าน {P} ข้อ {'' if not F else f'ตก {len(F)} ข้อ'}")
    if F:
        print("\n" + "\n".join(F) + "\n")
        sys.exit(1)


if __name__ == "__main__":
    main()
