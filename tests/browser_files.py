import io
import os
import re
import sys
import shutil
import zipfile
import pathlib
import tempfile
import xml.dom.minidom as minidom

from playwright.sync_api import sync_playwright, TimeoutError as PWTimeout

import openpyxl
import fitz            # pymupdf — ตรวจ PDF จริง (จำนวนหน้า/ข้อความ/รูปฝัง)
import docx as pydocx   # python-docx — ตรวจ Word จริง
from PIL import Image

# ── ตั้งค่าพื้นฐาน — อ่าน BASE จาก env เสมอ ห้าม hardcode port/URL ──────────────
ROOT = pathlib.Path(__file__).resolve().parent.parent
BASE = os.environ.get("FK_BASE", "http://localhost:8899")
SAMPLES = ROOT / "samples"

# รายชื่อเครื่องมือจริงจาก registry.js — ใช้แค่เช็คว่า id ที่เทสอ้างถึงมีอยู่จริง
# (ห้าม hardcode จำนวนเครื่องมือ — ดึงจากไฟล์จริงเสมอ)
REG_TEXT = (ROOT / "src/registry.js").read_text(encoding="utf-8")
TOOL_IDS = set(re.findall(r'id:"([\w-]+)"', REG_TEXT))

SIF_TIMEOUT = 10_000   # เวลารอ input[type=file] ก่อนใส่ไฟล์ (สั้นพอให้เคสพังไม่ต้องรอ 30s เต็ม)

WORKDIR = pathlib.Path(tempfile.mkdtemp(prefix="filekit_e2e_"))
BROKEN_DIR = WORKDIR / "broken"
OUT_DIR = WORKDIR / "out"
BROKEN_DIR.mkdir(parents=True, exist_ok=True)
OUT_DIR.mkdir(parents=True, exist_ok=True)


def log(msg=""):
    print(msg, flush=True)


def goto(pg, tool_id):
    assert tool_id in TOOL_IDS, f"ไม่พบเครื่องมือ id={tool_id} ใน src/registry.js (เทสอ้าง id ผิด)"
    # ไป about:blank ก่อนเสมอ — เพราะเป็น SPA hash-router, goto ตรง ๆ ระหว่างสอง #/route
    # บางทีไม่ทำให้ mount ใหม่ครบ (ตามแบบแผนของ tests/browser_smoke.py)
    pg.goto("about:blank")
    pg.goto(f"{BASE}/#/{tool_id}", wait_until="networkidle")


def dl(pg, trigger, filename, timeout=20_000):
    with pg.expect_download(timeout=timeout) as di:
        trigger()
    d = di.value
    out = OUT_DIR / filename
    d.save_as(str(out))
    return out


NO_RAW_ERROR_MARKERS = [
    "undefined", "NaN", "[object Object]", "TypeError", "ReferenceError",
    "Uncaught", " at http", ".js:", "SyntaxError", "null is not",
]


def assert_readable_thai_error(text, ctx):
    assert text and text.strip(), f"{ctx}: ข้อความว่างเปล่า — ควรบอกผู้ใช้ว่าเกิดอะไรขึ้น"
    for bad in NO_RAW_ERROR_MARKERS:
        assert bad not in text, f"{ctx}: โชว์ error ดิบ ({bad!r}) ปนอยู่ในข้อความที่ผู้ใช้เห็น: {text}"
    assert re.search(r"[ก-๙]", text), f"{ctx}: ไม่มีข้อความไทยเลยในสิ่งที่ผู้ใช้เห็น: {text}"


# ─────────────────────────────────────────────────────────────────────────
# สร้างไฟล์เสีย 5 ชนิด — เก็บนอกโปรเจกต์ (temp dir) ห้ามใส่ลง samples/
# ─────────────────────────────────────────────────────────────────────────
def make_broken_files():
    files = {}

    p = BROKEN_DIR / "เสีย-pdf-ปลอม.pdf"
    p.write_bytes(b"this is not a real pdf file, just garbage bytes 123456789 !!! ")
    files["garbage_pdf"] = p

    p = BROKEN_DIR / "ว่างเปล่า-ศูนย์ไบต์.xlsx"
    p.write_bytes(b"")
    files["empty_xlsx"] = p

    p = BROKEN_DIR / "ไทยเพี้ยน-TIS620-กำหนดเอง.csv"
    txt = "ชื่อ,แผนก,ตำแหน่ง\nกิตติ ใจงาม,ฝ่ายขาย,พนักงาน\nมานี รักเรียน,บัญชี,หัวหน้า\n"
    p.write_bytes(txt.encode("tis-620"))
    files["tis620_csv"] = p

    p = BROKEN_DIR / "รูปพัง-header-เพี้ยน.jpg"
    p.write_bytes(bytes([0xFF, 0xD8, 0xFF]) + os.urandom(600))
    files["broken_jpg"] = p

    p = BROKEN_DIR / "เวิร์ดปลอม-zip-เปล่า.docx"
    with zipfile.ZipFile(p, "w") as z:
        z.writestr("hello.txt", "นี่ไม่ใช่ไฟล์ Word จริง — แค่ zip เปล่า ๆ ที่เปลี่ยนนามสกุลมา")
    files["fake_docx"] = p

    return files


# ─────────────────────────────────────────────────────────────────────────
# ① เคสฟีเจอร์จริง — ใส่ไฟล์ตัวอย่าง → กดปุ่ม → ดาวน์โหลด → ตรวจเนื้อในไฟล์จริง
#    (ไม่ใช่แค่ len(bytes) > 0 — เปิดด้วยไลบรารีจริงแล้วอ่านค่าจริง)
# ─────────────────────────────────────────────────────────────────────────
def case_excel_csv(pg):
    goto(pg, "excel-csv")
    pg.set_input_files("input[type=file]", str(SAMPLES / "ตัวอย่าง-ข้อมูลใบเสนอราคา.xlsx"), timeout=SIF_TIMEOUT)
    pg.wait_for_selector(".file-row", timeout=15_000)
    pg.get_by_role("button", name="แปลงไฟล์").click()
    pg.wait_for_selector(".results .result", timeout=20_000)
    out = dl(pg, lambda: pg.locator(".results .result button").first.click(), "excelcsv_out.csv")
    data = out.read_bytes()
    assert data[:3] == b"\xef\xbb\xbf", "CSV ที่ได้ต้องมี UTF-8 BOM (กันไทยเพี้ยนตอนเปิดด้วย Excel)"
    text = data.decode("utf-8-sig")
    header = text.splitlines()[0]
    for col in ["เลขที่", "ชื่อลูกค้า", "ชื่อสินค้า"]:
        assert col in header, f"หัวคอลัมน์ {col!r} หายไปหลังแปลง: {header}"
    assert len(text.splitlines()) >= 3, "แถวข้อมูลหายไปหลังแปลงเป็น CSV"


def case_pdf_merge(pg):
    files = [SAMPLES / "ตัวอย่าง-รายงานประจำเดือน.pdf", SAMPLES / "ตัวอย่าง-ใบปะหน้าเอกสาร.pdf"]
    expected_pages = sum(fitz.open(str(f)).page_count for f in files)
    goto(pg, "pdf-merge")
    pg.set_input_files("input[type=file]", [str(f) for f in files], timeout=SIF_TIMEOUT)
    pg.wait_for_selector(".file-row", timeout=15_000)
    assert pg.locator(".file-row").count() == 2, "ไฟล์ที่ใส่ไม่ครบ 2 แถวในรายการ"
    pg.get_by_role("button", name="รวมไฟล์").click()
    pg.wait_for_selector(".results .result", timeout=20_000)
    out = dl(pg, lambda: pg.locator(".results .result button").first.click(), "pdfmerge_out.pdf")
    doc = fitz.open(str(out))
    assert doc.page_count == expected_pages, f"หน้าหลังรวมไม่ครบ: ได้ {doc.page_count} ต้องการ {expected_pages}"
    assert doc[0].get_text().strip(), "หน้าแรกของไฟล์ที่รวมแล้วไม่มีข้อความเลย"


def case_image_resize(pg):
    src = SAMPLES / "ตัวอย่าง-รูปภาพ-1.jpg"
    goto(pg, "image-resize")
    pg.set_input_files("input[type=file]", str(src), timeout=SIF_TIMEOUT)
    pg.wait_for_selector(".file-row", timeout=15_000)
    pg.get_by_role("button", name="ย่อและบีบอัด").click()
    pg.wait_for_selector(".status-wrap .status.show", timeout=20_000)
    out = dl(pg, lambda: pg.get_by_role("button", name="ดาวน์โหลดรูป").click(), "imgresize_out.jpg")
    im = Image.open(out)
    im.load()
    assert max(im.size) <= 1600, f"ควรถูกจำกัดด้านยาวสุดที่ 1600px แต่ได้ขนาด {im.size}"
    assert im.size[0] > 0 and im.size[1] > 0, f"ขนาดรูปผิดปกติ: {im.size}"


def case_thai_encoding(pg):
    goto(pg, "thai-encoding")
    pg.set_input_files("input[type=file]", str(SAMPLES / "ตัวอย่าง-ไทยเพี้ยน-แบบ TIS620.csv"), timeout=SIF_TIMEOUT)
    pg.wait_for_selector(".enc-card", timeout=15_000)
    card_text = pg.locator(".enc-card").inner_text()
    assert ("TIS-620" in card_text) or ("Windows-874" in card_text), f"ไม่ตรวจพบว่าเป็น TIS-620: {card_text}"
    out = dl(pg, lambda: pg.get_by_role("button", name="บันทึกเป็น UTF-8").click(), "thaienc_out.csv")
    data = out.read_bytes()
    assert data[:3] == b"\xef\xbb\xbf", "ไฟล์ที่ซ่อมแล้วต้องมี UTF-8 BOM"
    text = data.decode("utf-8-sig")
    assert re.search(r"[ก-๙]", text), "ถอดรหัสแล้วไม่มีอักษรไทยที่อ่านออกเลย"
    assert "�" not in text, "ยังมีอักขระเพี้ยน (mojibake) หลงเหลืออยู่หลังซ่อม"


def case_thai_id(pg):
    goto(pg, "thai-id")
    pg.set_input_files("input[type=file]", str(SAMPLES / "ตัวอย่าง-วันที่ไทยและเลขบัตร.xlsx"), timeout=SIF_TIMEOUT)
    pg.wait_for_selector(".stats .stat", timeout=15_000)
    chips = pg.locator(".stats").inner_text()
    assert "ถูกต้อง" in chips, f"ไม่พบผลตรวจ 'ถูกต้อง' เลย: {chips}"
    out = dl(pg, lambda: pg.get_by_role("button", name="ดาวน์โหลดเป็น Excel").click(), "thaiid_out.xlsx")
    wb = openpyxl.load_workbook(out)
    ws = wb.active
    header = [c.value for c in ws[1]]
    col_idx = next((i for i, h in enumerate(header) if h and "ผลตรวจ" in str(h)), None)
    assert col_idx is not None, f"ไม่พบคอลัมน์ 'ผลตรวจ' ในไฟล์ผลลัพธ์: {header}"
    values = [row[col_idx] for row in ws.iter_rows(min_row=2, values_only=True)]
    assert any(v == "ถูกต้อง" for v in values), f"ไม่มีแถวไหนตรวจว่า 'ถูกต้อง' เลย: {values}"
    assert any(v and v != "ถูกต้อง" for v in values), \
        f"ไม่มีแถวไหนตรวจว่าไม่ถูกต้องเลย (ไฟล์ตัวอย่างมีเลขบัตรผิด/ไม่ครบปนอยู่): {values}"


def case_thai_date(pg):
    goto(pg, "thai-date")
    pg.set_input_files("input[type=file]", str(SAMPLES / "ตัวอย่าง-วันที่ไทยและเลขบัตร.xlsx"), timeout=SIF_TIMEOUT)
    pg.wait_for_selector(".stats .stat", timeout=15_000)
    chips = pg.locator(".stats").inner_text()
    assert "แปลงได้" in chips, f"ไม่พบผลแปลงวันที่เลย: {chips}"
    out = dl(pg, lambda: pg.get_by_role("button", name="ดาวน์โหลดเป็น Excel").click(), "thaidate_out.xlsx")
    wb = openpyxl.load_workbook(out)
    ws = wb.active
    header = [c.value for c in ws[1]]
    col_idx = next((i for i, h in enumerate(header) if h and "ค.ศ." in str(h)), None)
    assert col_idx is not None, f"ไม่พบคอลัมน์ผลลัพธ์ 'วันที่ (ค.ศ.)': {header}"
    vals = [row[col_idx] for row in ws.iter_rows(min_row=2, values_only=True)]
    good = [v for v in vals if isinstance(v, str) and re.match(r"^\d{4}-\d{2}-\d{2}$", v)]
    assert len(good) >= 3, f"แปลงวันที่สำเร็จน้อยผิดปกติ ({len(good)} แถว): {vals}"
    for v in good:
        year = int(v[:4])
        assert 2000 <= year <= 2100, f"ปีหลังแปลง พ.ศ.→ค.ศ. ดูไม่สมเหตุสมผล: {v}"


def case_word_clean(pg):
    goto(pg, "word-clean")
    pg.set_input_files("input[type=file]", str(SAMPLES / "ตัวอย่าง-ใบเสนอราคา.docx"), timeout=SIF_TIMEOUT)
    pg.wait_for_selector(".clean-card", timeout=15_000)
    report = pg.locator(".clean-card").inner_text()
    assert ("ผู้เขียน" in report) or ("รหัสรอบการบันทึก" in report), \
        f"ไม่พบร่องรอย metadata ที่ควรเจอในไฟล์ตัวอย่าง (คาดว่ามี author/rsid): {report}"
    pg.get_by_role("button", name="ล้าง").click()
    pg.wait_for_selector(".results .result", timeout=20_000)
    out = dl(pg, lambda: pg.locator(".results .result button").first.click(), "wordclean_out.docx")

    with zipfile.ZipFile(out) as z:
        names = z.namelist()
        assert "word/document.xml" in names, "docx ผลลัพธ์ไม่มี word/document.xml เลย — ไม่ใช่ .docx ที่ถูกต้อง"
        core_xml = z.read("docProps/core.xml")
        # ยืนยันว่าฟีเจอร์ "ล้างชื่อผู้เขียน" ทำงานจริง — เช็คด้วย regex ตรง ๆ กันไว้ก่อน
        # เผื่อ XML ผิดรูปแบบ (ดูบั๊กด้านล่าง) จะได้ไม่บังฟีเจอร์ที่ทำถูกอยู่
        assert re.search(rb"<dc:creator\s*/>|<dc:creator>\s*</dc:creator>", core_xml), \
            f"ชื่อผู้เขียนไม่ถูกล้างใน docProps/core.xml: {core_xml[:300]!r}"

        # ── บั๊กจริงที่เจอระหว่างเขียนเทสนี้ (ไม่ใช่เทสเขียนผิด) ──────────────────
        # src/docxclean.js:155-156 — ตอนติ๊ก "ล้างชื่อผู้เขียน บริษัท และเวลาที่ใช้ทำ"
        # (ติ๊กอยู่โดย default) โค้ดเขียน docProps/core.xml ใหม่โดยขาดช่องว่างระหว่าง
        # attribute: `<?xml version="1.0"encoding="UTF-8"standalone="yes"?>` และ
        # `xmlns:cp="..."xmlns:dc="..."` ซึ่งผิดสเปก XML (ต้องมี whitespace คั่น
        # attribute เสมอ) ทำให้ parser มาตรฐาน (lxml/python-docx และอาจรวมถึง Word
        # เองในบางเงื่อนไข) เปิดไฟล์ core.xml ไม่ผ่าน — assert นี้จะ "แดง" จนกว่าจะแก้
        try:
            minidom.parseString(core_xml)
        except Exception as e:
            raise AssertionError(
                "บั๊กจริงใน src/docxclean.js:155-156 — docProps/core.xml ที่เขียนใหม่ตอนล้าง "
                "metadata ขาดช่องว่างระหว่าง attribute ทำให้เป็น XML ผิดรูปแบบ "
                f"(parser มาตรฐานเปิดไม่ผ่าน): {e} · เนื้อจริง: {core_xml[:200]!r}"
            )

    # ยืนยันซ้ำด้วย python-docx ตัวจริง (ถ้าบั๊กด้านบนยังอยู่ บรรทัดนี้จะ raise เช่นกัน)
    doc = pydocx.Document(str(out))
    assert doc.paragraphs, "เปิดไฟล์ที่ล้างแล้วด้วย python-docx ไม่พบย่อหน้าเลย"


def case_pdf_to_text(pg):
    src = SAMPLES / "ตัวอย่าง-รายงานประจำเดือน.pdf"
    expected_pages = fitz.open(str(src)).page_count
    goto(pg, "pdf-to-text")
    pg.set_input_files("input[type=file]", str(src), timeout=SIF_TIMEOUT)
    pg.wait_for_selector(".file-row", timeout=15_000)
    pg.get_by_role("button", name="ดึงข้อความ").click()
    pg.wait_for_selector(".results .actions", timeout=20_000)
    status = pg.locator(".status-wrap .status").inner_text()
    assert str(expected_pages) in status, f"จำนวนหน้าที่ดึงข้อความไม่ตรง (คาด {expected_pages}): {status}"
    out = dl(pg, lambda: pg.get_by_role("button", name="ดาวน์โหลด .txt").click(), "pdftotext_out.txt")
    data = out.read_bytes()
    assert data[:3] == b"\xef\xbb\xbf", "ไฟล์ .txt ที่ได้ต้องมี UTF-8 BOM"
    text = data.decode("utf-8-sig")
    assert re.search(r"[ก-๙]", text), "ไม่มีข้อความไทยในไฟล์ข้อความที่ดึงออกมาเลย"
    assert len(text.strip()) > 20, "เนื้อหาที่ดึงออกมาสั้นผิดปกติ — น่าจะดึงไม่สำเร็จจริง"


def case_images_to_pdf(pg):
    files = [SAMPLES / "ตัวอย่าง-รูปภาพ-1.jpg", SAMPLES / "ตัวอย่าง-รูปภาพ-2.png"]
    goto(pg, "images-to-pdf")
    pg.set_input_files("input[type=file]", [str(f) for f in files], timeout=SIF_TIMEOUT)
    pg.wait_for_selector(".file-row", timeout=15_000)
    assert pg.locator(".file-row").count() == 2, "ไฟล์รูปที่ใส่ไม่ครบ 2 แถว"
    pg.get_by_role("button", name="สร้างไฟล์ PDF").click()
    pg.wait_for_selector(".results .result", timeout=20_000)
    out = dl(pg, lambda: pg.locator(".results .result button").first.click(), "imgtopdf_out.pdf")
    doc = fitz.open(str(out))
    assert doc.page_count == 2, f"ควรได้ PDF 2 หน้า (1 รูป = 1 หน้า) แต่ได้ {doc.page_count}"
    pix = doc[0].get_pixmap()
    assert pix.width > 10 and pix.height > 10, "หน้าแรกของ PDF ที่ได้แทบไม่มีเนื้อหา"


def case_powerpoint_to_word(pg):
    goto(pg, "powerpoint-to-word")
    pg.set_input_files("input[type=file]", str(SAMPLES / "ตัวอย่าง-นำเสนอบริษัท.pptx"), timeout=SIF_TIMEOUT)
    pg.wait_for_selector(".file-row", timeout=15_000)
    pg.get_by_role("button", name="แปลงเป็น Word").click()
    pg.wait_for_selector(".results button", timeout=20_000)
    out = dl(pg, lambda: pg.locator(".results button", has_text="ดาวน์โหลด").first.click(), "ppt2word_out.docx")
    doc = pydocx.Document(str(out))
    full_text = "\n".join(p.text for p in doc.paragraphs)
    assert re.search(r"[ก-๙]", full_text), "ไม่มีข้อความไทยเลยในเอกสาร Word ที่แปลงออกมา"
    assert len(doc.paragraphs) >= 3, f"เนื้อหาน้อยผิดปกติหลังแปลงจาก PowerPoint: {len(doc.paragraphs)} ย่อหน้า"


def case_pdf_to_images(pg):
    src = SAMPLES / "ตัวอย่าง-รายงานประจำเดือน.pdf"
    expected_pages = fitz.open(str(src)).page_count
    goto(pg, "pdf-to-images")
    pg.set_input_files("input[type=file]", str(src), timeout=SIF_TIMEOUT)
    pg.wait_for_selector(".file-row", timeout=15_000)
    pg.get_by_role("button", name="แปลงเป็นรูป").click()
    pg.wait_for_selector(".results button", timeout=25_000)
    out = dl(pg, lambda: pg.locator(".results button", has_text="ZIP").first.click(), "pdf2img_out.zip")
    with zipfile.ZipFile(out) as z:
        names = z.namelist()
        assert len(names) == expected_pages, f"จำนวนรูปไม่ตรงจำนวนหน้า PDF: ได้ {len(names)} ต้องการ {expected_pages}"
        for n in names:
            im = Image.open(io.BytesIO(z.read(n)))
            im.verify()


def case_pdf_watermark(pg):
    src = SAMPLES / "ตัวอย่าง-รายงานประจำเดือน.pdf"
    expected_pages = fitz.open(str(src)).page_count
    goto(pg, "pdf-watermark")
    pg.set_input_files("input[type=file]", str(src), timeout=SIF_TIMEOUT)
    pg.wait_for_selector(".file-row", timeout=15_000)
    pg.get_by_role("button", name="ใส่ลายน้ำ").click()
    pg.wait_for_selector(".result", timeout=20_000)
    out = dl(pg, lambda: pg.get_by_role("button", name="ดาวน์โหลด").click(), "wm_out.pdf")
    doc = fitz.open(str(out))
    assert doc.page_count == expected_pages, f"จำนวนหน้าหลังใส่ลายน้ำไม่ตรง: {doc.page_count} vs {expected_pages}"
    assert len(doc[0].get_images()) >= 1, "หน้าแรกไม่มีภาพลายน้ำฝังอยู่เลย (ลายน้ำวาดเป็น PNG ฝังทับ)"


def case_pdf_split(pg):
    src = SAMPLES / "ตัวอย่าง-รายงานประจำเดือน.pdf"
    expected_pages = fitz.open(str(src)).page_count
    goto(pg, "pdf-split")
    pg.set_input_files("input[type=file]", str(src), timeout=SIF_TIMEOUT)
    pg.wait_for_selector(".file-row", timeout=15_000)
    pg.locator('input[type=radio][value="every"]').check()
    pg.locator("input[type=number]").fill("1")
    pg.wait_for_selector(".sp-count", timeout=10_000)
    assert f"{expected_pages} ไฟล์" in pg.locator(".sp-count").inner_text(), \
        f"ตัวอย่างจำนวนไฟล์ที่จะได้ไม่ตรง: {pg.locator('.sp-count').inner_text()}"
    pg.get_by_role("button", name="แยกไฟล์").click()
    pg.wait_for_selector(".results button", timeout=20_000)
    out = dl(pg, lambda: pg.locator(".results button", has_text="ZIP").first.click(), "pdfsplit_out.zip")
    with zipfile.ZipFile(out) as z:
        names = z.namelist()
        assert len(names) == expected_pages, f"จำนวนไฟล์ที่แยกได้ไม่ตรงจำนวนหน้า: {len(names)} vs {expected_pages}"
        for n in names:
            data = z.read(n)
            assert data[:4] == b"%PDF", f"{n} ไม่ใช่ไฟล์ PDF จริง (header: {data[:8]!r})"
            sub = fitz.open(stream=data, filetype="pdf")
            assert sub.page_count == 1, f"{n} ควรมี 1 หน้า ได้ {sub.page_count}"


def case_image_convert(pg):
    src = SAMPLES / "ตัวอย่าง-รูปภาพ-2.png"
    goto(pg, "image-convert")
    pg.set_input_files("input[type=file]", str(src), timeout=SIF_TIMEOUT)
    pg.wait_for_selector(".file-row", timeout=15_000)
    # ‼️ ต้องจำกัดขอบเขตไว้ในแผงเครื่องมือ เพราะหน้าแรกมี <select> เรียงลำดับ
    #    ที่ยังอยู่ใน DOM (ซ่อนด้วย CSS) locator("select") เปล่า ๆ จึงเจอ 2 ตัวแล้วพัง
    pg.locator(".panel select").first.select_option("webp")
    pg.get_by_role("button", name="แปลงไฟล์").click()
    pg.wait_for_selector(".results .result", timeout=20_000)
    out = dl(pg, lambda: pg.locator(".results .result button").first.click(), "imgconvert_out.webp")
    im = Image.open(out)
    im.load()
    assert im.format == "WEBP", f"ชนิดไฟล์ผลลัพธ์ไม่ใช่ WEBP ตามที่เลือก: {im.format}"


# ─────────────────────────────────────────────────────────────────────────
# ② เคสไฟล์เสีย 5 ชนิด — ต้องขึ้นข้อความไทยที่คนอ่านรู้เรื่อง ไม่ค้าง ไม่โชว์ error ดิบ
# ─────────────────────────────────────────────────────────────────────────
def case_broken_garbage_pdf(pg, path):
    goto(pg, "pdf-to-text")
    pg.set_input_files("input[type=file]", str(path), timeout=SIF_TIMEOUT)
    pg.wait_for_selector(".file-row", timeout=15_000)
    pg.get_by_role("button", name="ดึงข้อความ").click()
    pg.wait_for_selector(".status-wrap .status.show.err", timeout=15_000)
    msg = pg.locator(".status-wrap .status").inner_text()
    assert_readable_thai_error(msg, "PDF ข้างในเป็นขยะ")
    assert ("เสียหาย" in msg) or ("ไม่ใช่ PDF" in msg), f"ข้อความไม่ได้บอกว่าไฟล์เสีย/ไม่ใช่ PDF จริง: {msg}"


def case_broken_empty_xlsx(pg, path):
    goto(pg, "excel-csv")
    pg.set_input_files("input[type=file]", str(path), timeout=SIF_TIMEOUT)
    pg.wait_for_selector(".file-row", timeout=15_000)
    pg.get_by_role("button", name="แปลงไฟล์").click()
    pg.wait_for_selector(".status-wrap .status.show.err", timeout=15_000)
    msg = pg.locator(".status-wrap .status").inner_text()
    assert_readable_thai_error(msg, "XLSX ขนาด 0 ไบต์")


def case_broken_tis620_csv(pg, path):
    goto(pg, "thai-encoding")
    pg.set_input_files("input[type=file]", str(path), timeout=SIF_TIMEOUT)
    pg.wait_for_selector(".enc-card", timeout=15_000)
    card = pg.locator(".enc-card").inner_text()
    assert ("TIS-620" in card) or ("Windows-874" in card), f"ไม่ตรวจพบว่าไฟล์นี้เป็น TIS-620: {card}"
    good_preview = pg.locator(".preview-text.good").inner_text()
    assert re.search(r"[ก-๙]", good_preview), f"หลังซ่อมแล้วยังไม่มีข้อความไทยอ่านออก: {good_preview}"
    assert "�" not in good_preview, f"หลังซ่อมยังมีอักขระเพี้ยนหลงเหลือ: {good_preview}"


def case_broken_jpg(pg, path):
    goto(pg, "image-resize")
    pg.set_input_files("input[type=file]", str(path), timeout=SIF_TIMEOUT)
    pg.wait_for_selector(".file-row", timeout=15_000)
    # ‼️ ".ws-empty" ใช้ซ้ำ 2 จุดในหน้านี้ (กล่อง "ยังไม่มีไฟล์" ของ workspace ที่ถูกซ่อนอยู่
    #    + กล่องข้อความพรีวิวพังของ image-resize) — เจาะจงด้วย has_text กันไป wait ตัวที่ซ่อนอยู่ค้าง
    err_box = pg.locator(".ws-empty", has_text="สร้างตัวอย่างไม่ได้")
    err_box.wait_for(state="visible", timeout=10_000)
    texts = err_box.inner_text()
    assert_readable_thai_error(texts, "JPG header เพี้ยน")
    assert "เสียหาย" in texts, f"ไม่ได้บอกผู้ใช้ว่าไฟล์รูปเสียหาย: {texts}"
    state = pg.locator(".file-row").get_attribute("data-state")
    assert state in ("pending", "error"), f"แถวไฟล์ค้างอยู่สถานะ {state!r} (ไม่ควรค้าง working)"


def case_broken_fake_docx(pg, path):
    goto(pg, "word-clean")
    pg.set_input_files("input[type=file]", str(path), timeout=SIF_TIMEOUT)
    pg.wait_for_selector(".file-row", timeout=15_000)
    pg.wait_for_selector(".status-wrap .status.show.err", timeout=15_000)
    msg = pg.locator(".status-wrap .status").inner_text()
    assert_readable_thai_error(msg, "DOCX ที่เป็น zip เปล่า")


# ─────────────────────────────────────────────────────────────────────────
# runner
# ─────────────────────────────────────────────────────────────────────────
def run_case(name, fn, *args):
    try:
        fn(*args)
        log(f"  ✅ {name}")
        return True
    except AssertionError as e:
        log(f"  ❌ {name}")
        log(f"     └─ {e}")
        return False
    except PWTimeout as e:
        log(f"  ❌ {name}")
        log(f"     └─ timeout: {str(e).splitlines()[0]}")
        return False
    except Exception as e:
        log(f"  ❌ {name}")
        log(f"     └─ {type(e).__name__}: {e}")
        return False


FEATURE_CASES = [
    ("excel-csv — Excel → CSV", case_excel_csv),
    ("pdf-merge — รวม 2 PDF เป็นเล่มเดียว", case_pdf_merge),
    ("image-resize — ย่อ+บีบอัดรูป", case_image_resize),
    ("thai-encoding — ซ่อม TIS-620 → UTF-8", case_thai_encoding),
    ("thai-id — ตรวจเลขบัตร ปชช. 13 หลัก", case_thai_id),
    ("thai-date — พ.ศ. → ค.ศ. ทั้งคอลัมน์", case_thai_date),
    ("word-clean — ล้างร่องรอยเอกสาร Word", case_word_clean),
    ("pdf-to-text — PDF → ข้อความ", case_pdf_to_text),
    ("images-to-pdf — รูปภาพ → PDF", case_images_to_pdf),
    ("powerpoint-to-word — PPTX → Word", case_powerpoint_to_word),
    ("pdf-to-images — PDF → รูปภาพ (ZIP)", case_pdf_to_images),
    ("pdf-watermark — ใส่ลายน้ำ PDF", case_pdf_watermark),
    ("pdf-split — แยกไฟล์ PDF (ZIP)", case_pdf_split),
    ("image-convert — แปลงชนิดไฟล์รูป", case_image_convert),
]

BROKEN_CASES = [
    ("ไฟล์เสีย: PDF ข้างในเป็นขยะ", case_broken_garbage_pdf, "garbage_pdf"),
    ("ไฟล์เสีย: XLSX ขนาด 0 ไบต์", case_broken_empty_xlsx, "empty_xlsx"),
    ("ไฟล์เสีย: CSV เข้ารหัส TIS-620", case_broken_tis620_csv, "tis620_csv"),
    ("ไฟล์เสีย: JPG header เพี้ยน", case_broken_jpg, "broken_jpg"),
    ("ไฟล์เสีย: DOCX ที่เป็น zip เปล่า", case_broken_fake_docx, "fake_docx"),
]


def run_main_suite(pg, broken):
    results = []
    log("═" * 66)
    log("① เทสฟีเจอร์จริง — ใส่ไฟล์ → กดปุ่ม → ดาวน์โหลด → ตรวจเนื้อในไฟล์จริง")
    log("═" * 66)
    for name, fn in FEATURE_CASES:
        results.append((name, run_case(name, fn, pg)))

    log("")
    log("═" * 66)
    log("② เคสไฟล์เสีย 5 ชนิด — ต้องขึ้นข้อความไทยอ่านรู้เรื่อง ไม่ค้าง ไม่ error ดิบ")
    log("═" * 66)
    for name, fn, key in BROKEN_CASES:
        results.append((name, run_case(name, fn, pg, broken[key])))

    passed = sum(1 for _, ok in results if ok)
    failed = len(results) - passed
    log("")
    log("━" * 66)
    log(f"สรุป: ผ่าน {passed} · ตก {failed}  (ทั้งหมด {len(results)} เคส)")
    if failed:
        log("เคสที่ตก:")
        for name, ok in results:
            if not ok:
                log(f"  ❌ {name}")
    return results


# ─────────────────────────────────────────────────────────────────────────
# ③ พิสูจน์ว่าเทสแดงได้เมื่อของพัง — ฉีดบั๊กปลอมด้วย page.route() แล้วต้องแดง
#    ถอดออกแล้วต้องกลับเขียว — เลือก 3 เคส คนละกลไก (abort ไลบรารี x2, สลับไฟล์เครื่องมือ x1)
# ─────────────────────────────────────────────────────────────────────────
STUB_JS = (
    "export function mount(tool){"
    " const d=document.createElement('div');"
    " d.className='tool-head';"
    " d.textContent='บั๊กปลอม: เครื่องมือนี้ถูกแทนที่ด้วย stub ที่ไม่มี dropzone จริง';"
    " return d;"
    "}"
)


def redproof(broken):
    log("")
    log("═" * 66)
    log("③ พิสูจน์ว่าเทสแดงได้เมื่อของพัง (page.route ฉีดบั๊กปลอม)")
    log("═" * 66)
    proof = []
    with sync_playwright() as p:
        browser = p.chromium.launch()
        pg = browser.new_page(viewport={"width": 1280, "height": 1000}, service_workers="block")

        # ── เคส 1/3: excel-csv — บล็อกไลบรารี xlsx ทั้ง vendor/ (local) และ CDN สำรอง ──
        log("-- เคส 1/3: excel-csv (บล็อกไลบรารี xlsx ทั้ง local + CDN สำรอง) --")
        pg.route("**/vendor/xlsx.full.min.js", lambda r: r.abort())
        pg.route("**cdnjs.cloudflare.com**xlsx**", lambda r: r.abort())
        red_ok = False
        try:
            case_excel_csv(pg)
            log("  ⚠️ ไม่แดงตามคาด — บั๊กปลอมไม่ทำงาน")
        except Exception as e:
            red_ok = True
            log(f"  ✅ แดงตามคาด: {type(e).__name__}: {str(e).splitlines()[0][:160]}")
        pg.unroute("**/vendor/xlsx.full.min.js")
        pg.unroute("**cdnjs.cloudflare.com**xlsx**")
        green_ok = run_case("excel-csv (ถอดบั๊กปลอมแล้ว — ต้องเขียว)", case_excel_csv, pg)
        proof.append(("excel-csv (block xlsx lib)", red_ok, green_ok))

        # ── เคส 2/3: pdf-merge — บล็อกไลบรารี pdf-lib ทั้ง local + CDN สำรอง ──────────
        log("-- เคส 2/3: pdf-merge (บล็อกไลบรารี pdf-lib ทั้ง local + CDN สำรอง) --")
        pg.route("**/vendor/pdf-lib.min.js", lambda r: r.abort())
        pg.route("**cdnjs.cloudflare.com**pdf-lib**", lambda r: r.abort())
        red_ok = False
        try:
            case_pdf_merge(pg)
            log("  ⚠️ ไม่แดงตามคาด — บั๊กปลอมไม่ทำงาน")
        except Exception as e:
            red_ok = True
            log(f"  ✅ แดงตามคาด: {type(e).__name__}: {str(e).splitlines()[0][:160]}")
        pg.unroute("**/vendor/pdf-lib.min.js")
        pg.unroute("**cdnjs.cloudflare.com**pdf-lib**")
        green_ok = run_case("pdf-merge (ถอดบั๊กปลอมแล้ว — ต้องเขียว)", case_pdf_merge, pg)
        proof.append(("pdf-merge (block pdf-lib lib)", red_ok, green_ok))

        # ── เคส 3/3: pdf-to-images — สลับไฟล์เครื่องมือเป็น stub ที่ไม่มี dropzone จริง ──
        log("-- เคส 3/3: pdf-to-images (ตอบไฟล์เครื่องมือเป็น stub ว่างเปล่า) --")
        pg.route("**/src/tools/pdf-to-images.js",
                  lambda r: r.fulfill(content_type="application/javascript", body=STUB_JS))
        red_ok = False
        try:
            case_pdf_to_images(pg)
            log("  ⚠️ ไม่แดงตามคาด — บั๊กปลอมไม่ทำงาน")
        except Exception as e:
            red_ok = True
            log(f"  ✅ แดงตามคาด: {type(e).__name__}: {str(e).splitlines()[0][:160]}")
        pg.unroute("**/src/tools/pdf-to-images.js")
        green_ok = run_case("pdf-to-images (ถอดบั๊กปลอมแล้ว — ต้องเขียว)", case_pdf_to_images, pg)
        proof.append(("pdf-to-images (swap stub module)", red_ok, green_ok))

        browser.close()

    log("")
    ok_all = all(r and g for _, r, g in proof)
    for name, r, g in proof:
        log(f"  {name}: แดงตอนพัง={'✅' if r else '❌'} · เขียวหลังถอดบั๊ก={'✅' if g else '❌'}")
    log(f"สรุปพิสูจน์แดง/เขียว: {'✅ ผ่านครบ 3/3' if ok_all else '❌ มีบางเคสไม่ผ่าน'}")
    return ok_all


def main():
    broken = make_broken_files()
    try:
        log(f"FK_BASE = {BASE}")
        log(f"WORKDIR (ไฟล์เสีย+ไฟล์ที่ดาวน์โหลด, อยู่นอกโปรเจกต์) = {WORKDIR}")
        log("")

        all_ok = True

        for round_no in (1, 2):
            log("")
            log("#" * 66)
            log(f"# รอบที่ {round_no}/2 ของชุดเต็ม (เช็คว่าไม่ flaky)")
            log("#" * 66)
            with sync_playwright() as p:
                browser = p.chromium.launch()
                pg = browser.new_page(viewport={"width": 1280, "height": 1000}, service_workers="block")
                results = run_main_suite(pg, broken)
                browser.close()
            if any(not ok for _, ok in results):
                all_ok = False

        redproof_ok = redproof(broken)

        log("")
        log("=" * 66)
        log(f"จบการรัน — ชุดเต็ม 2 รอบ: {'มีเคสตก (ดูรายละเอียดด้านบน)' if not all_ok else 'ผ่านครบทุกเคสทั้ง 2 รอบ'}")
        log(f"พิสูจน์แดง/เขียว 3 เคส: {'ผ่านครบ' if redproof_ok else 'ไม่ครบ'}")
        return 0 if (redproof_ok) else 1
    finally:
        shutil.rmtree(WORKDIR, ignore_errors=True)


if __name__ == "__main__":
    sys.exit(main())
