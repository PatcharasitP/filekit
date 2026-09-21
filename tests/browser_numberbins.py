# เครื่องมือ "จัดกลุ่มตัวเลขเป็นช่วง" ตรวจของที่ผู้ใช้เอาไปใช้ต่อจริง
#
# ‼️ ที่เทสนี้จับ ไม่ใช่แค่ "กดแล้วไม่ error" แต่คือ
#    ① ค่าบนหน้าจอตรงกับที่คำนวณอิสระด้วย Python ทุกกลุ่ม (ถ้ากติกาขอบเพี้ยนจะแดงทันที)
#    ② ไฟล์ที่ดาวน์โหลดมี "ป้าย + เลขเรียง" ถูกต้องครบทุกแถว ไม่ใช่แค่บนจอ
#    ③ ล็อกรายตัวแล้วลำดับการเรียงไม่พัง (สิ่งที่พี่ปอนด์สั่งไว้ตรง ๆ)
#    ④ โค้ด Power Query ที่คัดลอกไป มีจุดตัดและรายการล็อกจริง
#
# ถ้าผิดจะรู้ได้ยังไง: ข้อมูลทดสอบมีค่าที่ "เท่ากับจุดตัดพอดี" (5.0) และค่าที่อ่านไม่ออก
# ถ้าใครเผลอเปลี่ยนกติกาขอบเป็นไม่รวม หรือกลืนแถวเสียเงียบ ๆ ข้อ ① กับ ② จะแดงทันที
#
# รัน: ../.venv/bin/python tests/browser_numberbins.py   (หรือผ่าน tests/run.sh browser_numberbins)
import os
import re
import pathlib
import socket
import subprocess
import sys
import tempfile
import time

import openpyxl
from playwright.sync_api import sync_playwright

sys.path.insert(0, str(pathlib.Path(__file__).resolve().parent))
import fkui  # noqa: E402  ตัวช่วยกลางที่รู้จักโครงหน้าเครื่องมือ v2

ROOT = pathlib.Path(__file__).resolve().parent.parent
TMP = pathlib.Path(tempfile.mkdtemp(prefix="filekit_nbins_"))
P, F = 0, []


def ck(name, got, want):
    global P
    ok_ = got == want
    if ok_:
        P += 1
    else:
        F.append(f"{name}\n      ได้    : {got!r}\n      ควรได้ : {want!r}")
    print(f"  {'✅' if ok_ else '❌'} {name}")


def ok(name, cond, note=""):
    global P
    if cond:
        P += 1
    else:
        F.append(f"{name} {note}")
    print(f"  {'✅' if cond else '❌'} {name}{(' ' + note) if note and not cond else ''}")


def free_port():
    s = socket.socket()
    s.bind(("127.0.0.1", 0))
    port = s.getsockname()[1]
    s.close()
    return port


# ระยะทางการกระจัดสถานีจริงจากชุดข้อมูลของงาน (Data/SiteRelocate_Mock) 40 คู่
# บวกกับดัก 3 แถว: ค่าเท่ากับจุดตัดพอดี, ค่าว่าง, ค่าที่อ่านไม่ออก
PAIRS = [
    ("BKK1000", 6.961), ("CMI1001", 4.601), ("PKT1002", 6.208), ("KKN1003", 5.024),
    ("NMA1004", 6.606), ("HYI1005", 2.235), ("CBI1006", 7.994), ("AYA1007", 4.596),
    ("UDN1008", 5.734), ("SKA1009", 3.097), ("BKK1010", 4.154), ("CMI1011", 4.916),
    ("PKT1012", 3.530), ("KKN1013", 0.498), ("NMA1014", 5.754), ("HYI1015", 3.929),
    ("CBI1016", 5.139), ("AYA1017", 7.936), ("UDN1018", 2.670), ("SKA1019", 7.149),
    ("BKK1020", 0.554), ("CMI1021", 4.349), ("PKT1022", 4.034), ("KKN1023", 4.594),
    ("NMA1024", 7.360), ("HYI1025", 4.263), ("CBI1026", 6.823), ("AYA1027", 4.134),
    ("UDN1028", 3.774), ("SKA1029", 7.508), ("BKK1030", 5.077), ("CMI1031", 5.196),
    ("PKT1032", 6.250), ("KKN1033", 1.337), ("NMA1034", 7.081), ("HYI1035", 7.504),
    ("CBI1036", 3.472), ("AYA1037", 0.409), ("UDN1038", 2.987), ("SKA1039", 4.253),
    ("EDGE0001", 5.0),        # เท่ากับจุดตัดพอดี ต้องอยู่ช่วงล่าง
    ("BLANK001", None),       # ค่าว่าง
    ("BAD00001", "ไม่ทราบ"),   # อ่านไม่ออก
]
BREAKS = [2, 5, 7]            # จุดตัดที่เทสนี้พิมพ์เอง (แทนค่าที่เครื่องเสนอ)


def band_of(v, breaks):
    """ตัวคำนวณอิสระของฝั่งเทส เขียนคนละวิธีกับ src/binkit.js โดยเจตนา"""
    if v is None or isinstance(v, str):
        return None
    for i, b in enumerate(breaks):
        if v <= b:
            return i
    return len(breaks)


def expected_counts(breaks):
    """คืนเป็น {เลขเรียงบนหน้าจอ: กี่แถว} — เลขเรียงเริ่มที่ 1 ส่วนค่าว่างเป็น 0 เหมือนที่ตกลงกันไว้"""
    counts = {}
    for _, v in PAIRS:
        i = band_of(v, breaks)
        key = 0 if i is None else i + 1
        counts[key] = counts.get(key, 0) + 1
    return counts


def make_file():
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "PAIR"
    ws.append(["SITE_CODE", "PROVINCE", "DISTANCE_KM"])
    for code, v in PAIRS:
        ws.append([code, "จังหวัดทดสอบ", v])
    p = TMP / "ระยะทางการกระจัด.xlsx"
    wb.save(p)
    return p


def groups_on_screen(pg):
    """ตารางกลุ่มบนจอ → [{sort, label, count, min, max}]"""
    return pg.evaluate("""() => [...document.querySelectorAll('.nb-groups tbody tr')].map(tr => {
      const td = tr.querySelectorAll('td');
      return {
        sort: +td[0].textContent.trim(),
        label: (td[1].querySelector('.nb-glabel') || td[1]).textContent.trim(),
        count: +td[2].textContent.replace(/,/g, ''),
        pct: +td[3].textContent,
        min: td[5] ? td[5].textContent.trim() : '',
        max: td[6] ? td[6].textContent.trim() : '',
      };
    })""")


def main():
    base = os.environ.get("FK_BASE")
    server = None
    if not base:
        port = free_port()
        server = subprocess.Popen([sys.executable, "-m", "http.server", str(port)],
                                  cwd=str(ROOT), stdout=subprocess.DEVNULL, stderr=subprocess.DEVNULL)
        base = f"http://localhost:{port}"
        time.sleep(1.5)

    xlsx = make_file()
    try:
        with sync_playwright() as pw:
            b = pw.chromium.launch()
            ctx = b.new_context(viewport={"width": 1400, "height": 1000}, accept_downloads=True,
                                timezone_id="Asia/Bangkok")
            pg = ctx.new_page()
            errs = []
            pg.on("pageerror", lambda e: errs.append(str(e)[:200]))
            pg.on("console", lambda m: errs.append(m.text[:200]) if m.type == "error" else None)

            pg.goto(f"{base}/#/number-bins", wait_until="networkidle")
            pg.wait_for_selector(".dz")
            pg.locator(".dz input[type=file]").first.set_input_files(str(xlsx))
            pg.wait_for_timeout(1800)

            print("\n── ① เปิดไฟล์แล้วต้องรู้เองว่าคอลัมน์ไหนคือตัวเลข ──")
            picked = pg.evaluate("() => {const s=document.querySelectorAll('.s2-stage select, .s2-side-bd select, .s2-stage select, .s2-side-bd select, .ws-left select');"
                                 "return [...s].map(x => x.options[x.selectedIndex]?.textContent || '')}")
            ok("เดาคอลัมน์ตัวเลขถูกเป็น DISTANCE_KM", "DISTANCE_KM" in picked, f"ได้ {picked}")
            ok("เดาคอลัมน์ชื่อรายการถูกเป็น SITE_CODE", "SITE_CODE" in picked, f"ได้ {picked}")
            chips = pg.evaluate("() => [...document.querySelectorAll('.nb-chip')].map(c => c.textContent)")
            n_num = len([1 for _, v in PAIRS if isinstance(v, (int, float))])
            ok(f"บอกว่าอ่านเป็นตัวเลขได้ {n_num} แถว",
               any("อ่านเป็นตัวเลขได้" in c and str(n_num) in c for c in chips), f"ได้ {chips}")
            ok("บอกว่ามีแถวว่าง 1 แถว", any("ว่าง" in c and "1" in c for c in chips), f"ได้ {chips}")
            ok("บอกว่าอ่านไม่ออก 1 แถว", any("อ่านไม่ออก" in c and "1" in c for c in chips), f"ได้ {chips}")
            ok("บอกเรื่องทศนิยมให้รู้ตัว", any("ทศนิยม" in c for c in chips), f"ได้ {chips}")

            print("\n── ② กราฟการกระจายต้องวาดของจริง ไม่ใช่กล่องเปล่า ──")
            bars = pg.evaluate("() => [...document.querySelectorAll('.nb-chart .bar')]"
                               ".map(r => +r.getAttribute('height'))")
            ok(f"มีแท่งการกระจายจริง {len(bars)} แท่ง", len(bars) >= 5, f"ได้ {len(bars)}")
            ok("แท่งมีความสูงมากกว่าศูนย์", bars and max(bars) > 10, f"สูงสุด {max(bars) if bars else 0}")

            print("\n── ③ พิมพ์จุดตัดเอง แล้วค่าทุกกลุ่มต้องตรงกับที่คำนวณอิสระ ──")
            pg.evaluate("""(txt) => {
              const inp = [...document.querySelectorAll('.s2-side-bd input.nb-num, .s2-stage input.nb-num, .ws-right input.nb-num')]
                .find(i => i.placeholder && i.placeholder.includes('0, 5, 10'));
              inp.value = txt;
              inp.dispatchEvent(new Event('change', { bubbles: true }));
            }""", ", ".join(str(b) for b in BREAKS))
            pg.wait_for_timeout(600)
            rows = groups_on_screen(pg)
            want = {k: v for k, v in expected_counts(BREAKS).items() if k > 0}
            got = {r["sort"]: r["count"] for r in rows if r["sort"] > 0}
            ck("จำนวนแถวของทุกกลุ่มตรงกับที่คำนวณเอง", got, want)
            ck("กลุ่มค่าว่างได้เลขเรียง 0 และมี 2 แถว (ว่าง + อ่านไม่ออก)",
               [(r["sort"], r["count"]) for r in rows if r["sort"] == 0], [(0, 2)])
            ck("ผลรวมทุกกลุ่มเท่ากับจำนวนแถวในไฟล์",
               sum(r["count"] for r in rows), len(PAIRS))
            labels = [r["label"] for r in rows if r["sort"] > 0]
            # ‼️ ข้อมูลชุดนี้มีทศนิยม ป้ายต้องเป็นแบบช่วงแท้เอง ไม่ใช่ "3-5" ซึ่งจะทำให้คนเข้าใจผิด
            #    (ค่า 2.235 ก็อยู่ในกลุ่มนั้น ทั้งที่ป้ายเขียนว่าเริ่มที่ 3)
            ck("ป้ายอัตโนมัติถูกต้องตามจุดตัด", labels, ["≤ 2", "> 2 ถึง 5", "> 5 ถึง 7", "> 7"])
            ok("ข้อมูลมีทศนิยม สวิตช์ป้ายจำนวนเต็มต้องถูกปิดให้เอง",
               pg.evaluate("() => document.querySelector('.s2-side-bd .nb-switch input, .s2-stage .nb-switch input, .ws-right .nb-switch input').checked") is False)
            ok("ชิปบอกด้วยว่าเปลี่ยนป้ายให้แล้วเพราะมีทศนิยม",
               any("ช่วงแท้" in c for c in
                   pg.evaluate("() => [...document.querySelectorAll('.nb-chip')].map(c => c.textContent)")))
            edge = [r for r in rows if r["sort"] == 2][0]
            ok("ค่า 5.0 ที่เท่ากับจุดตัดพอดี อยู่ช่วงล่าง (ค่าสูงสุดของกลุ่มนี้คือ 5)",
               edge["max"] in ("5", "5.0", "5.000"), f"ได้ {edge['max']}")

            print("\n── ④ เปลี่ยนวิธีเสนอจุดตัด แล้วจุดตัดต้องเปลี่ยนจริง ──")
            def breaks_now():
                return pg.evaluate("""() => ([...document.querySelectorAll('.s2-side-bd input.nb-num, .s2-stage input.nb-num, .ws-right input.nb-num')]
                  .find(i => i.placeholder && i.placeholder.includes('0, 5, 10')) || {}).value""")
            def pick_method(v):
                pg.evaluate("""(v) => {
                  /* ‼️ หาช่องเลือกจากตัวเลือกที่มันมี ห้ามนับลำดับ (22/09/2026)
                     v2 รวมแผงซ้ายกับขวาไว้แผงเดียว ช่องแรกจึงกลายเป็นช่องเลือกชีต
                     เทสเลยไปตั้งค่าชีตแทนวิธีแบ่ง จุดตัดไม่เปลี่ยนเลยทั้งสามวิธี */
                  const s = [...document.querySelectorAll('.s2-side-bd select, .ws-right select')]
                    .find((x) => [...x.options].some((o) => o.value === 'quantile'));
                  s.value = v; s.dispatchEvent(new Event('change', { bubbles: true }));
                }""", v)
                pg.wait_for_timeout(500)
                return breaks_now()

            b_round = pick_method("round")
            b_quant = pick_method("quantile")
            b_nat = pick_method("natural")
            ok("วิธีเลขกลมให้จุดตัดที่เป็นเลขกลมจริง",
               all(float(x) == round(float(x), 1) for x in b_round.split(", ") if x), f"ได้ {b_round}")
            ok("วิธีจำนวนเท่ากันให้จุดตัดคนละชุดกับเลขกลม", b_quant != b_round, f"ทั้งคู่ได้ {b_round}")
            ok("วิธีช่องว่างธรรมชาติก็ให้อีกชุดหนึ่ง", b_nat not in ("", b_round), f"ได้ {b_nat}")
            counts_q = {r["sort"]: r["count"] for r in groups_on_screen(pg) if r["sort"] > 0}
            ok("ทุกกลุ่มยังมีแถวรวมกันครบ",
               sum(counts_q.values()) + 2 == len(PAIRS), f"ได้ {sum(counts_q.values()) + 2}")

            print("\n── ⑤ ล็อกรายตัว ต้องชนะจุดตัด และลำดับต้องไม่พัง ──")
            pg.evaluate("""(txt) => {
              const inp = [...document.querySelectorAll('.s2-side-bd input.nb-num, .s2-stage input.nb-num, .ws-right input.nb-num')]
                .find(i => i.placeholder && i.placeholder.includes('0, 5, 10'));
              inp.value = txt; inp.dispatchEvent(new Event('change', { bubbles: true }));
            }""", ", ".join(str(b) for b in BREAKS))
            pg.wait_for_timeout(400)

            def add_lock(code, group):
                pg.evaluate("""([code, group]) => {
                  const ins = [...document.querySelectorAll('.s2-side-bd input.nb-num, .s2-stage input.nb-num, .ws-right input.nb-num')];
                  const k = ins.find(i => i.placeholder && i.placeholder.includes('SKA1029'));
                  const g = ins.find(i => i.placeholder && (i.placeholder.includes('หมวด') || i.placeholder.includes('group')));
                  k.value = code; g.value = group;
                }""", [code, group])
                # ‼️ ข้ามปุ่มหัวข้อกลุ่มที่พับได้ ไม่งั้นจะไปกดพับกลุ่ม "ล็อกรายตัว" แทน
                #    ต้องตรงคำว่า ล็อก ทั้งปุ่ม เพราะหัวข้อกลุ่มก็มีคำนี้อยู่ในชื่อ
                pg.locator(".s2-side-bd button:visible:not(.ws-fold-btn), .ws-right button:not(.ws-fold-btn)",
                           has_text=re.compile(r"^\s*ล็อก\s*$")).first.click()
                pg.wait_for_timeout(500)

            add_lock("KKN1013", "> 7")          # ค่า 0.498 แต่บังคับไปกลุ่มบนสุดที่มีอยู่แล้ว
            add_lock("SKA1029", "เฝ้าระวัง")     # ค่า 7.508 บังคับไปหมวดใหม่
            add_lock("AYA1037", "เฝ้าระวัง")     # ค่า 0.409 หมวดใหม่เดียวกัน

            rows = groups_on_screen(pg)
            by_label = {r["label"]: r for r in rows}
            ok("หมวดใหม่ที่ล็อกโผล่ในตารางกลุ่ม", "เฝ้าระวัง" in by_label, f"ได้ {list(by_label)}")
            ck("หมวดใหม่ได้เลขเรียงต่อท้ายชุดเดิม (5)", by_label.get("เฝ้าระวัง", {}).get("sort"), 5)
            ck("หมวดใหม่มี 2 แถวตามที่ล็อกไว้", by_label.get("เฝ้าระวัง", {}).get("count"), 2)
            want_top = expected_counts(BREAKS)[4] + 1 - 1   # +KKN1013 ที่ล็อกเข้ามา, -SKA1029 ที่ย้ายออก
            ck("กลุ่ม > 7 นับใหม่ถูกหลังมีคนย้ายเข้าและย้ายออก",
               by_label.get("> 7", {}).get("count"), want_top)
            ck("เลขเรียงของกลุ่มเดิมไม่ขยับเลย",
               [(r["label"], r["sort"]) for r in rows if r["sort"] in (1, 2, 3, 4)],
               [("≤ 2", 1), ("> 2 ถึง 5", 2), ("> 5 ถึง 7", 3), ("> 7", 4)])
            sorts = [r["sort"] for r in rows if r["sort"] > 0]
            ok("ไม่มีสองกลุ่มใช้เลขเรียงซ้ำกัน", len(sorts) == len(set(sorts)), f"ได้ {sorts}")

            print("\n── ⑥ ตัวอย่างผลลัพธ์รายแถว ต้องบอกด้วยว่าแถวไหนถูกล็อก ──")
            pg.locator(".seg-item", has_text="ตัวอย่างผลลัพธ์").first.click()
            pg.wait_for_timeout(400)
            prow = pg.evaluate("""() => {
              const out = {};
              for (const tr of document.querySelectorAll('.nb-rows tbody tr')) {
                const td = tr.querySelectorAll('td');
                out[td[0].textContent.replace('ล็อก', '').trim()] =
                  { v: td[1].textContent.trim(), band: td[2].textContent.trim(), sort: +td[3].textContent };
              }
              return out;
            }""")
            ck("แถวที่ถูกล็อกไปหมวดใหม่ ได้ป้ายและเลขเรียงของหมวดนั้น",
               (prow.get("SKA1029") or {}).get("band"), "เฝ้าระวัง")
            ck("แถวที่ล็อกไปกลุ่มเดิม ได้เลขเรียงของกลุ่มเดิม ไม่ใช่เลขใหม่",
               (prow.get("KKN1013") or {}).get("sort"), 4)
            ok("แถวที่ถูกล็อกมีป้ายกำกับให้เห็นบนจอ",
               pg.evaluate("() => !!document.querySelector('.nb-tag')"))

            print("\n── ⑦ โค้ด Power Query ต้องมีค่าจริง ไม่ใช่โครงเปล่า ──")
            pg.locator(".seg-item", has_text="Power Query").first.click()
            pg.wait_for_timeout(500)
            code = pg.evaluate("() => document.querySelector('.nb-code code')?.textContent || ''")
            ok("มีจุดตัดจริงในโค้ด", "Breaks = {2, 5, 7}" in code, f"ได้ {code[:120]!r}")
            ok("ป้ายในโค้ดตรงกับป้ายบนจอ (ช่วงแท้ ไม่ใช่จำนวนเต็ม)", '"> 2 ถึง 5"' in code, f"ได้ {code[:200]!r}")
            ok("มีรายการที่ล็อกไว้ในโค้ด", '#"SKA1029" = "เฝ้าระวัง"' in code)
            ok("อ้างคอลัมน์คีย์ที่เลือกไว้จริง", 'Record.FieldOrDefault(_, "SITE_CODE", "")' in code)
            ok("ตั้งชื่อคอลัมน์ผลลัพธ์ตามชื่อคอลัมน์ในไฟล์",
               '"DISTANCE_KM Band"' in code and '"DISTANCE_KM Band Sort"' in code)
            ok("ใช้กติกาขอบบนรวม (v > b)", "v > b" in code)

            print("\n── ⑧ ไฟล์ที่ดาวน์โหลด ต้องถูกครบทุกแถว (ของจริงที่ผู้ใช้เอาไปใช้) ──")
            with pg.expect_download() as dl:
                fkui.dl_button(pg, "ดาวน์โหลดเป็น Excel").click()
            out = TMP / "out.xlsx"
            dl.value.save_as(str(out))
            wb2 = openpyxl.load_workbook(out)
            ws2 = wb2.active
            head = [c.value for c in ws2[1]]
            ck("คอลัมน์เดิมอยู่ครบ แล้วต่อท้ายด้วยป้ายกับเลขเรียง", head,
               ["SITE_CODE", "PROVINCE", "DISTANCE_KM", "DISTANCE_KM Band", "DISTANCE_KM Band Sort"])
            got_rows = [(r[0], r[2], r[3], r[4]) for r in ws2.iter_rows(min_row=2, values_only=True)]
            ck("จำนวนแถวในไฟล์ครบ", len(got_rows), len(PAIRS))
            locks = {"KKN1013": ("> 7", 4), "SKA1029": ("เฝ้าระวัง", 5), "AYA1037": ("เฝ้าระวัง", 5)}
            labels4 = ["≤ 2", "> 2 ถึง 5", "> 5 ถึง 7", "> 7"]
            bad = []
            for code_, v, band, srt in got_rows:
                if code_ in locks:
                    want_b, want_s = locks[code_]
                elif v is None or isinstance(v, str):
                    want_b, want_s = "ไม่มีข้อมูล", 0
                else:
                    i = band_of(v, BREAKS)
                    want_b, want_s = labels4[i], i + 1
                if (band, srt) != (want_b, want_s):
                    bad.append(f"{code_} ค่า {v} ได้ {band}/{srt} ควรเป็น {want_b}/{want_s}")
            ck("ทุกแถวในไฟล์ได้ป้ายและเลขเรียงถูกต้อง", bad, [])

            print("\n── ⑨ ไม่มี error ในคอนโซล ──")
            real = [e for e in errs if "favicon" not in e.lower()]
            ck("เงียบสนิทตลอดการใช้งาน", real, [])

            b.close()
    finally:
        if server:
            server.terminate()

    print(f"\nสรุป: ผ่าน {P} ข้อ, ตก {len(F)} ข้อ")
    for f in F:
        print("  ❌ " + f)
    sys.exit(1 if F else 0)


if __name__ == "__main__":
    main()
