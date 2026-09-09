# -*- coding: utf-8 -*-
"""
เทสโซ่ (chain) — ใช้เครื่องมือ FileKit ต่อกันเป็นทอด ๆ เหมือนงานเอกสารจริง
แทนที่จะเทสแต่ละเครื่องมือแยกกัน (มีเทสอื่นทำแล้ว) ไฟล์นี้ "ดาวน์โหลดผลลัพธ์จริง
แล้วอัปโหลดกลับเข้าเครื่องมือถัดไป" จำลองสิ่งที่ผู้ใช้จริงทำ — จับบั๊กที่โผล่เฉพาะ
ตอนใช้ต่อเนื่อง (ไฟล์ผลลัพธ์ของตัวหนึ่งเปิดไม่ได้ในอีกตัว) ซึ่งเทสแยกส่วนจับไม่ได้

รัน: ../.venv/bin/python tests/browser_chain.py
     (เสิร์ฟเองที่ port 8925 ถ้ายังไม่มีอะไรฟังอยู่ที่พอร์ตนั้น — ตั้ง FK_BASE เพื่อยิงที่อื่นแทน)
"""
import io
import os
import re
import socket
import shutil
import subprocess
import sys
import tempfile
import time
import zipfile
import pathlib

from playwright.sync_api import sync_playwright

import openpyxl
import fitz             # pymupdf
import docx as pydocx   # python-docx
from PIL import Image

ROOT = pathlib.Path(__file__).resolve().parent.parent
SAMPLES = ROOT / "samples"
PORT = 8925
BASE = os.environ.get("FK_BASE", f"http://localhost:{PORT}").rstrip("/")

REG_TEXT = (ROOT / "src/registry.js").read_text(encoding="utf-8")
TOOL_IDS = set(re.findall(r'id:"([\w-]+)"', REG_TEXT))

WORKDIR = pathlib.Path(tempfile.mkdtemp(prefix="filekit_chain_"))
OUT_DIR = WORKDIR / "out"
GEN_DIR = WORKDIR / "gen"
OUT_DIR.mkdir(parents=True, exist_ok=True)
GEN_DIR.mkdir(parents=True, exist_ok=True)

STEP_TIMEOUT = 60_000  # ms — ต่อขั้น (บางเครื่องมือโหลดไลบรารีหนักตอนแรก)


def log(msg=""):
    print(msg, flush=True)


def port_open(host, port):
    with socket.socket(socket.AF_INET, socket.SOCK_STREAM) as s:
        s.settimeout(0.3)
        try:
            return s.connect_ex((host, port)) == 0
        except OSError:
            return False


def start_server_if_needed():
    """เสิร์ฟเองที่ port 8925 ถ้า BASE ชี้มาที่พอร์ตนี้และยังไม่มีอะไรฟังอยู่"""
    local = BASE in (f"http://localhost:{PORT}", f"http://127.0.0.1:{PORT}")
    if not local:
        log(f"  (ใช้ FK_BASE ภายนอก: {BASE} — ไม่เปิดเซิร์ฟเวอร์เอง)")
        return None
    if port_open("127.0.0.1", PORT):
        log(f"  (มีอะไรบางอย่างฟังอยู่ที่พอร์ต {PORT} แล้ว — ใช้ของเดิม ไม่เปิดซ้ำ)")
        return None
    proc = subprocess.Popen(
        [sys.executable, "-m", "http.server", str(PORT)],
        cwd=str(ROOT), stdout=subprocess.DEVNULL, stderr=subprocess.DEVNULL,
    )
    for _ in range(100):
        if port_open("127.0.0.1", PORT):
            break
        time.sleep(0.1)
    else:
        proc.terminate()
        raise RuntimeError(f"เปิดเซิร์ฟเวอร์ที่พอร์ต {PORT} ไม่สำเร็จภายในเวลาที่กำหนด")
    log(f"  เปิดเซิร์ฟเวอร์ท้องถิ่นที่พอร์ต {PORT} แล้ว (pid={proc.pid})")
    return proc


def goto(pg, tool_id):
    assert tool_id in TOOL_IDS, f"ไม่พบเครื่องมือ id={tool_id} ใน src/registry.js"
    pg.goto("about:blank")
    pg.goto(f"{BASE}/#/{tool_id}", wait_until="networkidle")
    pg.wait_for_selector(".tool-head", timeout=STEP_TIMEOUT)


def dl(pg, trigger, filename, timeout=STEP_TIMEOUT):
    with pg.expect_download(timeout=timeout) as di:
        trigger()
    d = di.value
    out = OUT_DIR / filename
    d.save_as(str(out))
    return out


def upload(pg, path, nth=0):
    pg.locator(".dz input[type=file]").nth(nth).set_input_files(str(path), timeout=STEP_TIMEOUT)


def new_page(browser):
    pg = browser.new_page(viewport={"width": 1360, "height": 980})
    errs = []
    pg.on("pageerror", lambda e: errs.append(str(e)))
    pg.on("console", lambda m: errs.append(m.text) if m.type == "error" else None)
    return pg, errs


def real_errors(errs):
    return [e for e in errs if "favicon" not in e.lower()]


# ─────────────────────────────────────────────────────────────────────────
# แต่ละโซ่คืน dict: {steps: [(ขั้น, ไฟล์เข้า, ไฟล์ออก, สิ่งที่ตรวจ, ผล)], ok, note}
# ─────────────────────────────────────────────────────────────────────────

def chain1_word_pdf_merge_compress(browser):
    """Word → PDF → รวม PDF → บีบอัด PDF → เปิดด้วย pymupdf ตรวจหน้า+ไทย"""
    steps = []
    pg, errs = new_page(browser)
    try:
        src1 = SAMPLES / "ตัวอย่าง-ใบเสนอราคา.docx"
        src2 = SAMPLES / "ตัวอย่าง-หนังสือแจ้งผลประเมิน.docx"

        # ① Word → PDF (แปลง 2 ไฟล์พร้อมกัน เอาไปรวมกันต่อขั้นถัดไป)
        goto(pg, "word-to-pdf")
        pg.locator(".dz input[type=file]").first.set_input_files(
            [str(src1), str(src2)], timeout=STEP_TIMEOUT)
        pg.wait_for_selector(".file-row", timeout=STEP_TIMEOUT)
        pg.get_by_role("button", name="แปลงเป็น PDF").click()
        pg.wait_for_selector(".results .result", timeout=STEP_TIMEOUT)
        pdf1 = dl(pg, lambda: pg.locator(".results .result button").nth(0).click(), "c1_word1.pdf")
        pdf2 = dl(pg, lambda: pg.locator(".results .result button").nth(1).click(), "c1_word2.pdf")
        p1n = fitz.open(str(pdf1)).page_count
        p2n = fitz.open(str(pdf2)).page_count
        steps.append(("Word → PDF (2 ไฟล์)", f"{src1.name}, {src2.name}",
                      f"{pdf1.name}({p1n}น.), {pdf2.name}({p2n}น.)",
                      "ทั้งสองไฟล์ดาวน์โหลดได้ เปิดด้วย pymupdf นับหน้าได้", "ผ่าน"))

        # ② รวม PDF (ใช้ผลลัพธ์จริงจากขั้นก่อน ไม่ใช่ตัวอย่างเดิม)
        goto(pg, "pdf-merge")
        pg.locator(".dz input[type=file]").first.set_input_files(
            [str(pdf1), str(pdf2)], timeout=STEP_TIMEOUT)
        pg.wait_for_selector(".file-row", timeout=STEP_TIMEOUT)
        assert pg.locator(".file-row").count() == 2, "ไฟล์ที่ใส่ไม่ครบ 2 แถว"
        pg.get_by_role("button", name="รวมไฟล์").click()
        pg.wait_for_selector(".results .result", timeout=STEP_TIMEOUT)
        merged = dl(pg, lambda: pg.locator(".results .result button").first.click(), "c1_merged.pdf")
        mdoc = fitz.open(str(merged))
        expect_pages = p1n + p2n
        ok_pages = mdoc.page_count == expect_pages
        steps.append(("รวม PDF", f"{pdf1.name} + {pdf2.name}", merged.name,
                      f"จำนวนหน้ารวม = {expect_pages}", "ผ่าน" if ok_pages else f"ตก — ได้ {mdoc.page_count} หน้า"))
        assert ok_pages, f"รวมแล้วหน้าไม่ครบ: ได้ {mdoc.page_count} ต้องการ {expect_pages}"

        # ③ บีบอัด PDF
        goto(pg, "pdf-compress")
        pg.locator(".dz input[type=file]").first.set_input_files(str(merged), timeout=STEP_TIMEOUT)
        pg.wait_for_selector(".file-row", timeout=STEP_TIMEOUT)
        pg.get_by_role("button", name="บีบอัดไฟล์").click()
        pg.wait_for_selector(".results .result", timeout=STEP_TIMEOUT)
        final = dl(pg, lambda: pg.locator(".results .result button").first.click(), "c1_final.pdf")

        # ④ เปิดไฟล์สุดท้ายด้วย pymupdf ตรวจจำนวนหน้า + มีข้อความไทย
        fdoc = fitz.open(str(final))
        pages_ok = fdoc.page_count == expect_pages
        thai_found = any(re.search(r"[ก-๙]", fdoc[i].get_text()) for i in range(fdoc.page_count))
        steps.append(("บีบอัด PDF + ตรวจปลายทาง", merged.name, final.name,
                      f"จำนวนหน้า={expect_pages}, มีข้อความไทย", "ผ่าน" if (pages_ok and thai_found) else "ตก"))
        assert pages_ok, f"บีบอัดแล้วหน้าหาย: ได้ {fdoc.page_count} ต้องการ {expect_pages}"
        assert thai_found, "ไฟล์สุดท้ายไม่มีข้อความไทยเลยหลังผ่าน 3 ขั้นตอน (บีบอัดอาจทำลายชั้นข้อความ)"

        real = real_errors(errs)
        assert not real, f"มี JS error ระหว่างโซ่: {real[0][:150]}"
        return {"ok": True, "steps": steps, "note": ""}
    except Exception as e:
        steps.append(("EXCEPTION", "-", "-", "-", f"ตก: {e}"))
        return {"ok": False, "steps": steps, "note": str(e)}
    finally:
        pg.close()


def chain2_excel_csv_encoding(browser):
    """Excel → CSV → ซ่อมไฟล์ไทยเพี้ยน → ตรวจว่าไทยยังอ่านออก"""
    steps = []
    pg, errs = new_page(browser)
    try:
        src = SAMPLES / "ตัวอย่าง-ข้อมูลใบเสนอราคา.xlsx"
        wb = openpyxl.load_workbook(src)
        expected_customer = str(wb.active["E2"].value or "")  # ชื่อลูกค้า

        goto(pg, "excel-csv")
        upload(pg, src)
        pg.wait_for_selector(".file-row", timeout=STEP_TIMEOUT)
        pg.get_by_role("button", name="แปลงไฟล์").click()
        pg.wait_for_selector(".results .result", timeout=STEP_TIMEOUT)
        csv1 = dl(pg, lambda: pg.locator(".results .result button").first.click(), "c2_out.csv")
        raw = csv1.read_bytes()
        has_thai_1 = bool(re.search(r"[ก-๙]", raw.decode("utf-8-sig", errors="replace")))
        steps.append(("Excel → CSV", src.name, csv1.name,
                      "ดาวน์โหลด CSV ได้ มีข้อความไทย", "ผ่าน" if has_thai_1 else "ตก"))
        assert has_thai_1, "CSV ที่แปลงจาก Excel ไม่มีข้อความไทยเลย"

        # ซ่อมไฟล์ไทยเพี้ยน — ป้อน CSV จริงที่เพิ่งดาวน์โหลด (อยู่แล้วเป็น UTF-8 ปกติ)
        goto(pg, "thai-encoding")
        upload(pg, csv1)
        pg.wait_for_selector(".enc-card", timeout=STEP_TIMEOUT)
        card_text = pg.locator(".enc-card").inner_text()
        fixed = dl(pg, lambda: pg.get_by_role("button", name="บันทึกเป็น UTF-8").click(), "c2_fixed.csv")
        data = fixed.read_bytes()
        text = data.decode("utf-8-sig", errors="replace")
        has_thai_2 = bool(re.search(r"[ก-๙]", text))
        no_mojibake = "�" not in text and expected_customer in text
        steps.append(("ซ่อมไฟล์ไทยเพี้ยน", csv1.name, fixed.name,
                      f"ตรวจพบ encoding: {card_text[:40]!r}; ไทยยังอ่านออก ไม่กลายเป็นตัวประหลาด",
                      "ผ่าน" if (has_thai_2 and no_mojibake) else "ตก"))
        assert has_thai_2, "ผ่านเครื่องมือซ่อมไทยแล้วไม่มีข้อความไทยเหลือเลย"
        assert no_mojibake, f"ข้อความไทยเพี้ยน/หายหลังผ่านเครื่องมือซ่อม (คาดพบ {expected_customer!r})"

        real = real_errors(errs)
        assert not real, f"มี JS error ระหว่างโซ่: {real[0][:150]}"
        return {"ok": True, "steps": steps, "note": ""}
    except Exception as e:
        steps.append(("EXCEPTION", "-", "-", "-", f"ตก: {e}"))
        return {"ok": False, "steps": steps, "note": str(e)}
    finally:
        pg.close()


def chain3_image_pdf_split_image(browser):
    """รูปภาพ → PDF → แยกหน้า PDF → PDF → รูปภาพ → เปิดด้วย PIL"""
    steps = []
    pg, errs = new_page(browser)
    try:
        img1 = SAMPLES / "ตัวอย่าง-รูปภาพ-1.jpg"
        img2 = SAMPLES / "ตัวอย่าง-รูปภาพ-2.png"

        goto(pg, "images-to-pdf")
        pg.locator(".dz input[type=file]").first.set_input_files(
            [str(img1), str(img2)], timeout=STEP_TIMEOUT)
        pg.wait_for_selector(".file-row", timeout=STEP_TIMEOUT)
        assert pg.locator(".file-row").count() == 2, "ไฟล์รูปที่ใส่ไม่ครบ 2 แถว"
        pg.get_by_role("button", name="สร้างไฟล์ PDF").click()
        pg.wait_for_selector(".results .result", timeout=STEP_TIMEOUT)
        pdf1 = dl(pg, lambda: pg.locator(".results .result button").first.click(), "c3_imgpdf.pdf")
        pdoc = fitz.open(str(pdf1))
        steps.append(("รูปภาพ → PDF", f"{img1.name}, {img2.name}", pdf1.name,
                      f"ได้ PDF {pdoc.page_count} หน้า (2 รูป)", "ผ่าน" if pdoc.page_count == 2 else "ตก"))
        assert pdoc.page_count == 2, f"ควรได้ PDF 2 หน้า ได้ {pdoc.page_count}"

        # แยกหน้า PDF — ดึงเฉพาะหน้า 1 (ช่วงหน้า ทำให้ได้ผลลัพธ์ไฟล์เดียว ไม่ใช่ ZIP)
        goto(pg, "pdf-split")
        upload(pg, pdf1)
        pg.wait_for_selector(".file-row", timeout=STEP_TIMEOUT)
        pg.locator('input[type=text]').first.fill("1")
        pg.wait_for_selector(".sp-count", timeout=STEP_TIMEOUT)
        assert "1 ไฟล์" in pg.locator(".sp-count").inner_text(), \
            f"ตัวอย่างจำนวนไฟล์ที่จะได้ไม่ตรง: {pg.locator('.sp-count').inner_text()}"
        pg.get_by_role("button", name="แยกไฟล์").click()
        pg.wait_for_selector(".results .result", timeout=STEP_TIMEOUT)
        split1 = dl(pg, lambda: pg.locator(".results .result button").first.click(), "c3_split_p1.pdf")
        sdoc = fitz.open(str(split1))
        steps.append(("แยกหน้า PDF (เอาแค่หน้า 1)", pdf1.name, split1.name,
                      "ได้ไฟล์เดียว 1 หน้า", "ผ่าน" if sdoc.page_count == 1 else "ตก"))
        assert sdoc.page_count == 1, f"แยกหน้าแล้วควรได้ 1 หน้า ได้ {sdoc.page_count}"

        # PDF → รูปภาพ
        goto(pg, "pdf-to-images")
        upload(pg, split1)
        pg.wait_for_selector(".file-row", timeout=STEP_TIMEOUT)
        pg.get_by_role("button", name="แปลงเป็นรูป").click()
        pg.wait_for_selector(".results .result", timeout=STEP_TIMEOUT)
        # ‼️ PDF ต้นทางมีแค่ 1 หน้า (เอามาจากขั้นแยกหน้าก่อนหน้า) → ผลลัพธ์มีรูปเดียว
        #    src/tools/pdf-to-images.js:65 โชว์ปุ่ม "ดาวน์โหลด ZIP" เฉพาะตอน made.length > 1 เท่านั้น
        #    รูปเดียวจึงมีแค่ปุ่มดาวน์โหลดเดี่ยวใน .results .result (ไม่ใช่ ZIP) — ไม่ใช่บั๊ก เป็นพฤติกรรมตั้งใจ
        img_out = dl(pg, lambda: pg.locator(".results .result button").first.click(), "c3_out.png")
        im = Image.open(img_out)
        im.load()
        ok_open = im.size[0] > 0 and im.size[1] > 0
        steps.append(("PDF → รูปภาพ", split1.name, img_out.name,
                      "รูปสุดท้ายเปิดได้ด้วย PIL ขนาด > 0", "ผ่าน" if ok_open else "ตก"))
        assert ok_open, "เปิดรูปผลลัพธ์สุดท้ายด้วย PIL ไม่ได้ หรือขนาดผิดปกติ"

        real = real_errors(errs)
        assert not real, f"มี JS error ระหว่างโซ่: {real[0][:150]}"
        return {"ok": True, "steps": steps, "note": ""}
    except Exception as e:
        steps.append(("EXCEPTION", "-", "-", "-", f"ตก: {e}"))
        return {"ok": False, "steps": steps, "note": str(e)}
    finally:
        pg.close()


def chain4_pdf_to_word(browser):
    """PDF → Word → เปิดด้วย python-docx ตรวจย่อหน้า"""
    steps = []
    pg, errs = new_page(browser)
    try:
        src = SAMPLES / "ตัวอย่าง-รายงานประจำเดือน.pdf"
        goto(pg, "pdf-to-word")
        upload(pg, src)
        pg.wait_for_selector(".file-row", timeout=STEP_TIMEOUT)
        pg.get_by_role("button", name="แปลงเป็น Word").click()
        pg.wait_for_selector(".results .result", timeout=STEP_TIMEOUT)
        out = dl(pg, lambda: pg.locator(".results .result button").first.click(), "c4_out.docx")

        doc = pydocx.Document(str(out))
        paras = [p.text for p in doc.paragraphs if p.text.strip()]
        full_text = "\n".join(paras)
        has_thai = bool(re.search(r"[ก-๙]", full_text))
        steps.append(("PDF → Word", src.name, out.name,
                      f"เปิดด้วย python-docx: {len(paras)} ย่อหน้าไม่ว่าง, มีข้อความไทย",
                      "ผ่าน" if (len(paras) >= 1 and has_thai) else "ตก"))
        assert len(paras) >= 1, "เปิดไฟล์ Word ที่แปลงแล้วไม่พบย่อหน้าที่มีเนื้อหาเลย"
        assert has_thai, "เอกสาร Word ที่แปลงแล้วไม่มีข้อความไทยเลย"

        real = real_errors(errs)
        assert not real, f"มี JS error ระหว่างโซ่: {real[0][:150]}"
        return {"ok": True, "steps": steps, "note": ""}
    except Exception as e:
        steps.append(("EXCEPTION", "-", "-", "-", f"ตก: {e}"))
        return {"ok": False, "steps": steps, "note": str(e)}
    finally:
        pg.close()


def chain5_ppt_pdf_watermark_sign(browser):
    """PowerPoint → PDF → ใส่ลายน้ำ → เซ็นชื่อ → ตรวจหน้าไม่หาย"""
    steps = []
    pg, errs = new_page(browser)
    try:
        src = SAMPLES / "ตัวอย่าง-นำเสนอบริษัท.pptx"
        with zipfile.ZipFile(src) as z:
            n_slides = len([n for n in z.namelist() if re.match(r"ppt/slides/slide\d+\.xml$", n)])

        goto(pg, "powerpoint-to-pdf")
        upload(pg, src)
        pg.wait_for_selector(".file-row", timeout=STEP_TIMEOUT)
        pg.get_by_role("button", name="สร้างไฟล์ PDF").click()
        pg.wait_for_selector(".results .result", timeout=STEP_TIMEOUT)
        pdf1 = dl(pg, lambda: pg.locator(".results .result button").first.click(), "c5_ppt.pdf")
        p1 = fitz.open(str(pdf1)).page_count
        steps.append(("PowerPoint → PDF", src.name, pdf1.name,
                      f"จำนวนหน้า = จำนวนสไลด์ ({n_slides})", "ผ่าน" if p1 == n_slides else f"ตก — ได้ {p1} หน้า"))
        assert p1 == n_slides, f"หน้า PDF ไม่ตรงจำนวนสไลด์: ได้ {p1} ต้องการ {n_slides}"

        goto(pg, "pdf-watermark")
        upload(pg, pdf1)
        pg.wait_for_selector(".file-row", timeout=STEP_TIMEOUT)
        pg.get_by_role("button", name="ใส่ลายน้ำ").click()
        pg.wait_for_selector(".results .result", timeout=STEP_TIMEOUT)
        pdf2 = dl(pg, lambda: pg.get_by_role("button", name="ดาวน์โหลด").click(), "c5_wm.pdf")
        p2 = fitz.open(str(pdf2)).page_count
        steps.append(("ใส่ลายน้ำ PDF", pdf1.name, pdf2.name,
                      f"จำนวนหน้าไม่เปลี่ยน ({n_slides})", "ผ่าน" if p2 == n_slides else f"ตก — ได้ {p2} หน้า"))
        assert p2 == n_slides, f"ใส่ลายน้ำแล้วหน้าหาย: ได้ {p2} ต้องการ {n_slides}"

        # เซ็นชื่อ — อัปโหลดรูปเป็นลายเซ็น แล้วคลิกวางบนหน้าที่ 1
        goto(pg, "pdf-sign")
        upload(pg, pdf2)
        pg.wait_for_selector(".sign-page", timeout=STEP_TIMEOUT)
        pg.wait_for_timeout(500)  # ให้ pdf.js เรนเดอร์หน้าแรกเสร็จก่อนวัดตำแหน่ง canvas
        pg.locator("input[type=file][accept='image/*']").set_input_files(
            str(SAMPLES / "ตัวอย่าง-รูปภาพ-2.png"), timeout=STEP_TIMEOUT)
        pg.wait_for_timeout(500)
        pg.locator(".sign-stage").click(position={"x": 120, "y": 120})
        pg.wait_for_selector(".sign-item", timeout=STEP_TIMEOUT)
        pg.get_by_role("button", name="บันทึกไฟล์เซ็นแล้ว").click()
        pg.wait_for_selector(".results .result", timeout=STEP_TIMEOUT)
        pdf3 = dl(pg, lambda: pg.get_by_role("button", name="ดาวน์โหลด").click(), "c5_signed.pdf")
        p3 = fitz.open(str(pdf3)).page_count
        steps.append(("เซ็นชื่อบน PDF", pdf2.name, pdf3.name,
                      f"จำนวนหน้าไม่เปลี่ยน ({n_slides})", "ผ่าน" if p3 == n_slides else f"ตก — ได้ {p3} หน้า"))
        assert p3 == n_slides, f"เซ็นชื่อแล้วหน้าหาย: ได้ {p3} ต้องการ {n_slides}"

        real = real_errors(errs)
        assert not real, f"มี JS error ระหว่างโซ่: {real[0][:150]}"
        return {"ok": True, "steps": steps, "note": ""}
    except Exception as e:
        steps.append(("EXCEPTION", "-", "-", "-", f"ตก: {e}"))
        return {"ok": False, "steps": steps, "note": str(e)}
    finally:
        pg.close()


def _make_address_xlsx():
    """สร้างไฟล์ Excel ที่มีคอลัมน์ที่อยู่ไทยรวมกัน (ตัวอย่างในโปรเจกต์ไม่มีคอลัมน์นี้เลย)"""
    path = GEN_DIR / "c6_address_src.xlsx"
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "ที่อยู่ลูกค้า"
    ws.append(["ชื่อ", "ที่อยู่"])
    ws.append(["สมชาย ใจดี", "123 หมู่ 4 ต.สุเทพ อ.เมืองเชียงใหม่ จ.เชียงใหม่ 50200"])
    ws.append(["วิภาวรรณ ศรีสุข", "45/6 ถ.สุขุมวิท แขวงบางนา เขตบางนา กรุงเทพมหานคร 10260"])
    wb.save(path)
    return path


def _make_mailmerge_template():
    """เทมเพลต Word ที่ใช้ตัวยึดตรงกับคอลัมน์ผลลัพธ์ของ thai-address ทุกตัว"""
    path = GEN_DIR / "c6_template.docx"
    doc = pydocx.Document()
    doc.add_paragraph("เรียนคุณ {{ชื่อ}}")
    doc.add_paragraph(
        "ที่อยู่ของท่านในระบบคือ ตำบล/แขวง {{ตำบล/แขวง}} อำเภอ/เขต {{อำเภอ/เขต}} "
        "จังหวัด {{จังหวัด}} รหัสไปรษณีย์ {{รหัสไปรษณีย์}}"
    )
    doc.save(path)
    return path


def chain6_excel_address_mailmerge(browser):
    """Excel → แยกที่อยู่ไทย → จดหมายเวียน Word → ตรวจว่ามีที่อยู่ที่แยกแล้วจริง"""
    steps = []
    pg, errs = new_page(browser)
    try:
        src = _make_address_xlsx()
        tpl = _make_mailmerge_template()

        goto(pg, "thai-address")
        upload(pg, src)
        pg.wait_for_selector(".stats .stat", timeout=STEP_TIMEOUT)
        chips = pg.locator(".stats").inner_text()
        out = dl(pg, lambda: pg.get_by_role("button", name="ดาวน์โหลดเป็น Excel").click(), "c6_split.xlsx")

        wbo = openpyxl.load_workbook(out)
        wso = wbo.active
        header = [c.value for c in wso[1]]
        need_cols = ["ตำบล/แขวง", "อำเภอ/เขต", "จังหวัด", "รหัสไปรษณีย์"]
        cols_ok = all(c in header for c in need_cols)
        idx = {c: header.index(c) for c in need_cols if c in header}
        row2 = [r for r in wso.iter_rows(min_row=2, max_row=2, values_only=True)][0]
        prov_ok = cols_ok and row2[idx["จังหวัด"]] == "เชียงใหม่"
        zip_ok = cols_ok and str(row2[idx["รหัสไปรษณีย์"]]) == "50200"
        steps.append(("แยกที่อยู่ไทย", f"{src.name} ({chips.strip()[:30]})", out.name,
                      "มีคอลัมน์ ตำบล/อำเภอ/จังหวัด/รหัสไปรษณีย์ และค่าตรงกับที่อยู่จริงแถวแรก",
                      "ผ่าน" if (cols_ok and prov_ok and zip_ok) else "ตก"))
        assert cols_ok, f"ไม่พบคอลัมน์ที่อยู่แยกแล้วครบ 4 ช่อง: {header}"
        assert prov_ok, f"จังหวัดแถวแรกไม่ตรง: ได้ {row2[idx['จังหวัด']]!r} ต้องการ 'เชียงใหม่'"
        assert zip_ok, f"รหัสไปรษณีย์แถวแรกไม่ตรง: ได้ {row2[idx['รหัสไปรษณีย์']]!r} ต้องการ '50200'"

        goto(pg, "word-mailmerge")
        pg.locator(".dz input[type=file]").nth(0).set_input_files(str(tpl), timeout=STEP_TIMEOUT)
        pg.wait_for_selector(".mm-step", timeout=STEP_TIMEOUT)
        pg.locator(".dz input[type=file]").nth(1).set_input_files(str(out), timeout=STEP_TIMEOUT)
        pg.wait_for_timeout(1200)
        locked4 = pg.locator(".mm-step").nth(3).get_attribute("data-locked")
        steps_note = f"ขั้น 4 ปลดล็อกหลังจับคู่คอลัมน์อัตโนมัติ (data-locked={locked4})"
        assert locked4 == "0", f"จับคู่คอลัมน์อัตโนมัติไม่สำเร็จ — ขั้นที่ 4 ยังล็อกอยู่ (data-locked={locked4})"
        pg.get_by_role("button", name="สร้างเอกสารทั้งชุด").click()
        pg.wait_for_selector(".results .result", timeout=STEP_TIMEOUT)
        merged_docx = dl(pg, lambda: pg.locator(".results .result button").first.click(), "c6_letter.docx")

        doc = pydocx.Document(str(merged_docx))
        full_text = "\n".join(p.text for p in doc.paragraphs)
        has_name = "สมชาย ใจดี" in full_text
        has_addr = "เชียงใหม่" in full_text and "50200" in full_text and "สุเทพ" in full_text
        no_leftover_tag = "{{" not in full_text
        steps.append(("จดหมายเวียน Word", f"{tpl.name} + {out.name}", merged_docx.name,
                      "เอกสารมีชื่อและที่อยู่ที่แยกแล้วจริง (ตำบล/จังหวัด/รหัสไปรษณีย์) ไม่เหลือตัวยึดดิบ",
                      "ผ่าน" if (has_name and has_addr and no_leftover_tag) else "ตก"))
        assert has_name, f"ไม่พบชื่อ 'สมชาย ใจดี' ในจดหมายที่สร้าง: {full_text[:200]!r}"
        assert has_addr, f"ไม่พบข้อมูลที่อยู่ที่แยกแล้วในจดหมาย: {full_text[:300]!r}"
        assert no_leftover_tag, f"ยังเหลือตัวยึด {{...}} ดิบในเอกสาร (จับคู่คอลัมน์ไม่ครบ): {full_text[:300]!r}"

        real = real_errors(errs)
        assert not real, f"มี JS error ระหว่างโซ่: {real[0][:150]}"
        return {"ok": True, "steps": steps, "note": ""}
    except Exception as e:
        steps.append(("EXCEPTION", "-", "-", "-", f"ตก: {e}"))
        return {"ok": False, "steps": steps, "note": str(e)}
    finally:
        pg.close()


def chain7_next_card_navigation(browser):
    """โซ่ที่เว็บแนะนำเอง — กดปุ่มใน 'ทำอะไรต่อดี' แล้วต้องเข้าเครื่องมือถัดไปได้ + ตรวจว่าไฟล์ตามไปด้วยไหม"""
    steps = []
    pg, errs = new_page(browser)
    limitations = []
    try:
        files = [SAMPLES / "ตัวอย่าง-รายงานประจำเดือน.pdf", SAMPLES / "ตัวอย่าง-ใบปะหน้าเอกสาร.pdf"]
        goto(pg, "pdf-merge")
        pg.locator(".dz input[type=file]").first.set_input_files([str(f) for f in files], timeout=STEP_TIMEOUT)
        pg.wait_for_selector(".file-row", timeout=STEP_TIMEOUT)
        pg.get_by_role("button", name="รวมไฟล์").click()
        pg.wait_for_selector(".results .result", timeout=STEP_TIMEOUT)

        n_next = pg.locator(".next-card").count()
        hrefs = pg.eval_on_selector_all(".next-card", "els => els.map(e => e.getAttribute('href'))")
        steps.append(("ทำ pdf-merge จนมีผลลัพธ์", f"{files[0].name}, {files[1].name}", "ไฟล์รวมแล้ว (ในหน่วยความจำ)",
                      f"แถว 'ทำอะไรต่อดี' มี ≥2 ตัวเลือก, ไม่ชี้กลับหาตัวเอง", "ผ่าน" if n_next >= 2 else "ตก"))
        assert n_next >= 2, f"แถว 'ทำอะไรต่อดี' มีแค่ {n_next} ตัวเลือก"
        assert "#/pdf-merge" not in hrefs, "แนะนำวนกลับหาตัวเอง"

        next_id = hrefs[0].replace("#/", "")
        pg.locator(".next-card").first.click()
        pg.wait_for_url(re.compile(re.escape(f"#/{next_id}") + "$"), timeout=STEP_TIMEOUT)
        pg.wait_for_selector(".tool-head", timeout=STEP_TIMEOUT)
        landed_ok = next_id in TOOL_IDS and pg.locator(".tool-head").count() == 1
        steps.append((f"คลิก next-card ตัวแรก → {next_id}", "-", "-",
                      "เข้าเครื่องมือถัดไปได้จริง หน้าเครื่องมือขึ้นครบ", "ผ่าน" if landed_ok else "ตก"))
        assert landed_ok, f"คลิกแล้วไม่เข้าเครื่องมือ {next_id} ให้ถูกต้อง"

        # ‼️ ไฟล์ผลลัพธ์ต้องตามไปเครื่องมือถัดไปเองโดยไม่ต้องดาวน์โหลดลงเครื่องก่อน
        #    (เดิมประกาศผลลัพธ์ตอน download() เท่านั้น ป้าย "พาไฟล์ไปด้วย" จึงไม่เคยโผล่
        #     ก่อนกดดาวน์โหลด = ฟีเจอร์นี้ใช้ไม่ได้จริงทั้งที่เขียนไว้แล้ว แก้ 09/09/2026
        #     ด้วย downloadButton() ใน src/ui.js ที่ประกาศตั้งแต่ตอนสร้างปุ่ม)
        #    ต้องรอจริง — เครื่องมือปลายทาง mount แบบ async แล้วค่อยหยิบไฟล์ที่ฝากไว้ใน microtask
        try:
            pg.wait_for_selector(".file-row", timeout=6000)
        except Exception:
            pass
        carried = pg.locator(".file-row").count() > 0 or pg.locator(".results .result").count() > 0
        steps.append(("ตรวจว่าไฟล์ตามไปด้วยไหม", "-", "-",
                      "ไฟล์ผลลัพธ์ต้องตามไปโดยไม่ต้องดาวน์โหลดก่อน",
                      "ผ่าน (ไฟล์ตามไปด้วย)" if carried else "ตก"))
        assert carried, (f"กด 'ทำอะไรต่อดี' ไป {next_id} แล้วไฟล์ผลลัพธ์ไม่ตามไปด้วย "
                         "(ป้าย 'พาไฟล์ไปด้วย' โชว์แต่ของไม่ไปจริง)")

        real = real_errors(errs)
        assert not real, f"มี JS error ระหว่างการนำทาง: {real[0][:150]}"
        return {"ok": True, "steps": steps, "note": "; ".join(limitations)}
    except Exception as e:
        steps.append(("EXCEPTION", "-", "-", "-", f"ตก: {e}"))
        return {"ok": False, "steps": steps, "note": str(e)}
    finally:
        pg.close()


CHAINS = [
    ("① Word → PDF → รวม PDF → บีบอัด PDF", chain1_word_pdf_merge_compress),
    ("② Excel → CSV → ซ่อมไฟล์ไทยเพี้ยน", chain2_excel_csv_encoding),
    ("③ รูปภาพ → PDF → แยกหน้า PDF → PDF → รูปภาพ", chain3_image_pdf_split_image),
    ("④ PDF → Word", chain4_pdf_to_word),
    ("⑤ PowerPoint → PDF → ลายน้ำ → เซ็นชื่อ", chain5_ppt_pdf_watermark_sign),
    ("⑥ Excel → แยกที่อยู่ไทย → จดหมายเวียน Word", chain6_excel_address_mailmerge),
    ("⑦ โซ่ที่เว็บแนะนำเอง (next-card)", chain7_next_card_navigation),
]


def main():
    log(f"FileKit chain test — BASE={BASE}")
    log(f"WORKDIR={WORKDIR}")
    server_proc = start_server_if_needed()
    results = []
    try:
        with sync_playwright() as p:
            browser = p.chromium.launch()
            try:
                for name, fn in CHAINS:
                    log(f"\n▶ {name}")
                    t0 = time.time()
                    r = fn(browser)
                    r["dur"] = time.time() - t0
                    r["name"] = name
                    results.append(r)
                    for step in r["steps"]:
                        mark = "✅" if step[4].startswith("ผ่าน") or step[4].startswith("ไฟล์ตามไปด้วย") or step[4].startswith("ไม่ตามไป") else "❌"
                        log(f"  {mark} {step[0]}: {step[4]}")
                    log(f"  {'✅ ผ่านทั้งโซ่' if r['ok'] else '❌ ตก'} — ใช้เวลา {r['dur']:.1f}s"
                        + (f" — {r['note']}" if r["note"] else ""))
            finally:
                browser.close()
    finally:
        if server_proc:
            server_proc.terminate()
            try:
                server_proc.wait(timeout=5)
            except subprocess.TimeoutExpired:
                server_proc.kill()
        shutil.rmtree(WORKDIR, ignore_errors=True)

    # ── ตารางโซ่ ─────────────────────────────────────────────────────────
    log("\n" + "═" * 90)
    log("ตารางโซ่ — ขั้นตอน · ไฟล์เข้า/ออก · สิ่งที่ตรวจ · ผล")
    log("═" * 90)
    for r in results:
        log(f"\n{r['name']}  ({r['dur']:.1f}s)  — {'ผ่าน' if r['ok'] else 'ตก'}")
        for i, (step, fin, fout, checked, res) in enumerate(r["steps"], 1):
            log(f"  {i}. {step}")
            log(f"     เข้า: {fin}")
            log(f"     ออก: {fout}")
            log(f"     ตรวจ: {checked}")
            log(f"     ผล: {res}")

    passed = sum(1 for r in results if r["ok"])
    log("\n" + "═" * 90)
    log(f"สรุป: ผ่าน {passed}/{len(results)} โซ่")
    for r in results:
        if r["note"]:
            log(f"  หมายเหตุ [{r['name']}]: {r['note']}")
    log("═" * 90)

    sys.exit(0 if passed == len(results) else 1)


if __name__ == "__main__":
    main()
